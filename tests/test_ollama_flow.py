"""
Integration and regression test suite for Ollama analysis flow and timeout handling.

This script tests:
1. Direct live Ollama suggestion generation with configured timeout, token bounds, and robust JSON recovery.
2. Background analysis job execution via FastAPI (/api/analyze) with polling and completion.
3. Fallback behavior when Ollama times out or encounters slow generation.

Run directly:
    python tests/test_ollama_flow.py
Or with pytest:
    pytest tests/test_ollama_flow.py
"""

import os
import sys
import time
import io
import json
import pytest
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from backend import main

client = TestClient(main.app)

SAMPLE_RESUME = """Jane Doe
Senior Software Engineer
San Francisco, CA | jane@example.com

Professional Summary
Experienced Software Engineer with 6 years building distributed systems and scalable web services in Python and Go.

Skills
Python, Go, Docker, Kubernetes, PostgreSQL, Git

Experience
Senior Software Engineer - Tech Solutions (2021 - Present)
- Designed and maintained high-throughput REST APIs handling 50M requests daily.
- Optimized database queries in PostgreSQL, reducing query latency by 30%.
- Led containerization migration to Kubernetes on AWS.
"""

SAMPLE_JD = """Senior Backend Engineer
We are looking for a Senior Backend Engineer to join our cloud platform team.
Requirements:
- 5+ years of experience with Python or Go.
- Experience with FastAPI, Docker, and Kubernetes.
- Strong knowledge of PostgreSQL and distributed caching (Redis).
- Experience designing resilient microservices.
"""


def create_sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Senior Software Engineer\nSan Francisco, CA | jane@example.com")
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph("Experienced Software Engineer with 6 years building distributed systems and scalable web services in Python and Go.")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, Go, Docker, Kubernetes, PostgreSQL, Git")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Senior Software Engineer - Tech Solutions (2021 - Present)")
    doc.add_paragraph("Designed and maintained high-throughput REST APIs handling 50M requests daily.")
    doc.add_paragraph("Optimized database queries in PostgreSQL, reducing query latency by 30%.")
    doc.add_paragraph("Led containerization migration to Kubernetes on AWS.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def is_ollama_available() -> bool:
    try:
        import urllib.request
        req = urllib.request.Request("http://localhost:11434/api/tags")
        with urllib.request.urlopen(req, timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            models = [m.get("name", "") for m in data.get("models", [])]
            return any("llama3.2" in m for m in models)
    except Exception:
        return False


def test_ollama_timeout_configuration():
    """Verify timeout and token limits are configurable and have safe defaults."""
    assert main.OLLAMA_TIMEOUT_SECONDS >= 300, "Ollama timeout should be at least 300 seconds"
    assert main.OLLAMA_MAX_RESPONSE_TOKENS >= 512, "Max response tokens should be at least 512"


def test_truncated_json_recovery():
    """Verify that truncated or unclosed JSON from Ollama is gracefully recovered."""
    truncated_raw = '''{
  "suggestions": [
    {
      "section": "Skills",
      "original_content": "Python, Go, Docker",
      "proposed_content": "Python, Go, FastAPI, Docker, Kubernetes",
      "why_this_may_help": "Adds FastAPI and Kubernetes.",
      "job_requirement": "FastAPI, Kubernetes",
      "expected_impact": "Improves skill alignment"
    },
    {
      "section": "Experience",
      "original_content": "Designed and maintained high-throughput REST APIs",
      "proposed_content": "Designed resilient microservices with FastAPI",
      "why_this_may_help": "Highlights resilient microservices.",
      "job_requirement": "resilient microservices",
      "expected_impact": "Direct match for cloud platform"
    },
    {
      "section": "Professional Summary",
      "original_content": "Experienced Software Engineer",
      "proposed_content": "Results-oriented Senior Backend Engineer",
      "why_this_may_help": "Emphasizes senior backend focus.",
      "job_requirement": "Senior Backend Engineer",
      "expected_impact": "Enhances overall profile appeal
'''
    recovered = main._extract_suggestions_from_ollama_response(truncated_raw)
    assert len(recovered) >= 2, f"Expected at least 2 recovered suggestions, got {len(recovered)}"
    assert recovered[0]["section"] == "Skills"
    assert recovered[1]["section"] == "Experience"


def test_background_analysis_fallback_on_timeout(monkeypatch):
    """Verify background job gracefully falls back to manual suggestions on timeout without failing with 504."""
    orig_call = main.call_ollama_suggestions

    def fake_timeout(*args, **kwargs):
        raise main.HTTPException(
            status_code=504,
            detail=f"The local Ollama model did not finish generating within {main.OLLAMA_TIMEOUT_SECONDS} seconds."
        )

    monkeypatch.setattr(main, "call_ollama_suggestions", fake_timeout)

    try:
        docx_bytes = create_sample_docx_bytes()
        resp = client.post(
            "/api/analyze",
            files={"file": ("test_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"job_description": SAMPLE_JD, "background": "true"},
        )
        assert resp.status_code == 200, resp.text
        job_id = resp.json()["job_id"]
        workflow_id = resp.json()["workflow_id"]

        # Poll job until finished
        completed = False
        for _ in range(30):
            job_res = client.get(f"/api/jobs/{job_id}")
            assert job_res.status_code == 200
            data = job_res.json()
            if data["status"] == "completed":
                completed = True
                assert data["progress"] == 100
                assert "manual suggestions" in data["message"].lower() or "completed" in data["message"].lower()
                break
            elif data["status"] == "failed":
                pytest.fail(f"Job unexpectedly failed with: {data.get('error')}")
            time.sleep(0.1)

        assert completed, "Background job should complete within polling window"

        # Verify workflow data has fallback suggestions and clear notice
        wf_res = client.get(f"/api/workflows/{workflow_id}")
        assert wf_res.status_code == 200
        wf = wf_res.json()
        assert wf["status"] == "completed"
        assert len(wf["suggestions"]) > 0
    finally:
        monkeypatch.setattr(main, "call_ollama_suggestions", orig_call)
        try:
            if "workflow_id" in locals():
                client.delete(f"/api/workflows/{workflow_id}")
        except Exception:
            pass


def test_targeted_micro_prompt_single_section(monkeypatch):
    """Verify call_ollama_single_section parses valid micro-suggestions and enforces budget."""
    orig_urlopen = main.urlopen
    mock_payload = {
        "response": json.dumps({
            "proposed_content": "Python, FastAPI, Kubernetes, PostgreSQL",
            "why_this_may_help": "Emphasizes modern Python frameworks.",
            "job_requirement": "FastAPI",
            "expected_impact": "Direct match for cloud requirements."
        })
    }
    class MockResp:
        def read(self):
            return json.dumps(mock_payload).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(main, "urlopen", lambda *args, **kwargs: MockResp())
    try:
        result, warnings = main.call_ollama_single_section(
            section="Skills",
            original_text="Python, Go, Docker, Kubernetes, PostgreSQL, Git",
            job_description=SAMPLE_JD,
        )
        assert result is not None
        assert result["section"] == "Skills"
        assert "FastAPI" in result["proposed_text"]
        assert len(result["proposed_text"]) <= len("Python, Go, Docker, Kubernetes, PostgreSQL, Git")
    finally:
        monkeypatch.setattr(main, "urlopen", orig_urlopen)


def test_tailor_section_endpoint(monkeypatch):
    """Verify POST /api/workflows/{workflow_id}/tailor-section updates the specific section."""
    orig_urlopen = main.urlopen
    mock_payload = {
        "response": json.dumps({
            "proposed_content": "Python, FastAPI, Kubernetes, PostgreSQL",
            "why_this_may_help": "Emphasizes modern Python frameworks.",
            "job_requirement": "FastAPI",
            "expected_impact": "Direct match for cloud requirements."
        })
    }
    class MockResp:
        def read(self):
            return json.dumps(mock_payload).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(main, "urlopen", lambda *args, **kwargs: MockResp())
    try:
        docx_bytes = create_sample_docx_bytes()
        resp = client.post(
            "/api/analyze",
            files={"file": ("test_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"job_description": SAMPLE_JD},
        )
        assert resp.status_code == 200
        wf_id = resp.json()["id"]

        tailor_res = client.post(
            f"/api/workflows/{wf_id}/tailor-section",
            json={"section": "Skills", "custom_instruction": "Highlight FastAPI"},
        )
        assert tailor_res.status_code == 200, tailor_res.text
        updated = tailor_res.json()
        skills_sug = next(s for s in updated["suggestions"] if s["section"] == "Skills")
        assert "FastAPI" in skills_sug["proposed_text"]
        assert skills_sug["status"] == "pending"
    finally:
        monkeypatch.setattr(main, "urlopen", orig_urlopen)


@pytest.mark.skipif(not is_ollama_available(), reason="Local Ollama with llama3.2 is not reachable")
def test_live_ollama_suggestions_generation():
    """Live integration test against the local running Ollama instance."""
    print("\n[LIVE TEST] Running direct call_ollama_suggestions against local Ollama...")
    t0 = time.time()
    suggestions, warnings = main.call_ollama_suggestions(SAMPLE_RESUME, SAMPLE_JD)
    duration = time.time() - t0
    print(f"[LIVE TEST] Completed in {duration:.2f}s with {len(suggestions)} suggestions.")
    print(f"[LIVE TEST] Warnings: {warnings}")
    assert isinstance(suggestions, list)
    assert duration < main.OLLAMA_TIMEOUT_SECONDS, f"Generation took {duration}s, exceeding timeout"
    for s in suggestions:
        print(f"  - [{s['section']}] {s['proposed_text'][:70]}...")
        assert s["section"] in main.SECTIONS
        assert s["proposed_text"]
        assert s["original_text"]


@pytest.mark.skipif(not is_ollama_available(), reason="Local Ollama with llama3.2 is not reachable")
def test_live_background_analysis_flow():
    """Live test of the background analysis API endpoint /api/analyze against local Ollama."""
    print("\n[LIVE TEST] Submitting real DOCX to /api/analyze with background=true...")
    docx_bytes = create_sample_docx_bytes()
    resp = client.post(
        "/api/analyze",
        files={"file": ("live_test_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"job_description": SAMPLE_JD + "\nFocus: Cloud Infrastructure and Microservices", "background": "true"},
    )
    assert resp.status_code == 200, resp.text
    job_id = resp.json()["job_id"]
    workflow_id = resp.json()["workflow_id"]
    print(f"[LIVE TEST] Job launched: job_id={job_id}, workflow_id={workflow_id}")

    try:
        # Poll /api/jobs/{job_id} until completed or timeout
        start_t = time.time()
        completed = False
        last_status = None
        last_progress = None
        while time.time() - start_t < main.OLLAMA_TIMEOUT_SECONDS + 30:
            job_res = client.get(f"/api/jobs/{job_id}")
            assert job_res.status_code == 200
            data = job_res.json()
            status = data.get("status")
            progress = data.get("progress")
            msg = data.get("message")
            if (status, progress) != (last_status, last_progress):
                print(f"  [Progress] Status: {status}, Progress: {progress}%, Msg: {msg}")
                last_status = status
                last_progress = progress
            if status == "completed":
                completed = True
                break
            elif status == "failed":
                pytest.fail(f"Background analysis failed: {data.get('error')}")
            time.sleep(0.5)

        assert completed, f"Background analysis did not complete within {main.OLLAMA_TIMEOUT_SECONDS + 30}s"
        duration = time.time() - start_t
        print(f"[LIVE TEST] Background job finished successfully in {duration:.2f}s!")

        wf_res = client.get(f"/api/workflows/{workflow_id}")
        assert wf_res.status_code == 200
        wf = wf_res.json()
        assert wf["status"] == "completed"
        assert len(wf["suggestions"]) > 0
        print(f"[LIVE TEST] Workflow contains {len(wf['suggestions'])} suggestions ready for user review:")
        for s in wf["suggestions"]:
            print(f"    * [{s['section']}] ({s['source']}): {s['proposed_text']}")
            print(f"      Length: {len(s['proposed_text'])} chars (Budget: {len(s['original_text'])})")
            if not s.get("is_skills_addition"):
                assert len(s["proposed_text"]) <= len(s["original_text"]), "Exceeded budget!"
            else:
                assert s["original_text"].rstrip(".,;") in s["proposed_text"], "Original skills were not preserved!"
    finally:
        try:
            client.delete(f"/api/workflows/{workflow_id}")
        except Exception:
            pass



if __name__ == "__main__":
    print("=" * 60)
    print("RUNNING OLLAMA FLOW AND TIMEOUT VERIFICATION")
    print("=" * 60)

    print("\n1. Testing timeout & token configuration...")
    test_ollama_timeout_configuration()
    print("   [PASS] OLLAMA_TIMEOUT_SECONDS =", main.OLLAMA_TIMEOUT_SECONDS)
    print("   [PASS] OLLAMA_MAX_RESPONSE_TOKENS =", main.OLLAMA_MAX_RESPONSE_TOKENS)

    print("\n2. Testing truncated JSON recovery...")
    test_truncated_json_recovery()
    print("   [PASS] Successfully recovered completed suggestions from truncated output")

    print("\n3. Testing background analysis timeout fallback...")
    class DummyMonkeypatch:
        def setattr(self, target, name, value):
            setattr(target, name, value)
    test_background_analysis_fallback_on_timeout(DummyMonkeypatch())
    print("   [PASS] Background job successfully fell back to manual suggestions without 504 failure")

    print("\n4. Testing targeted micro-prompt single section...")
    test_targeted_micro_prompt_single_section(DummyMonkeypatch())
    print("   [PASS] Targeted micro-prompt correctly tailored section within budget")

    print("\n5. Testing /tailor-section API endpoint...")
    test_tailor_section_endpoint(DummyMonkeypatch())
    print("   [PASS] POST /api/workflows/{id}/tailor-section updated section successfully")

    if is_ollama_available():
        print("\n6. Testing LIVE Ollama generation against localhost:11434...")
        test_live_ollama_suggestions_generation()
        print("   [PASS] Live Ollama generation succeeded!")

        print("\n7. Testing LIVE Background Analysis API flow (/api/analyze)...")
        test_live_background_analysis_flow()
        print("   [PASS] Live background analysis API flow succeeded!")
    else:
        print("\n6. [SKIP] Local Ollama is not active or llama3.2 is not found.")

    print("\n" + "=" * 60)
    print("ALL OLLAMA FLOW TESTS PASSED SUCCESSFULLY!")
    print("=" * 60)
