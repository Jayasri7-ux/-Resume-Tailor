from io import BytesIO
import json

import fitz
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient

from backend import main


JOB_DESCRIPTION = (
    "Quality associate performs quality audits, communicates with stakeholders, "
    "uses SQL and Microsoft Excel, and meets productivity benchmarks."
)
client = TestClient(main.app)


@pytest.fixture
def isolated_workspace(tmp_path, monkeypatch):
    data_root = tmp_path / "files"
    data_root.mkdir()
    test_engine = main.create_engine(
        f"sqlite:///{tmp_path / 'workflows.sqlite'}",
        connect_args={"check_same_thread": False},
    )
    main.Base.metadata.create_all(test_engine)
    monkeypatch.setattr(main, "DATA_ROOT", data_root)
    monkeypatch.setattr(main, "engine", test_engine)
    monkeypatch.setattr(main.app.state, "master_resume", None)
    monkeypatch.setattr(main, "OLLAMA_API_URL", "")
    with main.analysis_jobs_lock:
        main.analysis_jobs.clear()
    yield
    test_engine.dispose()


def upload_docx():
    doc = Document()
    doc.add_paragraph("Jordan Example", style="Title")
    doc.add_paragraph("Skills", style="Heading 1")
    skills = doc.add_paragraph("SQL, Excel, reporting and dashboard tools")
    skills.alignment = WD_ALIGN_PARAGRAPH.CENTER
    skills.runs[0].bold = True
    doc.add_paragraph("Experience", style="Heading 1")
    doc.add_paragraph("Performed quality audits and communicated with stakeholders.")
    content = BytesIO()
    doc.save(content)
    return content.getvalue()


def upload_pdf(scanned=False):
    doc = fitz.open()
    page = doc.new_page()
    if not scanned:
        page.insert_text((72, 72), "Jordan Example\nSkills\nSQL, Excel\nExperience")
        page.insert_text((72, 112), "Performed quality audits and communicated with stakeholders.")
    data = doc.tobytes()
    doc.close()
    return data


def analyze(filename, content, job_description=JOB_DESCRIPTION):
    response = client.post(
        "/api/analyze",
        files={"file": (filename, content)},
        data={"job_description": job_description, "job_title": "Quality Associate", "company_name": "Example"},
    )
    assert response.status_code == 200, response.text
    return response.json()


def resolve_all(workflow_id):
    response = client.post(
        f"/api/workflows/{workflow_id}/suggestions/bulk",
        json={"action": "approve"},
    )
    assert response.status_code == 200, response.text
    return response.json()


def remove_workflow(workflow_id):
    response = client.delete(f"/api/workflows/{workflow_id}")
    assert response.status_code == 204


def test_master_resume_is_loaded_and_used_without_upload(isolated_workspace):
    master_path = main.DATA_ROOT / "master_resume.docx"
    master_path.write_bytes(upload_docx())
    loaded = main.initialize_master_resume(master_path)
    assert loaded is not None
    assert loaded["paragraphs"]
    with main.Session(main.engine) as db:
        stored = db.get(main.WorkspaceResumeRecord, "default")
        assert stored is not None
        assert "Jordan Example" in stored.payload
    default_resume = client.get("/api/resume/default")
    assert default_resume.status_code == 200
    assert default_resume.content == master_path.read_bytes()
    health = client.get("/api/health").json()
    assert health["default_resume_available"] is True
    assert health["ollama_timeout_seconds"] == main.OLLAMA_TIMEOUT_SECONDS

    response = client.post(
        "/api/analyze",
        data={"job_description": JOB_DESCRIPTION, "job_title": "Quality Associate"},
    )
    assert response.status_code == 200, response.text
    workflow = response.json()
    assert workflow["filename"] == "master_resume.docx"
    original = client.get(f"/api/workflows/{workflow['id']}/original")
    assert original.status_code == 200
    assert "Jordan Example" in "\n".join(
        paragraph.text for paragraph in Document(BytesIO(original.content)).paragraphs
    )
    assert original.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    remove_workflow(workflow["id"])


def test_docx_approval_edit_export_and_delete(isolated_workspace):
    original_bytes = upload_docx()
    workflow = analyze("resume.docx", original_bytes)
    workflow_id = workflow["id"]
    assert workflow["demo_mode"] is True
    assert workflow["analysis"]["score"] < 100
    assert workflow["suggestions"]
    assert client.post(f"/api/workflows/{workflow_id}/finalize").status_code == 409
    assert client.get(f"/api/workflows/{workflow_id}/download/docx").status_code == 409
    skills = next(item for item in workflow["suggestions"] if item["section"] == "Skills")
    edited = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "edit", "proposed_text": "SQL, Excel, productivity benchmarks"},
    )
    assert edited.status_code == 200, edited.text
    assert next(item for item in edited.json()["suggestions"] if item["id"] == skills["id"])["status"] == "pending"

    resolve_all(workflow_id)
    approve = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "approve"},
    )
    assert approve.status_code == 200, approve.text
    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")
    assert finalized.status_code == 200, finalized.text
    result = finalized.json()
    assert result["finalized"] is True
    assert result["final_score"] > result["analysis"]["score"]

    original = client.get(result["download_urls"]["original"])
    tailored_docx = client.get(result["download_urls"]["docx"])
    tailored_pdf = client.get(result["download_urls"]["pdf"])
    assert original.status_code == tailored_docx.status_code == tailored_pdf.status_code == 200
    assert original.content == original_bytes
    source_doc = Document(BytesIO(original.content))
    tailored_doc = Document(BytesIO(tailored_docx.content))
    source_skills = next(p for p in source_doc.paragraphs if p.text.startswith("SQL, Excel,"))
    tailored_skills = next(p for p in tailored_doc.paragraphs if "productivity benchmarks" in p.text)
    assert tailored_skills.alignment == source_skills.alignment
    assert tailored_skills.runs[0].bold is True
    assert result["layout_review"]
    pdf = fitz.open(stream=tailored_pdf.content, filetype="pdf")
    assert "productivity benchmarks" in pdf[0].get_text()
    pdf.close()

    remove_workflow(workflow_id)
    assert client.get(f"/api/workflows/{workflow_id}").status_code == 404


def test_post_export_reconciles_suggestion_state_and_returns_sanitized_attachments(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    pending_export = client.post(
        f"/api/workflows/{workflow_id}/download/docx",
        json={
            "suggestions": [
                {
                    "id": suggestion["id"],
                    "status": suggestion["status"],
                    "proposed_text": suggestion["proposed_text"],
                }
                for suggestion in workflow["suggestions"]
            ]
        },
    )
    assert pending_export.status_code == 409
    skills = next(item for item in workflow["suggestions"] if item["section"] == "Skills")
    edited_text = "• SQL, Excel, productivity benchmarks"
    updated = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "approve", "proposed_text": edited_text},
    )
    assert updated.status_code == 200, updated.text

    for suggestion in updated.json()["suggestions"]:
        if suggestion["status"] == "pending":
            rejected = client.patch(
                f"/api/workflows/{workflow_id}/suggestions/{suggestion['id']}",
                json={"action": "reject"},
            )
            assert rejected.status_code == 200, rejected.text

    current = client.get(f"/api/workflows/{workflow_id}").json()
    export_state = {
        "suggestions": [
            {
                "id": suggestion["id"],
                "status": suggestion["status"],
                "proposed_text": suggestion["proposed_text"],
            }
            for suggestion in current["suggestions"]
        ]
    }
    docx_response = client.post(
        f"/api/workflows/{workflow_id}/download/docx",
        json=export_state,
    )
    assert docx_response.status_code == 200, docx_response.text
    assert "attachment" in docx_response.headers["content-disposition"]
    exported_docx = Document(BytesIO(docx_response.content))
    assert "- SQL, Excel, productivity benchmarks" in "\n".join(
        paragraph.text for paragraph in exported_docx.paragraphs
    )

    pdf_response = client.post(
        f"/api/workflows/{workflow_id}/download/pdf",
        json=export_state,
    )
    assert pdf_response.status_code == 200, pdf_response.text
    assert "attachment" in pdf_response.headers["content-disposition"]
    pdf = fitz.open(stream=pdf_response.content, filetype="pdf")
    assert "SQL, Excel, productivity benchmarks" in "\n".join(
        page.get_text() for page in pdf
    )
    pdf.close()
    remove_workflow(workflow_id)


def test_direct_download_routes_preserve_original_template_and_layout(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    blocked = client.post(
        "/api/download/pdf",
        json={"workflow_id": workflow_id, "text": workflow["final_text"]},
    )
    assert blocked.status_code == 409
    for suggestion in workflow["suggestions"]:
        response = client.patch(
            f"/api/workflows/{workflow_id}/suggestions/{suggestion['id']}",
            json={"action": "reject"},
        )
        assert response.status_code == 200, response.text

    current = client.get(f"/api/workflows/{workflow_id}").json()
    preview_text = current["final_text"]
    first_line, remaining_text = preview_text.split("\n", 1)
    submitted_text = f"**{first_line}**\n{remaining_text}"

    docx_response = client.post(
        "/api/download/docx",
        json={"workflow_id": workflow_id, "text": submitted_text},
    )
    assert docx_response.status_code == 200, docx_response.text
    assert "attachment" in docx_response.headers["content-disposition"]
    docx = Document(BytesIO(docx_response.content))
    docx_lines = [paragraph.text for paragraph in docx.paragraphs]
    assert docx_lines[0] == first_line
    assert "Skills" in docx_lines
    assert "Experience" in docx_lines
    source_docx = Document(BytesIO(upload_docx()))
    source_skills = next(paragraph for paragraph in source_docx.paragraphs if paragraph.text.startswith("SQL, Excel,"))
    exported_skills = next(paragraph for paragraph in docx.paragraphs if paragraph.text.startswith("SQL, Excel,"))
    assert exported_skills.alignment == source_skills.alignment
    assert exported_skills.runs[0].bold is True

    pdf_response = client.post(
        "/api/download/pdf",
        json={"workflow_id": workflow_id, "text": submitted_text},
    )
    assert pdf_response.status_code == 200, pdf_response.text
    assert "attachment" in pdf_response.headers["content-disposition"]
    pdf = fitz.open(stream=pdf_response.content, filetype="pdf")
    pdf_text = "\n".join(page.get_text() for page in pdf)
    assert first_line in pdf_text
    assert "Skills" in pdf_text
    assert len(pdf) == 1
    pdf.close()

    workflow_folder = main.DATA_ROOT / workflow_id
    assert not (workflow_folder / "tailored.docx").exists()
    assert not (workflow_folder / "tailored.pdf").exists()

    state_response = client.post(
        f"/api/workflows/{workflow_id}/download/docx",
        json={
            "suggestions": [
                {
                    "id": suggestion["id"],
                    "status": suggestion["status"],
                    "proposed_text": suggestion["proposed_text"],
                }
                for suggestion in current["suggestions"]
            ]
        },
    )
    assert state_response.status_code == 200, state_response.text
    assert "attachment" in state_response.headers["content-disposition"]
    state_docx = Document(BytesIO(state_response.content))
    state_skills = next(paragraph for paragraph in state_docx.paragraphs if paragraph.text.startswith("SQL, Excel,"))
    assert state_skills.alignment == source_skills.alignment
    assert state_skills.runs[0].bold is True
    state_pdf = fitz.open(workflow_folder / "tailored.pdf")
    assert len(state_pdf) == 1
    state_pdf.close()
    remove_workflow(workflow_id)


def test_docx_replacement_preserves_runs_and_embeds_styled_hyperlinks():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Reach ").bold = True
    paragraph.add_run("https://linkedin.com/in/jaya").italic = True

    main.replace_docx_paragraph(
        paragraph,
        "Reach https://linkedin.com/in/jaya",
    )

    hyperlinks = paragraph._p.xpath("./w:hyperlink")
    assert len(hyperlinks) == 1
    hyperlink = hyperlinks[0]
    hyperlink_run = hyperlink.xpath(".//w:r")[0]
    hyperlink_properties = hyperlink_run.find(main.qn("w:rPr"))
    assert hyperlink_properties.find(main.qn("w:i")) is not None
    assert hyperlink_properties.find(main.qn("w:color")).get(main.qn("w:val")) == "0563C1"
    assert hyperlink_properties.find(main.qn("w:u")).get(main.qn("w:val")) == "single"
    relationship_id = hyperlink.get(main.qn("r:id"))
    assert paragraph.part.rels[relationship_id].target_ref == "https://linkedin.com/in/jaya"
    assert paragraph._p.xpath("./w:r")[0].find(main.qn("w:rPr")).find(main.qn("w:b")) is not None

    hyperlink_properties.find(main.qn("w:color")).set(main.qn("w:val"), "C00000")
    hyperlink_properties.find(main.qn("w:u")).set(main.qn("w:val"), "double")
    main.replace_docx_paragraph(paragraph, paragraph.text)
    assert hyperlink_properties.find(main.qn("w:color")).get(main.qn("w:val")) == "C00000"
    assert hyperlink_properties.find(main.qn("w:u")).get(main.qn("w:val")) == "double"
    assert paragraph.part.rels[relationship_id].target_ref == "https://linkedin.com/in/jaya"


def test_flatten_docx_paragraphs_deduplicates_merged_table_cells():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Horizontally merged"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Vertically merged"
    table.cell(1, 0).add_paragraph("Distinct paragraph")

    paragraphs = main.flatten_docx_paragraphs(doc)
    text = [paragraph.text for paragraph in paragraphs]

    assert text.count("Horizontally merged") == 1
    assert text.count("Vertically merged") == 1
    assert text.count("Distinct paragraph") == 1


def test_pdf_fallback_compacts_empty_paragraphs_to_one_page():
    paragraphs = []
    sections = ("OBJECTIVE", "EDUCATION", "TECHNICAL SKILLS", "INTERNSHIPS", "PROJECTS", "AWARDS", "CERTIFICATIONS")
    for index in range(40):
        if index < len(sections):
            paragraphs.append(sections[index])
        paragraphs.append(
            "Used Python and analytics tools to automate reporting and review detailed data records for project work."
        )
        paragraphs.append("")

    pdf = fitz.open(stream=main.build_pdf_bytes(paragraphs), filetype="pdf")

    assert len(pdf) == 1
    assert "CERTIFICATIONS" in pdf[0].get_text()
    pdf.close()


def test_docx_export_rejects_replacements_over_the_existing_text_budget(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    skills = next(item for item in workflow["suggestions"] if item["section"] == "Skills")
    edit = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "edit", "proposed_text": "SQL, Excel, this replacement is much too long"},
    )
    assert edit.status_code == 200, edit.text
    approve = client.post(f"/api/workflows/{workflow_id}/suggestions/bulk", json={"action": "approve"})
    assert approve.status_code == 200, approve.text
    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")
    assert finalized.status_code == 422
    assert "layout budget" in finalized.json()["detail"]
    remove_workflow(workflow_id)


def test_docx_export_refuses_a_multi_page_result(isolated_workspace, monkeypatch):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    rejected = client.post(
        f"/api/workflows/{workflow_id}/suggestions/bulk",
        json={"action": "reject"},
    )
    assert rejected.status_code == 200, rejected.text
    monkeypatch.setattr(
        main,
        "inspect_pdf_layout",
        lambda _path: {"page_count": 2, "blank_pages": [], "overflow_blocks": 0, "overlapping_blocks": 0},
    )

    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")

    assert finalized.status_code == 422
    assert "one page" in finalized.json()["detail"]
    assert not (main.DATA_ROOT / workflow_id / "tailored.docx").exists()
    remove_workflow(workflow_id)


def test_pdf_analysis_export_and_scanned_pdf_error(isolated_workspace):
    workflow = analyze("resume.pdf", upload_pdf())
    workflow_id = workflow["id"]
    assert workflow["extension"] == ".pdf"
    assert workflow["analysis"]["score"] > 0
    skills = workflow["suggestions"][0]
    replacement = skills["original_text"][:-1] + (
        "!" if skills["original_text"].endswith(".") else "."
    )
    edited = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "edit", "proposed_text": replacement},
    )
    assert edited.status_code == 200, edited.text
    for other in workflow["suggestions"]:
        if other["id"] != skills["id"]:
            response = client.patch(
                f"/api/workflows/{workflow_id}/suggestions/{other['id']}",
                json={"action": "reject"},
            )
            assert response.status_code == 200, response.text
    approved = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "approve"},
    )
    assert approved.status_code == 200, approved.text
    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")
    assert finalized.status_code == 200, finalized.text
    downloaded_docx = client.get(finalized.json()["download_urls"]["docx"])
    downloaded_pdf = client.get(finalized.json()["download_urls"]["pdf"])
    assert replacement in "\n".join(
        paragraph.text for paragraph in Document(BytesIO(downloaded_docx.content)).paragraphs
    )
    pdf = fitz.open(stream=downloaded_pdf.content, filetype="pdf")
    assert replacement in "\n".join(page.get_text() for page in pdf)
    pdf.close()
    remove_workflow(workflow_id)

    scanned = client.post(
        "/api/analyze",
        files={"file": ("scan.pdf", upload_pdf(scanned=True))},
        data={"job_description": JOB_DESCRIPTION},
    )
    assert scanned.status_code == 422
    assert "scanned" in scanned.json()["detail"].lower()


def test_rejected_changes_keep_original_and_instructions_return_to_review(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    suggestion = workflow["suggestions"][0]
    response = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{suggestion['id']}",
        json={"action": "reject"},
    )
    assert response.status_code == 200, response.text
    undone = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{suggestion['id']}",
        json={"action": "undo"},
    )
    assert undone.status_code == 200, undone.text
    assert next(item for item in undone.json()["suggestions"] if item["id"] == suggestion["id"])["status"] == "pending"

    updated = client.post(
        f"/api/workflows/{workflow_id}/instructions",
        json={"instruction": "Please shorten my summary without changing any facts."},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["demo_mode"] is True
    assert updated.json()["instructions"] == "Please shorten my summary without changing any facts."
    assert all(item["status"] == "pending" for item in updated.json()["suggestions"])
    suggestion = updated.json()["suggestions"][0]
    reject = client.post(
        f"/api/workflows/{workflow_id}/suggestions/bulk",
        json={"action": "reject"},
    )
    assert reject.status_code == 200, reject.text
    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")
    assert finalized.status_code == 200, finalized.text
    assert finalized.json()["final_text"] == finalized.json()["original_text"]
    assert finalized.json()["final_score"] == finalized.json()["analysis"]["score"]
    remove_workflow(workflow_id)


def test_unavailable_ollama_returns_visible_api_error(isolated_workspace, monkeypatch):
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")

    def unavailable(*_args):
        raise main.HTTPException(status_code=502, detail="Ollama is unavailable.")

    monkeypatch.setattr(main, "call_ollama_suggestions", unavailable)
    response = client.post(
        "/api/analyze",
        files={"file": ("resume.docx", upload_docx())},
        data={"job_description": JOB_DESCRIPTION},
    )
    assert response.status_code == 502
    assert "Ollama is unavailable" in response.json()["detail"]


def test_zero_overlap_still_offers_source_text_for_manual_editing(isolated_workspace):
    suggestions = main.make_demo_suggestions(
        ["Experience", "Coordinated vendor schedules, maintained records, and resolved delivery issues."],
        "Analytical capacity, critical thinking, attention to detail, and Excel tracking.",
    )

    assert len(suggestions) == 1
    suggestion = suggestions[0]
    assert suggestion["section"] == "Experience"
    assert suggestion["original_text"]
    assert suggestion["proposed_text"] == ""
    assert suggestion["requirement"] == ""
    assert suggestion["source"] == "demo"
    assert "manual editing" in suggestion["reason"]


def test_manual_fallback_cannot_be_approved_without_a_rewrite(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx(), "analytical capacity macros mathematics")
    suggestion = workflow["suggestions"][0]
    assert suggestion["proposed_text"] == ""

    approve = client.patch(
        f"/api/workflows/{workflow['id']}/suggestions/{suggestion['id']}",
        json={"action": "approve"},
    )
    assert approve.status_code == 422
    assert "no rewritten text" in approve.json()["detail"]

    unchanged_edit = client.patch(
        f"/api/workflows/{workflow['id']}/suggestions/{suggestion['id']}",
        json={"action": "edit", "proposed_text": suggestion["original_text"]},
    )
    assert unchanged_edit.status_code == 200, unchanged_edit.text
    approve_unchanged = client.patch(
        f"/api/workflows/{workflow['id']}/suggestions/{suggestion['id']}",
        json={"action": "approve"},
    )
    assert approve_unchanged.status_code == 422
    remove_workflow(workflow["id"])


def test_keyword_extraction_keeps_short_and_compound_job_terms():
    details = main.score_resume(
        "C# development and Node.js services",
        "AI, UX, QA, C#, C++, R, Go, CI/CD, .NET, Node.js, cloud-native architecture.",
    )

    assert {"ai", "ux", "qa", "c#", "c++", "r", "go", "ci/cd", ".net", "node.js"} <= set(
        details["matched_keywords"] + details["missing_keywords"]
    )
    assert {"c#", "node.js"} <= set(details["matched_keywords"])
    assert {"ai", "ux", "qa", "c++", "r", "go", "ci/cd", ".net"} <= set(
        details["missing_keywords"]
    )


def test_score_counts_curated_related_skill_equivalents_without_forcing_floor():
    details = main.score_resume(
        "Analysis capability, problem-solving, precision",
        "Analytical capacity, critical thinking, attention to detail.",
    )

    assert details["score"] == 100
    assert details["exact_matched_count"] == 0
    assert details["semantic_matched_count"] == 6
    assert details["missing_keywords"] == []
    assert details["semantic_matches"]["critical"] == ["problem-solving"]
    assert "not machine-understood meaning" in details["method"]

    unrelated = main.score_resume(
        "Managed schedules and maintained records",
        "Analytical capacity, critical thinking, attention to detail.",
    )
    assert unrelated["score"] == 0
    assert len(unrelated["missing_keywords"]) == unrelated["requirement_count"]


def test_empty_ollama_result_returns_manual_review_during_analysis(
    isolated_workspace, monkeypatch
):
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")
    monkeypatch.setattr(main, "call_ollama_suggestions", lambda *_args: ([], []))
    response = client.post(
        "/api/analyze",
        files={"file": ("resume.docx", upload_docx())},
        data={"job_description": "Analytical capacity, critical thinking, attention to detail."},
    )
    assert response.status_code == 200, response.text
    workflow = response.json()
    assert workflow["demo_mode"] is False
    assert workflow["ai_provider"] is None
    assert len(workflow["suggestions"]) == 1
    assert workflow["suggestions"][0]["source"] == "manual"
    assert workflow["suggestions"][0]["proposed_text"] == ""
    assert workflow["suggestions"][0]["original_text"]
    assert "Ollama did not return a valid" in workflow["suggestion_warnings"][0]
    remove_workflow(workflow["id"])


def test_applying_instructions_replaces_stale_queue_items(isolated_workspace, monkeypatch):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    paragraphs = workflow["original_text"].splitlines()
    skills_index = paragraphs.index("SQL, Excel, reporting and dashboard tools")
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")

    captured_instruction = {}

    def make_instruction_suggestions(resume_text, job_description, instruction=""):
        captured_instruction["value"] = instruction
        return ([
            {
                "id": main.suggestion_id("Skills"),
                "section": "Skills",
                "paragraph_index": skills_index,
                "original_text": paragraphs[skills_index],
                "suggested_text": "SQL and Excel reporting tools",
                "proposed_text": "SQL and Excel reporting tools",
                "reason": "A concise edit for review.",
                "requirement": "SQL, Excel",
                "expected_impact": "Keeps existing tools visible.",
                "status": "pending",
                "source": "ollama",
            }
        ], [])

    monkeypatch.setattr(
        main,
        "call_ollama_suggestions",
        make_instruction_suggestions,
    )

    instruction = "Make my verified skills concise."
    updated = client.post(
        f"/api/workflows/{workflow_id}/instructions",
        json={"instruction": instruction},
    )

    assert updated.status_code == 200, updated.text
    assert updated.json()["demo_mode"] is False
    assert captured_instruction["value"] == instruction
    assert updated.json()["instructions"] == instruction
    assert len(updated.json()["suggestions"]) == 1
    assert updated.json()["suggestions"][0]["section"] == "Skills"
    assert updated.json()["suggestions"][0]["source"] == "ollama"
    remove_workflow(workflow_id)


def test_empty_ollama_result_keeps_instructions_and_offers_manual_review(isolated_workspace, monkeypatch):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")
    monkeypatch.setattr(main, "call_ollama_suggestions", lambda *_args: ([], []))

    updated = client.post(
        f"/api/workflows/{workflow_id}/instructions",
        json={"instruction": "Emphasize my verified skills."},
    )

    assert updated.status_code == 200, updated.text
    refreshed = updated.json()
    assert refreshed["instructions"] == "Emphasize my verified skills."
    assert refreshed["demo_mode"] is False
    assert refreshed["ai_provider"] is None
    assert len(refreshed["suggestions"]) == 1
    assert refreshed["suggestions"][0]["source"] == "manual"
    assert refreshed["suggestions"][0]["proposed_text"] == ""
    assert "custom instruction was saved" in refreshed["suggestions"][0]["reason"]
    current = client.get(f"/api/workflows/{workflow_id}").json()
    assert current["instructions"] == refreshed["instructions"]
    assert current["suggestions"] == refreshed["suggestions"]
    remove_workflow(workflow_id)


def test_zero_overlap_manual_edit_updates_export_and_final_score(isolated_workspace):
    job_description = "analytical capacity economics macros mathematics"
    content = upload_docx()
    document = Document(BytesIO(content))
    for paragraph in document.paragraphs:
        if paragraph.text == "SQL, Excel, reporting and dashboard tools":
            paragraph.text = (
                "Reviewed weekly operational reports, coordinated detailed vendor schedules, "
                "maintained customer records, and documented team resolutions."
            )
        elif paragraph.text == "Performed quality audits and communicated with stakeholders.":
            paragraph.text = "Managed daily records, resolved service issues, and coordinated team schedules."
    rewritten = BytesIO()
    document.save(rewritten)

    workflow = analyze("resume.docx", rewritten.getvalue(), job_description)
    workflow_id = workflow["id"]
    assert workflow["analysis"]["score"] == 0
    assert workflow["suggestions"]
    suggestion = workflow["suggestions"][0]
    assert suggestion["proposed_text"] == ""

    proposed_text = (
        "Reviewed weekly economics reports, coordinated detailed vendor schedules, "
        "maintained customer records, and documented team resolutions."
    )
    assert len(proposed_text) <= len(suggestion["original_text"])
    edited = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{suggestion['id']}",
        json={"action": "edit", "proposed_text": proposed_text},
    )
    assert edited.status_code == 200, edited.text
    approved = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{suggestion['id']}",
        json={"action": "approve"},
    )
    assert approved.status_code == 200, approved.text
    for other in approved.json()["suggestions"]:
        if other["status"] == "pending":
            response = client.patch(
                f"/api/workflows/{workflow_id}/suggestions/{other['id']}",
                json={"action": "reject"},
            )
            assert response.status_code == 200, response.text

    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")
    assert finalized.status_code == 200, finalized.text
    result = finalized.json()
    assert result["final_score"] > result["analysis"]["score"]
    assert "economics reports" in result["final_text"]
    assert "economics" in result["final_score_details"]["matched_keywords"]
    assert result["analysis"]["missing_keywords"] == [
        "analytical", "capacity", "economics", "macros", "mathematics"
    ]
    assert "economics" not in result["final_score_details"]["missing_keywords"]
    remove_workflow(workflow_id)


def test_ollama_404_identifies_the_configured_model(monkeypatch):
    def missing_model(*_args, **_kwargs):
        raise main.HTTPError("https://example.invalid", 404, "Not found", {}, None)

    monkeypatch.setattr(main, "urlopen", missing_model)
    monkeypatch.setattr(main, "OLLAMA_MODEL", "missing-model")

    with pytest.raises(main.HTTPException) as exc_info:
        main.request_ollama_json(main.Request("http://localhost:11434/api/generate"))

    assert exc_info.value.status_code == 502
    assert "missing-model" in exc_info.value.detail
    assert "ollama pull missing-model" in exc_info.value.detail


def test_upload_validation(isolated_workspace):
    unsupported = client.post(
        "/api/analyze",
        files={"file": ("resume.txt", b"text")},
        data={"job_description": JOB_DESCRIPTION},
    )
    assert unsupported.status_code == 415

    invalid_docx = client.post(
        "/api/analyze",
        files={"file": ("resume.docx", b"PK invalid")},
        data={"job_description": JOB_DESCRIPTION},
    )
    assert invalid_docx.status_code == 422


def test_approved_section_replacement_and_live_dynamic_score(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    baseline_score = workflow["analysis"]["score"]
    assert baseline_score < 100

    skills = next(item for item in workflow["suggestions"] if item["section"] == "Skills")
    # Commit approved edit with missing JD keywords ("productivity", "benchmarks").
    tailored_text_block = "SQL, Excel, productivity benchmarks."
    response = client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "approve", "proposed_text": tailored_text_block},
    )
    assert response.status_code == 200, response.text
    data = response.json()

    # 1. Check approved sections replacement in payload
    assert tailored_text_block in data["final_text"]
    assert tailored_text_block in data["tailored_text"]
    assert data["final_text"] != data["original_text"]

    # 2. Check dynamic live score calculation progressing past baseline with positive variance
    assert data["final_score"] is not None
    assert data["tailored_score"] is not None
    assert data["final_score"] > baseline_score
    point_variance = data["final_score"] - baseline_score
    assert point_variance > 0
    assert "productivity" in data["final_score_details"]["matched_keywords"]
    assert "benchmarks" in data["final_score_details"]["matched_keywords"]

    persisted = client.get(f"/api/workflows/{workflow_id}")
    assert persisted.status_code == 200, persisted.text
    assert tailored_text_block in persisted.json()["final_text"]
    assert persisted.json()["final_score"] == data["final_score"]
    remove_workflow(workflow_id)


def test_response_recalculates_live_score_from_all_approved_paragraph_edits():
    paragraphs = [
        "Jordan Example",
        "Skills",
        "Python and SQL, with documented experience maintaining accurate reports for large operational teams.",
        "Experience",
        "Coordinated communication across managers, documented process improvements, and supported multi-team project delivery.",
    ]
    job_description = (
        "Python SQL productivity benchmarks stakeholder communication "
        "project delivery analytics"
    )
    suggestions = [
        {
            "id": "skills-0",
            "section": "Skills",
            "paragraph_index": 2,
            "original_text": paragraphs[2],
            "suggested_text": "Python, SQL, productivity benchmarks, analytics, and reporting for large operational teams.",
            "proposed_text": "Python, SQL, productivity benchmarks, analytics, and reporting for large operational teams.",
            "status": "approved",
            "source": "ollama",
        },
        {
            "id": "experience-0",
            "section": "Experience",
            "paragraph_index": 4,
            "original_text": paragraphs[4],
            "suggested_text": "Delivered stakeholder communication, project delivery, and benchmarks through verified cross-team coordination.",
            "proposed_text": "Delivered stakeholder communication, project delivery, and benchmarks through verified cross-team coordination.",
            "status": "approved",
            "source": "ollama",
        },
    ]
    payload = {
        "paragraphs": paragraphs,
        "job_description": job_description,
        "suggestions": suggestions,
        "job_title": "",
        "company_name": "",
        "suggestion_warnings": [],
        "limitations": [],
        "instructions": "",
        "finalized": False,
        "layout_review": [],
        "analysis": main.score_resume("\n".join(paragraphs), job_description),
        "final_score": 0,
        "final_score_details": None,
        "tailored_score": 0,
        "tailored_score_details": None,
        "final_text": "\n".join(paragraphs),
        "tailored_text": "\n".join(paragraphs),
    }
    record = main.WorkflowRecord(
        id="test-workflow",
        filename="resume.docx",
        extension=".docx",
        created_at="2026-01-01T00:00:00+00:00",
    )

    result = main.response_payload(record, payload)

    assert result["tailored_text"] != result["original_text"]
    assert "productivity benchmarks" in result["tailored_text"]
    assert "stakeholder communication" in result["tailored_text"]
    assert result["tailored_score"] > result["analysis"]["score"]
    assert result["tailored_score"] == result["final_score"]
    assert result["tailored_score_details"]["score"] == result["tailored_score"]
    assert result["tailored_score_details"]["missing_keywords"] == []


def test_generative_rewrite_inverse_reconciliation_zero_overlap(monkeypatch):
    resume_text = (
        "Taylor Example\n"
        "Summary\n"
        "Customer service representative supporting inbound caller inquiries.\n"
        "Experience\n"
        "Automated support workflows with Python and Kubernetes, documenting deployments, ticket resolution, "
        "and operational reporting for distributed service teams."
    )
    job_description = (
        "Senior Cloud Engineer who builds Python microservices, manages Kubernetes clusters, "
        "and optimizes PostgreSQL database queries for high availability."
    )
    fake_ollama_payload = {
        "response": json.dumps({
            "suggestions": [{
                "section": "Experience",
                "original_content": "Automated support workflows with Python and Kubernetes, documenting deployments, ticket resolution, and operational reporting for distributed service teams.",
                "proposed_content": "Used Python and Kubernetes to automate cloud support workflows and document deployment issues.",
                "why_this_may_help": "Reframes verified Python and Kubernetes workflow experience in concise, role-relevant language.",
                "job_requirement": "Python, Kubernetes",
                "expected_impact": "Makes supported Python and Kubernetes experience more explicit.",
            }]
        })
    }

    class MockResponse:
        def read(self):
            return json.dumps(fake_ollama_payload).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(main, "urlopen", lambda *args, **kwargs: MockResponse())

    validated, warnings = main.call_ollama_suggestions(resume_text, job_description)
    assert len(validated) == 1
    suggestion = validated[0]
    assert suggestion["source"] == "ollama"
    assert suggestion["section"] == "Experience"
    assert suggestion["reason"].startswith("Reframes verified Python")
    assert suggestion["requirement"] == "Python, Kubernetes"
    assert "Python" in suggestion["proposed_text"]
    assert "Kubernetes" in suggestion["proposed_text"]
    assert suggestion["proposed_text"] != suggestion["original_text"]

    fake_ollama_payload["response"] = json.dumps({
        "suggestions": [{
            "section": "Experience",
            "original_content": suggestion["original_text"],
            "proposed_content": "x" * (len(suggestion["original_text"]) + 1),
            "why_this_may_help": "Long test value.",
            "job_requirement": "Python",
            "expected_impact": "None",
        }]
    })
    over_budget, warnings = main.call_ollama_suggestions(resume_text, job_description)
    assert not over_budget
    assert any("character budget" in warning for warning in warnings)

    fake_ollama_payload["response"] = json.dumps({
        "suggestions": [{
            "section": "Experience",
            "original_content": suggestion["original_text"],
            "proposed_content": suggestion["original_text"],
            "why_this_may_help": "Unchanged test value.",
            "job_requirement": "Python",
            "expected_impact": "None",
        }]
    })
    unchanged, warnings = main.call_ollama_suggestions(resume_text, job_description)
    assert not unchanged
    assert any("unchanged" in warning for warning in warnings)


def test_ollama_prompt_receives_every_missing_keyword_in_one_local_call(monkeypatch):
    keywords = [f"specialty{i}" for i in range(31)]
    captured_prompt = {}
    response_payload = {"response": json.dumps({"suggestions": []})}

    class MockResponse:
        def read(self):
            return json.dumps(response_payload).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *_args):
            pass

    def capture_request(request, **_kwargs):
        body = json.loads(request.data)
        captured_prompt["request_count"] = captured_prompt.get("request_count", 0) + 1
        captured_prompt["url"] = request.full_url
        captured_prompt["model"] = body["model"]
        captured_prompt["text"] = body["prompt"]
        captured_prompt["temperature"] = body["options"]["temperature"]
        captured_prompt["num_predict"] = body["options"]["num_predict"]
        captured_prompt["format"] = body["format"]
        return MockResponse()

    monkeypatch.setattr(main, "urlopen", capture_request)
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")
    monkeypatch.setattr(main, "OLLAMA_MODEL", "llama3.2")
    main.call_ollama_suggestions(
        "entry level resume",
        " ".join(keywords)
        + " analytical capacity critical thinking attention to detail macros mathematics",
        "Emphasize verified SQL work and keep the summary concise.",
    )

    prompt_text = captured_prompt["text"]
    context_text = prompt_text.split("COMPLETE_RESUME_AND_JOB_CONTEXT_JSON:\n", 1)[1]
    context = json.loads(context_text)
    assert all(keyword in context["missing_job_description_keywords"] for keyword in keywords)
    assert all(section in context["resume_sections"] for section in main.SECTIONS)
    assert [
        paragraph["text"]
        for section in context["resume_sections"].values()
        for paragraph in section
    ] == ["entry level resume"]
    assert context["user_instructions"] == (
        "Emphasize verified SQL work and keep the summary concise."
    )
    assert captured_prompt["request_count"] == 1
    assert captured_prompt["url"] == "http://localhost:11434/api/generate"
    assert captured_prompt["model"] == "llama3.2"
    assert captured_prompt["format"] == "json"
    assert captured_prompt["num_predict"] == main.OLLAMA_MAX_RESPONSE_TOKENS
    assert f"only the {main.OLLAMA_MAX_SUGGESTIONS} most relevant suggestions" in prompt_text
    assert "every Experience bullet, Projects entry, and Education entry" in prompt_text
    assert "apply it to every relevant suggestion" in prompt_text
    assert "meaningfully different" in prompt_text
    assert "CRITICAL LAYOUT CONSTRAINT: The final output must fit on exactly ONE page." in prompt_text
    assert "aim for no more than 80% of max_characters" in prompt_text
    assert "omit that section instead of returning an over-budget rewrite" in prompt_text
    assert "never return the original unchanged" in prompt_text
    assert all(
        phrase in context["job_description"]
        for phrase in (
            "analytical capacity",
            "critical thinking",
            "attention to detail",
            "macros",
            "mathematics",
        )
    )
    assert captured_prompt["temperature"] == 0.35
    assert "original_content" in prompt_text
    assert "proposed_content" in prompt_text
    assert "why_this_may_help" in prompt_text


def test_ollama_suggestion_count_is_bounded(monkeypatch):
    originals = {
        section: (
            f"Documented and coordinated verified {section.casefold()} tasks, "
            "reviewed detailed project information, and prepared concise reports."
        )
        for section in main.SECTIONS
    }
    replacements = {
        section: f"Verified {section.casefold()} work and prepared concise reports."
        for section in main.SECTIONS
    }
    paragraphs = []
    for section in main.SECTIONS:
        paragraphs.extend((section, originals[section]))

    class MockResponse:
        def read(self):
            return json.dumps({
                "response": json.dumps({
                    "suggestions": [
                        {
                            "section": section,
                            "original_content": original,
                            "proposed_content": proposed,
                        }
                        for section, original, proposed in (
                            (section, originals[section], replacements[section])
                            for section in main.SECTIONS
                        )
                    ]
                })
            }).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *_args):
            pass

    monkeypatch.setattr(main, "urlopen", lambda *_args, **_kwargs: MockResponse())
    suggestions, _warnings = main.call_ollama_suggestions(
        "\n".join(paragraphs),
        "Python SQL reporting automation and data analysis.",
    )

    assert len(suggestions) == main.OLLAMA_MAX_SUGGESTIONS == 4
    assert [item["section"] for item in suggestions] == main.SECTIONS[:4]


def test_ollama_generation_uses_one_local_request_without_rate_limit_retries(monkeypatch):
    attempts = []
    response_payload = {"response": json.dumps({"suggestions": []})}

    class MockResponse:
        def read(self):
            return json.dumps(response_payload).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *_args):
            pass

    def local_response(request, timeout):
        attempts.append(request)
        assert timeout == main.OLLAMA_TIMEOUT_SECONDS
        return MockResponse()

    monkeypatch.setattr(main, "urlopen", local_response)
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")

    suggestions, warnings = main.call_ollama_suggestions(
        "A candidate resume with verified SQL reporting experience.",
        "SQL reporting analyst role.",
    )

    assert suggestions == []
    assert warnings == []
    assert len(attempts) == 1


def test_ollama_timeout_is_reported_as_slow_generation(monkeypatch):
    def timeout(*_args, **_kwargs):
        raise TimeoutError("timed out")

    monkeypatch.setattr(main, "urlopen", timeout)
    with pytest.raises(main.HTTPException) as exc_info:
        main.request_ollama_json(main.Request("http://localhost:11434/api/generate"))

    assert exc_info.value.status_code == 504
    assert "did not finish generating" in exc_info.value.detail
    assert str(main.OLLAMA_TIMEOUT_SECONDS) in exc_info.value.detail


def test_ollama_http_error_is_reported_without_retry(monkeypatch):
    attempts = []

    def always_rate_limited(*_args, **_kwargs):
        attempts.append(None)
        raise main.HTTPError(
            "https://example.invalid",
            429,
            "Too many requests",
            {"Retry-After": "0"},
            None,
        )

    monkeypatch.setattr(main, "urlopen", always_rate_limited)
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")

    with pytest.raises(main.HTTPException) as exc_info:
        main.call_ollama_suggestions(
            "A candidate resume with verified SQL reporting experience.",
            "SQL reporting analyst role.",
        )

    assert len(attempts) == 1
    assert exc_info.value.status_code == 502
    assert "Ollama request failed (HTTP 429)" in exc_info.value.detail


def test_end_to_end_data_stream_convergence_docx_and_pdf(isolated_workspace):
    workflow = analyze("resume.docx", upload_docx())
    workflow_id = workflow["id"]
    assert workflow["download_urls"] is None
    skills = next(item for item in workflow["suggestions"] if item["section"] == "Skills")
    tailored_text_block = "SQL, Excel, productivity benchmarks"

    # Commit approved replacement
    client.patch(
        f"/api/workflows/{workflow_id}/suggestions/{skills['id']}",
        json={"action": "approve", "proposed_text": tailored_text_block},
    )
    # Reject other pending to satisfy review requirement
    for other in workflow["suggestions"]:
        if other["id"] != skills["id"]:
            client.patch(
                f"/api/workflows/{workflow_id}/suggestions/{other['id']}",
                json={"action": "reject"},
            )

    finalized = client.post(f"/api/workflows/{workflow_id}/finalize")
    assert finalized.status_code == 200, finalized.text
    result = finalized.json()
    assert result["finalized"] is True
    assert result["download_urls"]["docx"].endswith("/download/docx")
    assert result["download_urls"]["pdf"].endswith("/download/pdf")

    # Check python-docx writer output
    tailored_docx_resp = client.get(result["download_urls"]["docx"])
    assert tailored_docx_resp.status_code == 200
    doc = Document(BytesIO(tailored_docx_resp.content))
    docx_text = "\n".join(p.text for p in doc.paragraphs)
    assert tailored_text_block in docx_text

    # Check ReportLab PDF builder output
    tailored_pdf_resp = client.get(result["download_urls"]["pdf"])
    assert tailored_pdf_resp.status_code == 200
    pdf = fitz.open(stream=tailored_pdf_resp.content, filetype="pdf")
    pdf_text = "\n".join(page.get_text() for page in pdf)
    pdf.close()
    assert "productivity benchmarks" in pdf_text
    remove_workflow(workflow_id)


def test_deterministic_weighted_scoring_categories_and_evidence_classification():
    resume = (
        "Jordan Example\n"
        "Skills\n"
        "Python, SQL, React, TypeScript, Docker\n"
        "Experience\n"
        "Performed quality audits and communicated with stakeholders to maintain benchmarks.\n"
        "Projects\n"
        "Built automated reporting pipeline prototype using Python and SQL.\n"
        "Education\n"
        "Bachelor of Science in Computer Science, Certified AWS Cloud Practitioner\n"
    )
    job_description = (
        "We are seeking a Quality Associate who uses Python, SQL, and Excel. "
        "Responsible for performing audits and communicating with stakeholders. "
        "Projects include building automated pipelines. "
        "Bachelor degree in Computer Science or related field required. "
        "Preferred qualifications: Master degree and Docker experience is a plus."
    )

    details1 = main.score_resume(resume, job_description)
    details2 = main.score_resume(resume, job_description)

    # Determinism
    assert details1["score"] == details2["score"]
    assert details1["categories"] == details2["categories"]
    assert details1["evidence_items"] == details2["evidence_items"]

    # Category weights
    categories = details1["categories"]
    assert categories["skills"]["weight"] == 0.30
    assert categories["skills"]["weight_percent"] == 30
    assert categories["responsibilities_and_experience"]["weight"] == 0.25
    assert categories["responsibilities_and_experience"]["weight_percent"] == 25
    assert categories["keyword_alignment"]["weight"] == 0.15
    assert categories["keyword_alignment"]["weight_percent"] == 15
    assert categories["projects"]["weight"] == 0.10
    assert categories["projects"]["weight_percent"] == 10
    assert categories["education_and_certifications"]["weight"] == 0.10
    assert categories["education_and_certifications"]["weight_percent"] == 10
    assert categories["preferred_qualifications"]["weight"] == 0.10
    assert categories["preferred_qualifications"]["weight_percent"] == 10

    # Evidence items classification
    classifications = {item["classification"] for item in details1["evidence_items"]}
    assert "fully evidenced" in classifications
    assert "not evidenced" in classifications

    # Check evidence structure
    for item in details1["evidence_items"]:
        assert "requirement" in item
        assert "category" in item
        assert item["classification"] in ("fully evidenced", "partially evidenced", "not evidenced", "uncertain")
        assert "evidence" in item
        assert "notes" in item

    # Check specific keyword evidence
    python_item = next(item for item in details1["evidence_items"] if item["requirement"] == "python")
    assert python_item["classification"] == "fully evidenced"
    assert "Python" in python_item["evidence"]


def test_candidate_profile_persistence_editing_and_extraction(isolated_workspace):
    # 1. Initial profile
    init_res = client.get("/api/profile")
    assert init_res.status_code == 200
    profile = init_res.json()
    assert "full_name" in profile
    assert "skills" in profile
    assert "experience" in profile

    # 2. Save edited profile
    updated_profile = {
        "full_name": "Alex Candidate",
        "headline": "Senior Data Analyst",
        "email": "alex@example.com",
        "phone": "(555) 123-4567",
        "location": "New York, NY",
        "summary": "Experienced data professional specializing in verified analytics and stakeholder reporting.",
        "skills": ["SQL", "Python", "Tableau", "Statistical Modeling"],
        "experience": [
            {
                "title": "Data Analyst",
                "company": "Acme Analytics",
                "location": "New York, NY",
                "dates": "2021 - Present",
                "bullets": ["Executed automated SQL reporting for operational stakeholders."],
            }
        ],
        "education": [
            {
                "institution": "City University",
                "degree": "B.S. in Mathematics",
                "field_of_study": "Mathematics",
                "graduation_year": "2021",
            }
        ],
        "projects": [
            {
                "name": "Metrics Engine",
                "description": "Designed real-time metrics tracking pipeline in Python.",
                "technologies": ["Python", "SQL"],
            }
        ],
        "certifications": ["Tableau Desktop Specialist"],
    }
    save_res = client.put("/api/profile", json=updated_profile)
    assert save_res.status_code == 200
    saved = save_res.json()
    assert saved["full_name"] == "Alex Candidate"
    assert "SQL" in saved["skills"]
    assert len(saved["experience"]) == 1

    # 3. Retrieve persistent profile
    persisted = client.get("/api/profile").json()
    assert persisted["full_name"] == "Alex Candidate"
    assert persisted["email"] == "alex@example.com"
    assert persisted["skills"] == updated_profile["skills"]

    # 4. Extract profile from resume DOCX without inventing facts
    docx_bytes = upload_docx()
    extract_res = client.post(
        "/api/profile/extract",
        files={"file": ("jordan.docx", docx_bytes)},
    )
    assert extract_res.status_code == 200
    extracted = extract_res.json()
    assert "Jordan" in extracted["full_name"]
    assert any("SQL" in s for s in extracted["skills"])
    assert any("quality audits" in b.lower() for exp in extracted["experience"] for b in exp["bullets"])

    # 5. Analyze using candidate profile directly (without file upload)
    analyze_res = client.post(
        "/api/analyze",
        data={
            "job_description": JOB_DESCRIPTION,
            "job_title": "Quality Analyst",
            "use_profile": "true",
        },
    )
    assert analyze_res.status_code == 200, analyze_res.text
    wf = analyze_res.json()
    assert "profile" in wf["filename"].lower() or "jordan" in wf["filename"].lower()
    assert wf["analysis"]["score"] > 0
    remove_workflow(wf["id"])


def test_background_analysis_job_polling_and_duplicate_protection(isolated_workspace, monkeypatch):
    monkeypatch.setattr(main, "OLLAMA_API_URL", "http://localhost:11434/api/generate")

    fake_payload = {
        "response": json.dumps({
            "suggestions": [{
                "section": "Skills",
                "original_content": "SQL, Excel, reporting and dashboard tools",
                "proposed_content": "SQL, Excel, reporting tools, and productivity benchmarks.",
                "why_this_may_help": "Emphasizes benchmarks.",
                "job_requirement": "productivity benchmarks",
                "expected_impact": "Direct alignment",
            }]
        })
    }

    class MockResp:
        def read(self):
            return json.dumps(fake_payload).encode("utf-8")
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(main, "urlopen", lambda *args, **kwargs: MockResp())

    # 1. Start background analysis
    response = client.post(
        "/api/analyze",
        files={"file": ("resume.docx", upload_docx())},
        data={"job_description": JOB_DESCRIPTION, "background": "true"},
    )
    assert response.status_code == 200, response.text
    job_info = response.json()
    job_id = job_info["job_id"]
    workflow_id = job_info["workflow_id"]
    assert job_id == workflow_id
    assert job_info["status"] in ("processing", "completed")

    # 2. Check duplicate job protection while active or newly registered
    dup_res = client.post(
        "/api/analyze",
        files={"file": ("resume.docx", upload_docx())},
        data={"job_description": JOB_DESCRIPTION, "background": "true"},
    )
    assert dup_res.status_code == 200
    dup_data = dup_res.json()
    assert dup_data["job_id"] == job_id or dup_data["workflow_id"] == workflow_id

    # 3. Poll job status endpoint
    import time
    for _ in range(10):
        status_res = client.get(f"/api/jobs/{job_id}")
        assert status_res.status_code == 200
        status_data = status_res.json()
        if status_data["status"] == "completed":
            break
        time.sleep(0.1)

    # 4. Check workflow status endpoint
    wf_status_res = client.get(f"/api/workflows/{workflow_id}/status")
    assert wf_status_res.status_code == 200
    assert wf_status_res.json()["status"] in ("processing", "completed")

    remove_workflow(workflow_id)


def test_end_to_end_tailoring_pipeline(isolated_workspace):
    # 1. Candidate Profile Setup & Save
    profile_data = {
        "full_name": "Alex Mercer",
        "headline": "Lead Distributed Systems Architect",
        "contact": "alex.mercer@example.com | 555-019-8822 | Austin, TX",
        "summary": "10+ years engineering resilient microservices, distributed caches, and cloud infrastructure.",
        "skills": ["Python", "FastAPI", "Go", "Docker", "Kubernetes", "PostgreSQL", "Kafka", "AWS"],
        "experience": [
            "Lead Systems Architect at Nexus Corp (2020-Present)\n- Scaled distributed event mesh handling 100k events/sec using Kafka and Go.\n- Reduced infrastructure latency by 35% through custom Redis caching layers."
        ],
        "education": ["B.S. in Computer Science, University of Texas at Austin (2015)"],
        "projects": ["Distributed Vector Index: open-source real-time search engine in Go."],
        "certifications": ["AWS Certified Solutions Architect Professional", "CKA"],
    }
    save_resp = client.put("/api/profile", json=profile_data)
    assert save_resp.status_code == 200, save_resp.text
    assert save_resp.json()["full_name"] == "Alex Mercer"

    # Verify retrieval
    get_profile = client.get("/api/profile")
    assert get_profile.status_code == 200
    assert "FastAPI" in get_profile.json()["skills"]

    # 2. Paste Job Description & Analyze with use_profile=True
    jd = (
        "Seeking Senior Distributed Systems Architect.\n"
        "Required Skills: Python, FastAPI, Docker, Kubernetes, PostgreSQL, Kafka, microservices.\n"
        "Responsibilities: Design resilient distributed systems, scale data pipelines, and mentor teams.\n"
        "Preferred Qualifications: Go, AWS Certified Solutions Architect, high-throughput caching."
    )
    analyze_resp = client.post(
        "/api/analyze",
        data={
            "job_description": jd,
            "job_title": "Senior Distributed Systems Architect",
            "company_name": "HyperScale Tech",
            "use_profile": "true",
        },
    )
    assert analyze_resp.status_code == 200, analyze_resp.text
    wf = analyze_resp.json()
    wf_id = wf["id"]

    # 3. Check Original Score & Deterministic Category Breakdown
    analysis = wf["analysis"]
    assert 0 <= analysis["score"] <= 100
    assert "category_scores" in analysis
    cat_names = [c["category"] for c in analysis["category_scores"]]
    expected_categories = [
        "Skills",
        "Responsibilities & Experience",
        "Keyword Alignment",
        "Projects",
        "Education & Certifications",
        "Preferred Qualifications",
    ]
    for exp in expected_categories:
        assert exp in cat_names, f"Missing category: {exp}"

    # Verify deterministic weights sum to 100%
    total_weights = sum(c["weight"] for c in analysis["category_scores"])
    assert round(total_weights, 1) == 100.0

    # 4. Check Requirement Evidence Classification
    assert "evidence_breakdown" in analysis
    valid_classifications = {"fully evidenced", "partially evidenced", "not evidenced", "uncertain"}
    for ev in analysis["evidence_breakdown"]:
        assert ev["classification"] in valid_classifications

    # 5. Review & Suggestion Edits
    assert len(wf["suggestions"]) > 0
    sugg = wf["suggestions"][0]
    sugg_id = sugg["id"]
    orig_text = sugg["original_text"]

    # Test longer replacement text exceeding character budget -> rejected with 422
    over_budget_text = orig_text + " Extra text making it much longer than the source paragraph."
    edit_long = client.patch(
        f"/api/workflows/{wf_id}/suggestions/{sugg_id}",
        json={"action": "edit", "proposed_text": over_budget_text},
    )
    assert edit_long.status_code == 200
    approve_long = client.patch(
        f"/api/workflows/{wf_id}/suggestions/{sugg_id}",
        json={"action": "approve"},
    )
    assert approve_long.status_code == 200
    export_attempt = client.post(
        f"/api/workflows/{wf_id}/download/docx",
        json={"suggestions": [{"id": sugg_id, "status": "approved", "proposed_text": over_budget_text}]},
    )
    assert export_attempt.status_code == 422
    assert "exceeds" in export_attempt.json()["detail"].lower()

    # Safe edit within character budget
    safe_text = orig_text[:len(orig_text) - 5]
    if not safe_text:
        safe_text = orig_text
    safe_edit = client.patch(
        f"/api/workflows/{wf_id}/suggestions/{sugg_id}",
        json={"action": "edit", "proposed_text": safe_text},
    )
    assert safe_edit.status_code == 200
    approve_safe = client.patch(
        f"/api/workflows/{wf_id}/suggestions/{sugg_id}",
        json={"action": "approve"},
    )
    assert approve_safe.status_code == 200

    # Reject other suggestions if present
    for other in wf["suggestions"]:
        if other["id"] != sugg_id:
            reject_res = client.patch(
                f"/api/workflows/{wf_id}/suggestions/{other['id']}",
                json={"action": "reject"},
            )
            assert reject_res.status_code == 200

    # 6. Finalize Resume & Verify Tailored Score Determinism
    finalize_res = client.post(f"/api/workflows/{wf_id}/finalize")
    assert finalize_res.status_code == 200, finalize_res.text
    final_wf = finalize_res.json()
    assert final_wf["finalized"] is True
    assert final_wf["tailored_score"] is not None
    assert "tailored_score_details" in final_wf
    tailored_details = final_wf["tailored_score_details"]
    assert "category_scores" in tailored_details
    assert "evidence_breakdown" in tailored_details

    # 7. DOCX Download Verification
    docx_res = client.post(
        f"/api/workflows/{wf_id}/download/docx",
        json={"suggestions": final_wf["suggestions"]},
    )
    assert docx_res.status_code == 200
    assert docx_res.content[:2] == b"PK"  # Valid zip/DOCX signature
    assert len(docx_res.content) > 1000

    # 8. PDF Download Verification
    pdf_res = client.post(
        f"/api/workflows/{wf_id}/download/pdf",
        json={"suggestions": final_wf["suggestions"]},
    )
    assert pdf_res.status_code == 200
    assert pdf_res.content[:5] == b"%PDF-"  # Valid PDF header
    assert len(pdf_res.content) > 500

    remove_workflow(wf_id)


def test_boilerplate_noise_filtering_eliminates_false_negatives():
    jd = (
        "Seeking a Full Stack Engineer. We offer competitive compensation, 401k matching, "
        "health, dental, and vision insurance with flexible PTO in a fast-paced and inclusive company culture. "
        "Requirements: Python, React, SQL, and data structures. 5+ years experience preferred."
    )
    keywords = main.extract_keywords(jd)
    # Perks and boilerplate noise MUST NOT be extracted as requirements
    assert "compensation" not in keywords
    assert "401k" not in keywords
    assert "health" not in keywords
    assert "dental" not in keywords
    assert "vision" not in keywords
    assert "fast-paced" not in keywords
    assert "culture" not in keywords
    assert "5+" not in keywords
    # Real technical requirements and multi-word entities MUST be preserved
    assert "python" in keywords
    assert "react" in keywords
    assert "sql" in keywords
    assert "data structures" in keywords


def test_tier1_role_eligibility_and_career_stage_detection():
    fresher_resume = (
        "Jaya Sri Palli\n"
        "B.Tech in Information Technology (2022 - 2026), CGPA: 8.77\n"
        "Skills: Python, React, Django, Flask, MongoDB, SQL, Data Structures\n"
        "Internships: Larsen & Toubro Full Stack Developer Intern"
    )
    jd = (
        "Full Stack Developer position building web applications using Python, React, and REST APIs."
    )
    analysis = main.score_resume(fresher_resume, jd, job_title="Full Stack Developer")
    assert analysis["career_stage"] == "Fresher / Entry-Level"
    assert analysis["tier1_role_match"]["eligibility_status"] == "High Eligibility"
    assert "Full Stack" in analysis["tier1_role_match"]["domain"]


def test_tool_confirmation_endpoint_live_rescore(isolated_workspace):
    # Upload resume
    res = analyze("Jordan.docx", upload_docx(), job_description="Developer uses SQL, Python, and Docker.")
    wf_id = res["id"]

    # Initial analysis
    initial_score = res["analysis"]["score"]
    discovered = res["analysis"].get("discovered_tools", [])
    assert any(t["tool"] == "docker" for t in discovered)

    # Candidate confirms Docker via interactive action
    confirm_res = client.post(
        f"/api/workflows/{wf_id}/confirm-tool",
        json={"tool": "docker", "status": "approved"},
    )
    assert confirm_res.status_code == 200
    updated_wf = confirm_res.json()
    assert "docker" in updated_wf["analysis"]["confirmed_tools"]
    assert updated_wf["analysis"]["score"] >= initial_score

    # Check discovered tools endpoint
    tools_res = client.get(f"/api/workflows/{wf_id}/discovered-tools")
    assert tools_res.status_code == 200
    tools_data = tools_res.json()
    assert "docker" in tools_data["confirmed_tools"]

    remove_workflow(wf_id)



