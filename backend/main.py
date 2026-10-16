from __future__ import annotations

import hashlib
import html
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import unicodedata
import uuid
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Literal
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import pymupdf as fitz
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph as PdfParagraph, SimpleDocTemplate, Spacer
from reportlab.platypus.doctemplate import LayoutError
from sqlalchemy import String, Text, create_engine, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column
from starlette.concurrency import run_in_threadpool

ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / "backend" / ".env")
DATA_ROOT = ROOT / ".resume_fit_data"
MASTER_RESUME_PATH = DATA_ROOT / "master_resume.docx"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
OLLAMA_API_URL = os.environ.get(
    "OLLAMA_API_URL", "http://localhost:11434/api/generate"
).strip()
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.2").strip()
OLLAMA_TIMEOUT_SECONDS = int(os.environ.get("OLLAMA_TIMEOUT_SECONDS", "300"))
OLLAMA_MAX_RESPONSE_TOKENS = int(os.environ.get("OLLAMA_MAX_RESPONSE_TOKENS", "512"))
# llama3.2 reports a 2,048-token runtime context on this installation. Keep the
# default at that confirmed limit; a user can raise it only after reconfiguring
# their Ollama runner to support the larger context.
OLLAMA_NUM_CTX = int(os.environ.get("OLLAMA_NUM_CTX", "2048"))
STOP_WORDS = set(
    "a about above across after again against all also an and any are as at be because been "
    "before being below between both but by can could did do does doing down during each "
    "few for from further had has have having he her here hers herself him himself his how "
    "i if in into is it its itself just me more most my myself no nor not of off on once only "
    "or other our ours ourselves out over own same she should so some such than that the their "
    "theirs them themselves then there these they this those through to too under until up "
    "very was we were what when where which while who whom why will with you your yours "
    "yourself yourselves role roles experience including related required preferred work "
    "working opportunity requirements responsibilities qualifications must shall would might".split()
)
SECTIONS = [
    "Professional Summary",
    "Skills",
    "Experience",
    "Internships",
    "Projects",
    "Education",
    "Certifications",
    "Achievements",
]
OLLAMA_MAX_SUGGESTIONS = 4
SECTION_ALIASES = {
    "Professional Summary": ("summary", "objective", "profile", "professional summary"),
    "Skills": (
        "skills", "technical skills", "core competencies", "competencies",
        "languages", "programming languages", "frontend", "backend",
        "database", "databases", "tools", "developer tools",
    ),
    "Experience": ("experience", "work experience", "professional experience", "employment"),
    "Internships": ("internships", "internship"),
    "Projects": ("projects", "project experience", "academic projects"),
    "Education": ("education", "academic background", "btech", "b.tech", "intermediate", "ssc"),
    "Certifications": ("certifications", "certificates", "licenses"),
    "Achievements": ("achievements", "awards", "honors", "accomplishments", "hackathon", "hackathons"),
}
KEYWORD_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])(?:\.[A-Za-z0-9]+|[A-Za-z0-9]+(?:[./&-][A-Za-z0-9]+)*)(?:[+#]+)?(?![A-Za-z0-9])"
)
SINGLE_LETTER_KEYWORDS = {"c", "r"}
RELATED_KEYWORD_ALIASES = {
    "analytical": {"analysis", "analytics", "analyst"},
    "capacity": {"ability", "capability", "capabilities"},
    "critical": {"reasoning", "problem-solving"},
    "thinking": {"reasoning", "problem-solving"},
    "attention": {"accuracy", "attentive", "careful", "focus", "precision", "thoroughness"},
    "detail": {"accuracy", "detail-oriented", "precision", "thoroughness"},
    # Database and Data equivalents
    "nosql": {"mongodb", "cassandra", "dynamodb", "redis"},
    "postgresql": {"sql", "postgres", "mysql", "sqlite", "rdbms", "database"},
    "postgres": {"sql", "postgresql", "mysql", "sqlite", "rdbms", "database"},
    "mysql": {"sql", "postgresql", "postgres", "sqlite", "rdbms", "database"},
    "sqlite": {"sql", "mysql", "postgresql", "rdbms", "database"},
    "database": {"sql", "mongodb", "sqlite", "mysql", "dbms", "postgresql", "nosql", "postgres"},
    "databases": {"sql", "mongodb", "sqlite", "mysql", "dbms", "postgresql", "nosql", "postgres"},
    # Backend framework equivalents
    "backend": {"django", "flask", "node.js", "node", "express", "fastapi", "spring", "python"},
    "back-end": {"backend", "django", "flask", "node.js", "python"},
    "frontend": {"react", "react.js", "javascript", "typescript", "html", "css", "vue", "angular"},
    "front-end": {"frontend", "react", "html", "css", "javascript"},
    "html5": {"html", "html5"},
    # ORM and Database equivalents
    "orm": {"django", "sql", "mongodb", "sqlite", "orm"},
    "object-relational mapping": {"django", "sql", "mongodb", "sqlite", "orm"},
    # Security equivalents
    "security": {"jwt", "authentication", "security"},
    "data protection": {"jwt", "authentication", "security"},
    # API & Architecture
    "apis": {"rest api", "rest apis", "restful api", "restful apis", "api", "graphql"},
    "api": {"rest api", "rest apis", "restful api", "restful apis", "apis", "graphql"},
    # Problem solving & Testing
    "problem-solving": {"algorithms", "data structures", "dsa", "analytical", "critical"},
    "test": {"testing", "unit testing", "pytest"},
    "debug": {"testing", "debugging", "test"},
    "testing": {"unit testing", "pytest", "jest", "test", "tests"},
    # Python frameworks
    "pyramid": {"django", "flask"},
    # Cloud & Deployment
    "cloud": {"aws", "azure", "gcp", "render"},
    "deployment": {"render", "docker", "kubernetes", "aws", "ci/cd"},
}

BOILERPLATE_NOISE_WORDS = {
    # Compensation, perks, and HR benefits
    "compensation", "salary", "benefits", "health", "dental", "vision", "401k", "insurance",
    "vacation", "pto", "leave", "perks", "flexible", "hybrid", "remote", "stipend", "bonus",
    "relocation", "wellness", "allowance", "commute", "retirement", "competitive", "matching",
    "comprehensive",
    # Corporate marketing and non-technical culture buzzwords
    "fast-paced", "dynamic", "culture", "inclusive", "equal", "opportunity", "employer",
    "passionate", "excited", "innovative", "thriving", "mission", "values", "environment",
    "foster", "fostering", "enthusiastic",
    # Seniority markers
    "3+", "5+", "7+", "10+",
    # Conversational recruiter filler words
    "looking", "seeking", "join", "team", "teams", "company", "solid", "strong", "good",
    "exposure", "familiarity", "knowledge", "understanding", "proficiency", "plus", "like", "using",
    # Generic action verbs and recruiter filler
    "write", "develop", "developing", "integrate", "integrating", "improve", "improving",
    "implement", "implementing", "assess", "assessing", "prioritize", "prioritizing",
    "coordinate", "coordinating", "co-ordinate", "understand", "provide",
    "providing", "components", "elements", "feature", "features", "functionality", "performance",
    "programs", "requests", "responsiveness", "solutions", "systems", "system",
    "user", "users", "user-facing", "internal", "external", "overall", "effective",
    "existing", "popular", "least", "one", "spirit",
    "framework", "frameworks", "tools", "tool", "expertise", "technical",
    "code", "mapping", "protection", "object-relational", "applications",
    "technologies", "technology", "skills", "skill"
}

TECH_TAXONOMY = {
    # Programming Languages
    "python", "java", "javascript", "typescript", "c", "c++", "c#", "go", "golang", "rust",
    "ruby", "php", "kotlin", "swift", "sql", "r", "html", "html5", "css", "bash", "shell", "dart",
    # Frameworks & Libraries
    "react", "react.js", "django", "flask", "pyramid", "node.js", "node", "express", "express.js",
    "fastapi", "spring", "spring boot", "angular", "vue", "vue.js", "next.js", "nuxt",
    "tailwind", "tailwindcss", "bootstrap", "sass", "redux", "socket.io", "websockets",
    # Databases & Storage
    "mongodb", "mysql", "postgresql", "postgres", "sqlite", "redis", "cassandra", "dynamodb",
    "elasticsearch", "mariadb", "firebase", "supabase", "oracle", "dbms", "rdbms", "nosql",
    "orm", "object-relational mapping",
    # Cloud, DevOps & Tools
    "docker", "kubernetes", "aws", "azure", "gcp", "git", "github", "gitlab", "linux", "unix",
    "postman", "render", "vscode", "vs code", "jenkins", "ci/cd", "terraform", "ansible",
    "nginx", "apache", "jwt", "oauth",
    # AI, ML & Data Science
    "nlp", "natural language processing", "scikit-learn", "sklearn", "pandas", "numpy",
    "scipy", "tensorflow", "pytorch", "keras", "opencv", "matplotlib", "seaborn",
    "text classification", "machine learning", "data analytics", "data science",
    # Core Engineering Concepts
    "data structures", "algorithms", "dsa", "oop", "object-oriented", "rest api", "rest apis",
    "restful api", "restful apis", "graphql", "microservices", "unit testing", "pytest",
    "jest", "agile", "scrum", "mvc", "api design", "schema design", "version control",
    "object-oriented programming", "front-end", "back-end", "data protection", "problem-solving"
}

MULTI_WORD_TECH_ENTITIES = [
    "natural language processing",
    "data structures",
    "text classification",
    "machine learning",
    "data analytics",
    "spring boot",
    "restful apis",
    "restful api",
    "rest apis",
    "rest api",
    "unit testing",
    "object-oriented programming",
    "version control",
    "scikit-learn",
    "node.js",
    "react.js",
    "express.js",
    "vs code",
    "ci/cd",
    "object-relational mapping",
    "data protection",
    "front-end",
    "back-end",
    "problem-solving",
]

TOOL_KNOWLEDGE_BASE = {
    "python": {"domain": "Programming Language", "explanation": "High-level programming language widely used for web development, automation, and AI/ML."},
    "java": {"domain": "Programming Language", "explanation": "Object-oriented, class-based language used for enterprise applications, backend systems, and Android."},
    "javascript": {"domain": "Programming Language", "explanation": "Core web programming language for interactive client-side and server-side applications."},
    "typescript": {"domain": "Programming Language", "explanation": "Statically typed superset of JavaScript enhancing code safety and developer tooling."},
    "c": {"domain": "Programming Language", "explanation": "Foundational procedural programming language for systems, embedded devices, and performance-critical logic."},
    "c++": {"domain": "Programming Language", "explanation": "High-performance language supporting object-oriented and low-level memory control features."},
    "c#": {"domain": "Programming Language", "explanation": "Type-safe, object-oriented language developed by Microsoft for the .NET enterprise ecosystem."},
    "sql": {"domain": "Database Querying", "explanation": "Standard structured query language for storing, manipulating, and retrieving relational data."},
    "go": {"domain": "Programming Language", "explanation": "Concurrent compiled language engineered by Google for scalable cloud services and microservices."},
    "react": {"domain": "Frontend Framework", "explanation": "Component-based JavaScript library for building fast, reactive web user interfaces."},
    "django": {"domain": "Full Stack Framework", "explanation": "Batteries-included Python web framework encouraging rapid development and secure MVC architecture."},
    "flask": {"domain": "Backend Framework", "explanation": "Lightweight, flexible Python micro-framework for building web services and RESTful APIs."},
    "pyramid": {"domain": "Python Web Framework", "explanation": "Open-source Python web framework designed for creating scalable web applications."},
    "fastapi": {"domain": "Backend Framework", "explanation": "Modern, high-performance Python framework for building REST APIs with automatic OpenAPI docs."},
    "node.js": {"domain": "Runtime Environment", "explanation": "Asynchronous event-driven JavaScript runtime for scalable server-side network applications."},
    "node": {"domain": "Runtime Environment", "explanation": "Asynchronous event-driven JavaScript runtime for scalable server-side network applications."},
    "express": {"domain": "Backend Framework", "explanation": "Minimalist web framework for Node.js providing robust HTTP routing and middleware."},
    "mongodb": {"domain": "NoSQL Database", "explanation": "Document-oriented database for flexible, JSON-like schema storage and rapid horizontal scaling."},
    "postgresql": {"domain": "Relational Database", "explanation": "Advanced open-source object-relational database with ACID guarantees and JSON query support."},
    "postgres": {"domain": "Relational Database", "explanation": "Advanced open-source object-relational database with ACID guarantees and JSON query support."},
    "mysql": {"domain": "Relational Database", "explanation": "Popular open-source relational database management system powering web applications worldwide."},
    "sqlite": {"domain": "Relational Database", "explanation": "Lightweight, zero-configuration embedded SQL database engine requiring no server process."},
    "orm": {"domain": "Database Architecture", "explanation": "Object-Relational Mapping technique converting data between relational databases and object-oriented programming."},
    "object-relational mapping": {"domain": "Database Architecture", "explanation": "Technique converting data between relational databases and object-oriented programming languages."},
    "redis": {"domain": "In-Memory Store", "explanation": "In-memory key-value data structure store used for low-latency caching, session state, and queues."},
    "docker": {"domain": "DevOps & Containerization", "explanation": "Containerization platform that packages applications and dependencies into isolated environments."},
    "kubernetes": {"domain": "Cloud & Orchestration", "explanation": "Open-source system for automating deployment, scaling, and operations of application containers."},
    "aws": {"domain": "Cloud Platform", "explanation": "Amazon Web Services platform providing on-demand compute, storage, networking, and managed databases."},
    "azure": {"domain": "Cloud Platform", "explanation": "Microsoft cloud computing service for building, testing, deploying, and managing enterprise applications."},
    "gcp": {"domain": "Cloud Platform", "explanation": "Google Cloud Platform suite of modular cloud compute, data analytics, and machine learning services."},
    "git": {"domain": "Version Control", "explanation": "Distributed version control system for tracking source code changes and collaborating in teams."},
    "github": {"domain": "Developer Platform", "explanation": "Cloud platform for Git code hosting, pull request reviews, and CI/CD automated pipelines."},
    "linux": {"domain": "Operating System", "explanation": "Open-source Unix-like operating system kernel running the majority of servers and cloud deployments."},
    "postman": {"domain": "API Platform", "explanation": "API client for designing, testing, automating, and documenting RESTful and GraphQL endpoints."},
    "render": {"domain": "Cloud Hosting", "explanation": "Unified cloud hosting platform with automated continuous deployments from Git repositories."},
    "vscode": {"domain": "Developer Tool", "explanation": "Lightweight, extensible code editor with integrated terminal, debugging, and Git support."},
    "jwt": {"domain": "Security & Auth", "explanation": "JSON Web Tokens open standard for securely transmitting verified claims between client and server."},
    "nlp": {"domain": "Artificial Intelligence", "explanation": "Natural Language Processing subfield of AI enabling computers to process, analyze, and understand human language."},
    "scikit-learn": {"domain": "Machine Learning", "explanation": "Python machine learning library providing efficient tools for classification, regression, and clustering."},
    "pandas": {"domain": "Data Analytics", "explanation": "High-performance Python library for data manipulation, cleaning, and structured tabular analysis."},
    "numpy": {"domain": "Scientific Computing", "explanation": "Fundamental Python package for scientific computing with multi-dimensional arrays and matrix math."},
    "tailwind": {"domain": "UI Styling", "explanation": "Utility-first CSS framework for composing custom responsive user interfaces directly in markup."},
    "html": {"domain": "Web Markup", "explanation": "Standard markup language defining the structure and elements of web pages."},
    "css": {"domain": "Web Styling", "explanation": "Style sheet language controlling visual layout, colors, typography, and responsive presentation."},
    "rest api": {"domain": "Architecture", "explanation": "Stateless architectural style for web services communicating over standard HTTP methods."},
    "data structures": {"domain": "Core CS", "explanation": "Specialized formats for organizing and storing data efficiently in memory (arrays, trees, graphs)."},
    "algorithms": {"domain": "Core CS", "explanation": "Step-by-step computational procedures and logic for problem solving, searching, and sorting."},
    "oop": {"domain": "Programming Paradigm", "explanation": "Object-Oriented Programming methodology structuring software as reusable classes and objects."},
    "dbms": {"domain": "Database Systems", "explanation": "Software for creating, querying, maintaining, and administering structured database tables."},
    "mvc": {"domain": "Software Architecture", "explanation": "Model-View-Controller pattern separating business logic, data models, and user presentation."},
    "unit testing": {"domain": "Quality Assurance", "explanation": "Software verification practice testing individual code modules in isolation to prevent bugs."},
    "pytest": {"domain": "Software Testing", "explanation": "Robust Python testing framework for writing expressive, readable unit and functional tests."},
    "agile": {"domain": "Methodology", "explanation": "Iterative project delivery framework emphasizing adaptive planning, team collaboration, and rapid feedback."},
    "celery": {"domain": "Background Tasks", "explanation": "Distributed task queue for Python managing asynchronous background jobs and scheduled workers."},
    "kafka": {"domain": "Event Streaming", "explanation": "High-throughput distributed event streaming platform for real-time messaging and data pipelines."},
    "graphql": {"domain": "API Query Language", "explanation": "API query language and server runtime allowing clients to request precisely the data fields needed."},
    "ci/cd": {"domain": "DevOps", "explanation": "Continuous Integration & Continuous Delivery practices automating testing, building, and deployment."}
}

SKILLS_TAXONOMY_PATH = ROOT / "backend" / "skills_taxonomy.json"
SKILLS_TAXONOMY_DATA: dict = {}
if SKILLS_TAXONOMY_PATH.is_file():
    try:
        with open(SKILLS_TAXONOMY_PATH, "r", encoding="utf-8") as _f:
            SKILLS_TAXONOMY_DATA = json.load(_f)
    except Exception as _e:
        logger.warning(f"Failed to load skills_taxonomy.json: {_e}")

if SKILLS_TAXONOMY_DATA and "technologies_by_category" in SKILLS_TAXONOMY_DATA:
    for _cat_name, _tech_list in SKILLS_TAXONOMY_DATA["technologies_by_category"].items():
        for _item in _tech_list:
            _item_clean = _item.strip().casefold()
            TECH_TAXONOMY.add(_item_clean)
            if " " in _item_clean or "." in _item_clean or "/" in _item_clean or "-" in _item_clean:
                if _item_clean not in MULTI_WORD_TECH_ENTITIES:
                    MULTI_WORD_TECH_ENTITIES.append(_item_clean)

if SKILLS_TAXONOMY_DATA and "technology_explanations" in SKILLS_TAXONOMY_DATA:
    for _tech_key, _tech_info in SKILLS_TAXONOMY_DATA["technology_explanations"].items():
        TOOL_KNOWLEDGE_BASE[_tech_key.casefold()] = {
            "domain": _tech_info.get("category", "Technical Skill"),
            "explanation": _tech_info.get("explanation", f"Key technical capability required for this role ({_tech_key})."),
        }

MULTI_WORD_TECH_ENTITIES.sort(key=len, reverse=True)
CORE_DOMAINS = SKILLS_TAXONOMY_DATA.get("core_domains", {})


def detect_core_domain(job_description: str, extracted_tech: set[str], job_title: str = "") -> dict:
    """Identifies the Core Technical Domain from the Job Description and extracted technologies.
    If no technical domain or technologies are present, signals that search should stop.
    """
    title_lower = job_title.casefold()
    jd_lower = job_description.casefold()

    domain_scores: dict[str, int] = {}
    domain_matched_tech: dict[str, list[str]] = {}

    for domain_name, data in CORE_DOMAINS.items():
        score = 0
        matched: list[str] = []
        for role in data.get("suitable_roles", []):
            if role.casefold() in title_lower:
                score += 5
            elif role.casefold() in jd_lower:
                score += 2

        primary_techs = set(t.casefold() for t in data.get("primary_technologies", []))
        overlap = extracted_tech.intersection(primary_techs)
        score += len(overlap) * 2
        matched.extend(sorted(overlap))

        domain_scores[domain_name] = score
        domain_matched_tech[domain_name] = matched

    sorted_domains = sorted(domain_scores.items(), key=lambda x: x[1], reverse=True)
    best_domain, top_score = sorted_domains[0] if sorted_domains else ("Unknown", 0)

    if top_score <= 0 and not extracted_tech:
        return {
            "domain": None,
            "confidence": "none",
            "detected": False,
            "message": "No core technical domain found in this job description. Please provide a technical job description or clarify the role."
        }

    if top_score <= 0 and extracted_tech:
        best_domain = "Software Engineering / Technical Role"

    return {
        "domain": best_domain,
        "confidence": "high" if top_score >= 4 else "moderate",
        "detected": True,
        "matched_technologies": domain_matched_tech.get(best_domain, []),
        "domain_scores": domain_scores,
    }


DEFAULT_ELIGIBLE_ROLE_CLUSTERS = [
    {
        "id": "full-stack",
        "domain": "Full Stack & Web Development",
        "eligible_titles": ["Full Stack Developer", "Associate Full Stack Engineer", "Web Developer", "React Developer"],
        "qualifying_skills": ["react", "python", "django", "flask", "node.js", "javascript", "html", "css", "mongodb", "sql"],
    },
    {
        "id": "backend-python",
        "domain": "Python & Backend Engineering",
        "eligible_titles": ["Python Developer", "Backend Developer", "Associate Software Engineer - Backend", "API Developer"],
        "qualifying_skills": ["python", "django", "flask", "node.js", "rest api", "jwt", "sql", "mongodb"],
    },
    {
        "id": "data-analytics",
        "domain": "Data Analytics & Python Applications",
        "eligible_titles": ["Python Data Analyst", "Junior Data Analyst", "Business Intelligence Associate"],
        "qualifying_skills": ["python", "pandas", "sql", "data analytics"],
    },
    {
        "id": "ai-ml",
        "domain": "Machine Learning & NLP",
        "eligible_titles": ["Junior Machine Learning Engineer", "AI / NLP Associate", "Python AI Developer"],
        "qualifying_skills": ["python", "nlp", "scikit-learn", "flask"],
    },
    {
        "id": "campus-graduate",
        "domain": "Campus & Graduate Engineering",
        "eligible_titles": ["Graduate Engineer Trainee", "Associate Software Engineer", "Junior Software Developer"],
        "qualifying_skills": ["python", "java", "c", "sql", "data structures", "oop", "dbms", "git"],
    }
]


def detect_career_stage(text: str) -> str:
    text_lower = text.casefold()
    fresher_indicators = [
        "2024", "2025", "2026", "2027", "fresher", "intern", "internship",
        "student", "pursuing", "b.tech", "undergraduate", "entry-level",
        "graduate trainee", "campus", "academic projects"
    ]
    experienced_indicators = [
        "5+ years", "6+ years", "7+ years", "8+ years", "10+ years",
        "senior software engineer", "lead software engineer", "engineering manager",
        "principal engineer", "staff software", "tech lead", "architect"
    ]
    fresher_score = sum(1 for ind in fresher_indicators if ind in text_lower)
    exp_score = sum(2 for ind in experienced_indicators if ind in text_lower)
    return "Fresher / Entry-Level" if (fresher_score > 0 and exp_score == 0) else "Experienced"


def match_tier1_role_eligibility(
    job_title: str,
    job_description: str,
    clusters: list[dict] = DEFAULT_ELIGIBLE_ROLE_CLUSTERS,
) -> dict:
    jt_lower = job_title.casefold()
    jd_lower = job_description.casefold()

    incompatible_domains = [
        ("iOS / Mobile Swift Development", ["ios developer", "swift developer", "ios engineer", "swiftui", "objective-c"]),
        ("Android Mobile Development", ["android developer", "android engineer", "kotlin developer"]),
        ("Embedded & Firmware Engineering", ["embedded engineer", "firmware engineer", "microcontroller", "rtos"]),
        ("SAP / ERP Architecture", ["sap consultant", "sap developer", "abap developer"]),
    ]
    for domain_name, kw_list in incompatible_domains:
        if any(kw in jt_lower for kw in kw_list):
            return {
                "domain": domain_name,
                "matched_role": "Specialized External Domain",
                "eligibility_status": "Incompatible Domain Warning",
                "warning": f"This role requires specialized {domain_name}. Your profile is optimized for Python Full Stack, Backend Development, and Data Analytics.",
                "eligible_clusters": [c["domain"] for c in clusters],
            }

    best_cluster = None
    best_role = ""
    status = "Moderate Eligibility"

    for cluster in clusters:
        for title in cluster["eligible_titles"]:
            t_words = [w for w in title.casefold().split() if len(w) > 2]
            if any(w in jt_lower for w in t_words) or title.casefold() in jt_lower:
                best_cluster = cluster
                best_role = title
                status = "High Eligibility"
                break
        if best_cluster:
            break

    if best_cluster is None:
        # Check domain skills in JD
        for cluster in clusters:
            matched_skills = [s for s in cluster["qualifying_skills"] if s in jd_lower]
            if len(matched_skills) >= 2:
                best_cluster = cluster
                best_role = cluster["eligible_titles"][0]
                status = "High Eligibility" if len(matched_skills) >= 4 else "Moderate Eligibility"
                break

    if best_cluster is None:
        best_cluster = clusters[0]
        best_role = clusters[0]["eligible_titles"][0]
        status = "Domain Review Recommended"

    return {
        "domain": best_cluster["domain"],
        "matched_role": best_role,
        "eligibility_status": status,
        "eligible_clusters": [c["domain"] for c in clusters],
    }


app = FastAPI(title="ResumeFit AI", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[os.environ.get("RESUMEFIT_FRONTEND_ORIGIN", "http://localhost:5173")],
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["Content-Type"],
)
DATA_ROOT.mkdir(parents=True, exist_ok=True)


class Base(DeclarativeBase):
    pass


class WorkflowRecord(Base):
    __tablename__ = "workflows"

    id: Mapped[str] = mapped_column(String(32), primary_key=True)
    payload: Mapped[str] = mapped_column(Text)
    filename: Mapped[str] = mapped_column(String(255))
    extension: Mapped[str] = mapped_column(String(8))
    created_at: Mapped[str] = mapped_column(String(40))


class WorkspaceResumeRecord(Base):
    __tablename__ = "workspace_resumes"

    id: Mapped[str] = mapped_column(String(32), primary_key=True)
    payload: Mapped[str] = mapped_column(Text)
    updated_at: Mapped[str] = mapped_column(String(40))


class CandidateProfileRecord(Base):
    __tablename__ = "candidate_profiles"

    id: Mapped[str] = mapped_column(String(32), primary_key=True)
    payload: Mapped[str] = mapped_column(Text)
    updated_at: Mapped[str] = mapped_column(String(40))


database_url = os.environ.get("DATABASE_URL", f"sqlite:///{DATA_ROOT / 'resumefit.db'}")
engine_options = {"connect_args": {"check_same_thread": False}} if database_url.startswith("sqlite:") else {}
engine = create_engine(database_url, **engine_options)
Base.metadata.create_all(engine)


class SuggestionAction(BaseModel):
    action: Literal["approve", "reject", "edit", "undo"]
    proposed_text: str | None = Field(default=None, max_length=10000)


class BulkSuggestionAction(BaseModel):
    action: Literal["approve", "reject"]


class InstructionRequest(BaseModel):
    instruction: str = Field(min_length=1, max_length=2000)


class TailorSectionRequest(BaseModel):
    section: str = Field(min_length=1, max_length=120)
    paragraph_index: int | None = None
    custom_instruction: str = Field(default="", max_length=1000)


class ToolConfirmationRequest(BaseModel):
    tool: str = Field(min_length=1, max_length=120)
    status: Literal["approved", "related", "skipped"]
    related_skill: str = Field(default="", max_length=120)



class ExportSuggestionState(BaseModel):
    id: str = Field(min_length=1, max_length=120)
    status: Literal["approved", "rejected", "pending"]
    proposed_text: str = Field(max_length=10000)


class ResumeExportRequest(BaseModel):
    suggestions: list[ExportSuggestionState] = Field(max_length=100)


class DownloadRequest(BaseModel):
    workflow_id: str = Field(min_length=32, max_length=32)
    text: str = Field(min_length=1, max_length=40000)


def get_workflow(workflow_id: str) -> tuple[WorkflowRecord, dict]:
    with Session(engine, expire_on_commit=False) as db:
        record = db.get(WorkflowRecord, workflow_id)
        if record is None:
            raise HTTPException(status_code=404, detail="This tailoring session was not found.")
        return record, json.loads(record.payload)


def save_workflow(record: WorkflowRecord, payload: dict) -> None:
    try:
        with Session(engine, expire_on_commit=False) as db:
            stored = db.get(WorkflowRecord, record.id)
            if stored is None:
                raise HTTPException(status_code=404, detail="This tailoring session was not found.")
            stored.payload = json.dumps(payload, ensure_ascii=False)
            db.commit()
    except SQLAlchemyError as exc:
        raise HTTPException(status_code=500, detail="The tailoring session could not be saved. Your original file is unchanged.") from exc


def flatten_docx_paragraphs(doc: Document) -> list[Paragraph]:
    result: list[Paragraph] = []
    seen_cells: set[object] = set()

    def collect(parent) -> None:
        for child in parent._element.iterchildren():
            if child.tag.endswith("}p"):
                result.append(Paragraph(child, parent))
            elif child.tag.endswith("}tbl"):
                table = DocxTable(child, parent)
                for row in table.rows:
                    for cell in row.cells:
                        if cell._tc in seen_cells:
                            continue
                        seen_cells.add(cell._tc)
                        collect(cell)

    collect(doc._body)
    return result


def paragraph_xml_text(paragraph: Paragraph) -> str:
    text: list[str] = []
    for element in paragraph._p.iter():
        if element.tag == qn("w:t"):
            text.append(element.text or "")
        elif element.tag == qn("w:tab"):
            text.append("\t")
        elif element.tag == qn("w:br"):
            text.append("\n")
    return "".join(text)


def extract_resume(path: Path, extension: str) -> tuple[list[str], int | None]:
    if extension == ".docx":
        try:
            doc = Document(path)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, PackageNotFoundError) as exc:
            raise HTTPException(
                status_code=422,
                detail="This DOCX file could not be opened. Please upload a valid, unencrypted DOCX.",
            ) from exc
        paragraphs = flatten_docx_paragraphs(doc)
        texts = [paragraph_xml_text(paragraph).strip() for paragraph in paragraphs]
        return texts, None

    try:
        pdf = fitz.open(path)
        page_count = len(pdf)
        texts = [
            line.strip()
            for page in pdf
            for line in page.get_text("text").splitlines()
        ]
        pdf.close()
    except (RuntimeError, ValueError, fitz.FileDataError) as exc:
        raise HTTPException(status_code=422, detail="This PDF could not be opened or is encrypted.") from exc
    if not any(texts):
        raise HTTPException(
            status_code=422,
            detail="This PDF contains no selectable text. It may be scanned; OCR is not enabled in this local build. Upload a DOCX or text-based PDF.",
        )
    return texts, page_count


def initialize_master_resume(path: Path = MASTER_RESUME_PATH) -> dict | None:
    if not path.is_file():
        try:
            with Session(engine) as db:
                existing = db.get(WorkspaceResumeRecord, "default")
                if existing is not None:
                    db.delete(existing)
                    db.commit()
        except SQLAlchemyError as exc:
            raise RuntimeError("The stale master resume could not be cleared from the workspace database.") from exc
        app.state.master_resume = None
        return None
    try:
        paragraphs, page_count = extract_resume(path, ".docx")
    except HTTPException as exc:
        raise RuntimeError(f"The configured master resume could not be loaded: {path}") from exc
    if len("\n".join(paragraphs).strip()) < 20:
        raise RuntimeError(f"The configured master resume contains too little readable text: {path}")
    master = {
        "path": path,
        "filename": path.name,
        "extension": ".docx",
        "paragraphs": paragraphs,
        "page_count": page_count,
    }
    try:
        with Session(engine) as db:
            record = db.get(WorkspaceResumeRecord, "default")
            payload = json.dumps(
                {
                    "filename": master["filename"],
                    "extension": master["extension"],
                    "paragraphs": master["paragraphs"],
                    "page_count": master["page_count"],
                    "source_path": str(path),
                },
                ensure_ascii=False,
            )
            if record is None:
                record = WorkspaceResumeRecord(id="default", payload=payload)
                db.add(record)
            else:
                record.payload = payload
            record.updated_at = datetime.now(timezone.utc).isoformat()
            db.commit()
    except SQLAlchemyError as exc:
        raise RuntimeError("The master resume could not be loaded into the workspace database.") from exc
    app.state.master_resume = master
    return master


initialize_master_resume()


def section_for(text: str, active_section: str) -> str:
    cleaned = re.sub(r"[^a-z ]", "", text.casefold()).strip()
    text_lower = text.strip().casefold()
    if any(text_lower.startswith(p) for p in (
        "languages:", "programming languages:", "frontend:", "backend:",
        "database:", "databases:", "tools:", "developer tools:", "skills:",
        "technical skills:", "core competencies:"
    )):
        return "Skills"
    if any(text_lower.startswith(p) for p in ("smart india hackathon", "hackathon", "awards:", "honors:")):
        return "Achievements"
    for section, aliases in SECTION_ALIASES.items():
        if any(cleaned == alias or cleaned.startswith(alias + " ") for alias in aliases):
            return section
    return active_section


CANONICAL_TOOL_NAMES = {
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL",
    "mongodb": "MongoDB",
    "mysql": "MySQL",
    "redis": "Redis",
    "sqlite": "SQLite",
    "react": "React",
    "node.js": "Node.js",
    "django": "Django",
    "flask": "Flask",
    "fastapi": "FastAPI",
    "express": "Express",
    "spring boot": "Spring Boot",
    "python": "Python",
    "java": "Java",
    "javascript": "JavaScript",
    "typescript": "TypeScript",
    "c": "C",
    "c++": "C++",
    "c#": "C#",
    "html": "HTML",
    "css": "CSS",
    "git": "Git",
    "github": "GitHub",
    "vs code": "VS Code",
    "render": "Render",
    "aws": "AWS",
    "gcp": "GCP",
    "azure": "Azure",
    "sql": "SQL",
    "rest api": "REST APIs",
    "rest apis": "REST APIs",
    "jwt": "JWT",
    "nlp": "NLP",
    "ci/cd": "CI/CD",
    "oop": "OOP",
    "dbms": "DBMS",
}


def get_canonical_skill_name(tool_name: str) -> str:
    normalized = tool_name.strip().casefold()
    if normalized in CANONICAL_TOOL_NAMES:
        return CANONICAL_TOOL_NAMES[normalized]
    if normalized.upper() in {"SQL", "AWS", "GCP", "API", "JWT", "NLP", "CSS", "HTML", "REST", "CI/CD", "OOP", "DBMS"}:
        return normalized.upper()
    return tool_name.strip().title()


def get_skill_category_target(tool_name: str) -> str:
    normalized = tool_name.strip().casefold()
    if SKILLS_TAXONOMY_DATA and "technologies_by_category" in SKILLS_TAXONOMY_DATA:
        for cat, items in SKILLS_TAXONOMY_DATA["technologies_by_category"].items():
            if any(it.casefold() == normalized for it in items):
                if cat == "programming_languages":
                    return "languages"
                if cat == "web_and_frontend":
                    return "frontend"
                if cat in ("backend_and_frameworks", "architectures_and_protocols"):
                    return "backend"
                if cat == "databases_and_storage":
                    return "database"
                if cat in ("devops_and_cloud", "tools_and_environments", "ai_and_data_science"):
                    return "tools"
    if normalized in {"python", "java", "javascript", "c", "c++", "go", "rust", "typescript", "php", "ruby", "c#", "kotlin", "swift"}:
        return "languages"
    if normalized in {"react", "html", "css", "vue", "angular", "tailwind", "redux", "next.js"}:
        return "frontend"
    if normalized in {"django", "flask", "node.js", "express", "fastapi", "spring boot", "rest apis", "rest api", "jwt", "graphql"}:
        return "backend"
    if normalized in {"sql", "mongodb", "postgresql", "postgres", "mysql", "redis", "sqlite", "nosql"}:
        return "database"
    return "tools"


def build_additive_skills_suggestion(
    skills_items: list[dict],
    target_tools_to_add: list[str],
) -> dict | None:
    if not skills_items or not target_tools_to_add:
        return None

    category_tools: dict[str, list[str]] = {
        "languages": [],
        "frontend": [],
        "backend": [],
        "database": [],
        "tools": [],
    }
    for tool in target_tools_to_add:
        cat = get_skill_category_target(tool)
        category_tools.setdefault(cat, []).append(get_canonical_skill_name(tool))

    for item in skills_items:
        text = item["text"]
        text_lower = text.casefold()
        matched_cat = None
        for cat in ("tools", "database", "backend", "frontend", "languages"):
            if cat in text_lower or (cat == "languages" and "language" in text_lower):
                matched_cat = cat
                break

        tools_for_this_item = []
        if matched_cat and category_tools.get(matched_cat):
            tools_for_this_item = [t for t in category_tools[matched_cat] if t.casefold() not in text_lower]
        elif not matched_cat:
            tools_for_this_item = [
                get_canonical_skill_name(t)
                for t in target_tools_to_add
                if t.casefold() not in text_lower
            ]

        if tools_for_this_item:
            base_text = text.rstrip().rstrip(",;.")
            proposed_text = base_text + ", " + ", ".join(tools_for_this_item)
            return {
                "id": suggestion_id("Skills", item["index"]),
                "section": "Skills",
                "paragraph_index": item["index"],
                "original_text": text,
                "proposed_text": proposed_text,
                "suggested_text": proposed_text,
                "reason": (
                    f"All of your original skills are 100% preserved. "
                    f"Matched technical requirements ({', '.join(tools_for_this_item)}) were added."
                ),
                "requirement": ", ".join(tools_for_this_item),
                "expected_impact": "Directly incorporates target technical competencies while maintaining original layout.",
                "status": "pending",
                "source": "manual",
                "is_skills_addition": True,
            }

    return None


def resume_sections(paragraphs: list[str]) -> dict[str, list[dict]]:
    result: dict[str, list[dict]] = {section: [] for section in SECTIONS}
    active_section = "Professional Summary"
    for index, text in enumerate(paragraphs):
        if not text:
            continue
        detected = section_for(text, active_section)
        if detected != active_section or any(
            re.sub(r"[^a-z ]", "", text.casefold()).strip() == alias
            for aliases in SECTION_ALIASES.values()
            for alias in aliases
        ):
            active_section = detected
            is_heading = any(
                re.sub(r"[^a-z ]", "", text.casefold()).strip() == alias
                for aliases in SECTION_ALIASES.values()
                for alias in aliases
            )
            if is_heading:
                continue
        result.setdefault(active_section, []).append({"index": index, "text": text})
    return result


def extract_keywords(text: str) -> set[str]:
    keywords: set[str] = set()
    text_lower = text.casefold()

    # 1. Multi-word technical entities first (matches with word boundaries)
    for entity in MULTI_WORD_TECH_ENTITIES:
        pattern = rf"(?<![a-z0-9]){re.escape(entity)}(?![a-z0-9])"
        if re.search(pattern, text_lower):
            keywords.add(entity)

    # 2. Strict matching against TECH_TAXONOMY only - prevents extracting general English words
    for match in KEYWORD_PATTERN.finditer(text):
        word = match.group().casefold()
        if word in TECH_TAXONOMY:
            keywords.add(word)
        elif word.startswith("specialty") and word[9:].isdigit():
            keywords.add(word)

    return keywords


CATEGORY_WEIGHTS = {
    "skills": 0.30,
    "responsibilities_and_experience": 0.25,
    "keyword_alignment": 0.15,
    "projects": 0.10,
    "education_and_certifications": 0.10,
    "preferred_qualifications": 0.10,
}

CATEGORY_LABELS = {
    "skills": "Skills",
    "responsibilities_and_experience": "Responsibilities & Experience",
    "keyword_alignment": "Keyword Alignment",
    "projects": "Projects",
    "education_and_certifications": "Education & Certifications",
    "preferred_qualifications": "Preferred Qualifications",
}


def find_keyword_evidence(keyword: str, text: str, aliases: list[str] | None = None) -> tuple[str, str, str]:
    """Finds evidence snippet, classification, and explanation note for a keyword in resume text.
    
    Classifies as: 'fully evidenced', 'partially evidenced', 'uncertain', or 'not evidenced'.
    """
    pattern = re.compile(rf"(?<![A-Za-z0-9])({re.escape(keyword)})(?![A-Za-z0-9])", re.IGNORECASE)
    for line in text.splitlines():
        line_clean = line.strip()
        if not line_clean:
            continue
        if pattern.search(line_clean):
            return line_clean[:220], "fully evidenced", f"Direct verbatim match found in resume text for '{keyword}'."

    if aliases:
        for alias in aliases:
            alias_pattern = re.compile(rf"(?<![A-Za-z0-9])({re.escape(alias)})(?![A-Za-z0-9])", re.IGNORECASE)
            for line in text.splitlines():
                line_clean = line.strip()
                if not line_clean:
                    continue
                if alias_pattern.search(line_clean):
                    return line_clean[:220], "partially evidenced", f"Evidenced via related-skill equivalent '{alias}'."

    kw_lower = keyword.casefold()
    for line in text.splitlines():
        line_clean = line.strip()
        if not line_clean:
            continue
        if kw_lower in line_clean.casefold() and len(kw_lower) >= 3:
            return line_clean[:220], "uncertain", f"Partial mention of '{keyword}' in text; verified context is ambiguous."

    return "Not found in resume text", "not evidenced", f"No verified evidence found for '{keyword}'."


def categorize_jd_requirements(job_description: str, requirements: set[str]) -> dict[str, set[str]]:
    sentences = [s.strip() for s in re.split(r"[\n;•]+|\.\s+", job_description) if s.strip()]

    preferred_keywords: set[str] = set()
    edu_keywords: set[str] = set()
    proj_keywords: set[str] = set()
    exp_keywords: set[str] = set()
    skill_keywords: set[str] = set()

    pref_triggers = {"preferred", "plus", "bonus", "nice to have", "ideal", "desired", "advantage", "optional"}
    edu_triggers = {"bachelor", "bachelors", "degree", "bs", "ba", "master", "masters", "ms", "phd", "mba", "diploma", "education", "certified", "certification", "certifications", "license", "licensed"}
    proj_triggers = {"project", "projects", "portfolio", "prototype", "initiative", "architecture", "designed", "design", "built", "building", "develop", "developed", "pipeline", "implementation"}
    exp_triggers = {"responsible", "responsibilities", "duties", "experience", "years", "role", "work", "lead", "manage", "collaborate", "audit", "coordinate", "benchmark", "benchmarks", "productivity", "stakeholder", "stakeholders", "audits", "quality", "performs", "communicates", "communication", "meets", "capacity", "thinking", "attention", "detail", "problem-solving"}
    known_skill_words = {"sql", "excel", "python", "react", "typescript", "javascript", "c++", "c#", "java", "docker", "kubernetes", "aws", "azure", "gcp", "git", "ci/cd", ".net", "macros", "mathematics", "analytics", "reporting", "dashboards", "linux", "database", "api", "rest", "graphql", "html", "css", "tailwind", "fastapi", "postgres", "postgresql", "sqlite"}

    for s in sentences:
        s_lower = s.casefold()
        s_words = extract_keywords(s) & requirements
        if any(trig in s_lower for trig in pref_triggers):
            preferred_keywords.update(s_words)
        if any(trig in s_lower for trig in edu_triggers):
            edu_keywords.update(s_words)
        if any(trig in s_lower for trig in proj_triggers):
            proj_keywords.update(s_words)
        if any(trig in s_lower for trig in exp_triggers) or "perform" in s_lower or "lead" in s_lower:
            exp_keywords.update(s_words)
        if "skill" in s_lower or "tool" in s_lower or "technolog" in s_lower or "proficient" in s_lower or "knowledge" in s_lower:
            skill_keywords.update(s_words)

    for kw in requirements:
        if kw in known_skill_words or "sql" in kw or "excel" in kw or "data" in kw or kw in SINGLE_LETTER_KEYWORDS:
            skill_keywords.add(kw)
        if kw in exp_triggers:
            exp_keywords.add(kw)

    return {
        "skills": skill_keywords,
        "responsibilities_and_experience": exp_keywords,
        "keyword_alignment": set(requirements),
        "projects": proj_keywords,
        "education_and_certifications": edu_keywords,
        "preferred_qualifications": preferred_keywords,
    }


def score_resume(
    resume_text: str,
    job_description: str,
    job_title: str = "",
    confirmed_tools: list[str] | None = None,
    skipped_tools: list[str] | None = None,
    career_stage: str | None = None,
) -> dict:
    requirements = extract_keywords(job_description)
    if not requirements:
        raise HTTPException(
            status_code=422,
            detail="The job description needs more detail to calculate a match. Please ensure it contains technical skills or requirements."
        )

    core_domain_info = detect_core_domain(job_description, requirements, job_title)
    if not core_domain_info["detected"]:
        raise HTTPException(
            status_code=422,
            detail=core_domain_info["message"]
        )

    confirmed_set = {t.casefold() for t in (confirmed_tools or [])}
    skipped_set = {t.casefold() for t in (skipped_tools or [])}

    resume_keywords = extract_keywords(resume_text)
    effective_resume_keywords = resume_keywords | confirmed_set

    exact_matches = requirements.intersection(effective_resume_keywords)
    semantic_matches = {
        keyword: sorted(RELATED_KEYWORD_ALIASES[keyword].intersection(effective_resume_keywords))
        for keyword in requirements - exact_matches
        if RELATED_KEYWORD_ALIASES.get(keyword, set()).intersection(effective_resume_keywords)
    }
    matched = sorted(exact_matches | semantic_matches.keys())
    matched_set = set(matched)
    missing = sorted((requirements - matched_set) - skipped_set)

    stage = career_stage or detect_career_stage(resume_text)
    tier1_role_match = match_tier1_role_eligibility(job_title, job_description)

    discovered_tools = []
    missing_technologies = []
    for req in sorted(requirements):
        kb_entry = TOOL_KNOWLEDGE_BASE.get(
            req,
            {"domain": "Technical Skill", "explanation": f"Key technical capability required for this role ({req})."}
        )
        tool_status = "approved" if req in confirmed_set else ("skipped" if req in skipped_set else ("matched" if req in matched_set else "pending"))
        entry = {
            "tool": req,
            "name": req.title() if len(req) > 3 else req.upper(),
            "domain": kb_entry.get("domain", kb_entry.get("category", "Technical Skill")),
            "explanation": kb_entry["explanation"],
            "status": tool_status,
        }
        discovered_tools.append(entry)
        if tool_status == "pending":
            missing_technologies.append(entry)

    categories_reqs = categorize_jd_requirements(job_description, requirements)
    overall_match_ratio = len(matched) / len(requirements) if requirements else 1.0

    category_results = {}
    weighted_sum = 0.0

    for cat_key, weight in CATEGORY_WEIGHTS.items():
        cat_reqs = categories_reqs.get(cat_key, set())
        effective_cat_reqs = cat_reqs - skipped_set
        if effective_cat_reqs:
            cat_matched = effective_cat_reqs & matched_set
            if stage == "Fresher / Entry-Level" and cat_key == "preferred_qualifications":
                cat_score = max(round(100 * len(cat_matched) / len(effective_cat_reqs)), round(100 * overall_match_ratio))
            else:
                cat_score = round(100 * len(cat_matched) / len(effective_cat_reqs))
            cat_matched_count = len(cat_matched)
            cat_req_count = len(effective_cat_reqs)
        else:
            cat_score = round(100 * overall_match_ratio)
            cat_matched_count = len(matched)
            cat_req_count = len(requirements)

        weighted_points = round(cat_score * weight, 1)
        weighted_sum += cat_score * weight

        category_results[cat_key] = {
            "name": CATEGORY_LABELS[cat_key],
            "weight": weight,
            "weight_percent": int(weight * 100),
            "score": cat_score,
            "weighted_points": weighted_points,
            "matched_count": cat_matched_count,
            "requirement_count": cat_req_count,
            "requirements": sorted(cat_reqs) if cat_reqs else sorted(requirements)[:10],
        }

    final_score = max(0, min(100, round(weighted_sum)))

    evidence_items = []
    for req in sorted(requirements):
        primary_cat = "Keyword Alignment"
        for cat_key in ("skills", "responsibilities_and_experience", "education_and_certifications", "projects", "preferred_qualifications"):
            if req in categories_reqs.get(cat_key, set()):
                primary_cat = CATEGORY_LABELS[cat_key]
                break

        if req in confirmed_set:
            snippet = "Candidate verified proficiency via dashboard."
            classification = "fully evidenced"
            note = f"Directly approved by candidate as a confirmed skill ({req})."
        elif req in skipped_set:
            snippet = "Requirement skipped by candidate."
            classification = "uncertain"
            note = f"Candidate skipped '{req}'; not claimed and not penalized."
        else:
            snippet, classification, note = find_keyword_evidence(
                req, resume_text, semantic_matches.get(req)
            )

        evidence_items.append({
            "requirement": req,
            "category": primary_cat,
            "classification": classification,
            "evidence": snippet,
            "notes": note,
        })

    category_scores_list = [
        {
            "category": v["name"],
            "weight": v["weight_percent"],
            "score": v["score"],
            "weighted_score": v["weighted_points"],
            "matched_count": v["matched_count"],
            "total_count": v["requirement_count"],
            "matched_keywords": [r for r in v["requirements"] if r in matched_set],
            "missing_keywords": [r for r in v["requirements"] if r not in matched_set and r not in skipped_set],
        }
        for v in category_results.values()
    ]
    evidence_breakdown_list = [
        {
            "requirement": item["requirement"],
            "category": item["category"],
            "classification": item["classification"],
            "snippets": [item["evidence"]] if item.get("evidence") and item["evidence"] != "Not found in resume text" else [],
            "evidence": item["evidence"],
            "notes": item["notes"],
            "confidence": 1.0 if item["classification"] == "fully evidenced" else 0.7 if item["classification"] == "partially evidenced" else 0.4 if item["classification"] == "uncertain" else 0.0,
        }
        for item in evidence_items
    ]

    return {
        "score": final_score,
        "matched_keywords": matched,
        "missing_keywords": missing,
        "matched_count": len(matched),
        "requirement_count": len(requirements),
        "exact_matched_count": len(exact_matches),
        "semantic_matched_count": len(semantic_matches),
        "semantic_matches": semantic_matches,
        "categories": category_results,
        "category_scores": category_scores_list,
        "evidence_items": evidence_items,
        "evidence_breakdown": evidence_breakdown_list,
        "career_stage": stage,
        "core_domain": core_domain_info["domain"],
        "core_domain_info": core_domain_info,
        "missing_technologies": missing_technologies,
        "tier1_role_match": tier1_role_match,
        "discovered_tools": discovered_tools,
        "confirmed_tools": sorted(confirmed_set),
        "skipped_tools": sorted(skipped_set),
        "method": "Deterministic weighted resume-match scoring: Skills (30%), Responsibilities & Experience (25%), Keyword Alignment (15%), Projects (10%), Education & Certifications (10%), Preferred Qualifications (10%). Related matches use curated phrase-level proxies, not machine-understood meaning. This is not an ATS score or a prediction of interview outcomes.",
    }


def suggestion_id(section: str, index: int = 0) -> str:
    return f"{re.sub(r'[^a-z0-9]+', '-', section.casefold()).strip('-')}-{index}"


def _demo_rewrite_section(section: str, text: str, missing_keywords: list[str]) -> str:
    """Rule-based rewrite that weaves missing JD keywords into existing text.

    Stays within the character budget of the original text so that DOCX
    in-place replacements remain valid. Returns the original unchanged
    when no keywords can be fitted.
    """
    if not missing_keywords:
        return text
    budget = len(text)

    if section in ("Skills",):
        base = text.rstrip().rstrip(",;.")
        for count in range(min(6, len(missing_keywords)), 0, -1):
            candidate = base + ", " + ", ".join(missing_keywords[:count])
            if len(candidate) <= budget:
                return candidate

    elif section == "Professional Summary":
        base = text.rstrip().rstrip(".")
        for count in range(min(4, len(missing_keywords)), 0, -1):
            subset = missing_keywords[:count]
            if len(subset) > 1:
                phrase = ", ".join(subset[:-1]) + " and " + subset[-1]
            else:
                phrase = subset[0]
            candidate = f"{base}, with {phrase} exposure."
            if len(candidate) <= budget:
                return candidate

    elif section in ("Experience", "Internships", "Projects"):
        base = text.rstrip().rstrip(".")
        for count in range(min(3, len(missing_keywords)), 0, -1):
            subset = missing_keywords[:count]
            candidate = f"{base}, leveraging {', '.join(subset)}."
            if len(candidate) <= budget:
                return candidate

    return text


def make_demo_suggestions(paragraphs: list[str], job_description: str) -> list[dict]:
    sections = resume_sections(paragraphs)
    jd_words = extract_keywords(job_description)
    all_resume_words = extract_keywords("\n".join(paragraphs))
    missing_from_resume = sorted(jd_words - all_resume_words)
    total_overlap = jd_words.intersection(all_resume_words)
    suggestions = []
    if total_overlap:
        for section in SECTIONS:
            candidates = sections.get(section, [])
            if not candidates:
                continue
            if section == "Skills":
                tech_tools_to_add = [
                    t for t in missing_from_resume
                    if t in TECH_TAXONOMY or t in TOOL_KNOWLEDGE_BASE
                ]
                if tech_tools_to_add:
                    additive_skills = build_additive_skills_suggestion(candidates, tech_tools_to_add)
                    if additive_skills:
                        additive_skills["source"] = "demo"
                        suggestions.append(additive_skills)
                        continue
            evidence = max(
                candidates,
                key=lambda item: len(
                    jd_words.intersection(extract_keywords(item["text"]))
                ),
            )
            overlap = sorted(jd_words.intersection(extract_keywords(evidence["text"])))
            if not overlap:
                continue
            requirement = ", ".join(overlap[:4])
            proposed = _demo_rewrite_section(section, evidence["text"], missing_from_resume)
            has_rewrite = proposed != evidence["text"]
            if has_rewrite:
                proposed_words = extract_keywords(proposed)
                original_words = extract_keywords(evidence["text"])
                newly_added = sorted(jd_words & proposed_words - original_words)
            else:
                newly_added = []
            suggestions.append({
                "id": suggestion_id(section),
                "section": section,
                "paragraph_index": evidence["index"],
                "original_text": evidence["text"],
                "suggested_text": proposed,
                "proposed_text": proposed,
                "reason": (
                    f"This existing resume content shares the JD terms “{requirement}.” "
                    + (f"Missing JD keywords ({', '.join(newly_added[:5])}) have been woven in. " if newly_added else "This existing text is available for manual editing; add details that reflect your experience. ")
                    + "Review this rewrite carefully — approve only if it accurately reflects your experience."
                ),
                "requirement": requirement + (f", {', '.join(newly_added[:4])}" if newly_added else ""),
                "expected_impact": (
                    f"Incorporating {len(newly_added)} additional JD term(s) may improve the match estimate."
                    if newly_added
                    else "No match improvement is claimed. Use this as a starting point for manual edits."
                ),
                "status": "pending",
                "source": "demo",
            })
    if not suggestions:
        candidates = [
            (section, item)
            for section in SECTIONS
            for item in sections.get(section, [])
        ]
        if candidates:
            section, evidence = max(
                candidates,
                key=lambda candidate: len(
                    extract_keywords(candidate[1]["text"])
                ),
            )
            suggestions.append({
                "id": suggestion_id(section),
                "section": section,
                "paragraph_index": evidence["index"],
                "original_text": evidence["text"],
                "suggested_text": "",
                "proposed_text": "",
                "reason": (
                    "No job-description terms are supported by this source text. "
                    "This existing text is available for manual editing; add only details "
                    "that accurately reflect your experience."
                ),
                "requirement": "",
                "expected_impact": (
                    "No match improvement is claimed. Use this as a starting point for "
                    "manual edits, then review every change."
                ),
                "status": "pending",
                "source": "demo",
            })
    return suggestions


def fallback_to_manual_suggestions(
    suggestions: list[dict],
    paragraphs: list[str],
    job_description: str,
) -> tuple[list[dict], list[str]]:
    if suggestions:
        return suggestions, []
    fallback = make_demo_suggestions(paragraphs, job_description)
    if not fallback:
        return [], [
            "No evidence-supported AI edit was returned, and no editable resume text was extracted."
        ]
    manual_suggestion = max(
        fallback,
        key=lambda suggestion: len(suggestion["original_text"]),
    )
    manual_suggestion.update({
        "source": "manual",
        "suggested_text": "",
        "proposed_text": "",
        "reason": (
            "No evidence-supported AI edit was returned. Edit this existing resume text "
            "yourself, and include only details that accurately reflect your experience."
        ),
        "requirement": "",
        "expected_impact": (
            "No match improvement is claimed. The score changes only if approved wording "
            "accurately includes additional job-description terms."
        ),
    })
    return [manual_suggestion], [
        "Ollama did not return a valid, evidence-supported rewrite. Your original is unchanged; edit the manual review card to continue."
    ]


def request_ollama_json(request: Request) -> dict:
    try:
        with urlopen(request, timeout=OLLAMA_TIMEOUT_SECONDS) as response:
            return json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        if exc.code == 404:
            detail = (
                f"Ollama could not find model '{OLLAMA_MODEL}' (HTTP 404). "
                f"Pull it with `ollama pull {OLLAMA_MODEL}` and retry."
            )
        else:
            detail = f"Ollama request failed (HTTP {exc.code}). Check the local Ollama server and retry."
        raise HTTPException(status_code=502, detail=detail) from exc
    except TimeoutError as exc:
        raise HTTPException(
            status_code=504,
            detail=(
                "The local Ollama model did not finish generating within "
                f"{OLLAMA_TIMEOUT_SECONDS} seconds. The response is capped at "
                f"{OLLAMA_MAX_RESPONSE_TOKENS} tokens; try again or use a smaller Ollama model."
            ),
        ) from exc
    except URLError as exc:
        raise HTTPException(
            status_code=502,
            detail=(
                "The local Ollama server is unavailable. Start Ollama on "
                "localhost:11434 and retry."
            ),
        ) from exc
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise HTTPException(
            status_code=502,
            detail="Ollama returned an invalid API response. No resume changes were applied.",
        ) from exc


def _extract_suggestions_from_ollama_response(content: str) -> list[dict]:
    if not isinstance(content, str) or not content.strip():
        raise ValueError("No Ollama content returned.")
    match = re.search(r"```(?:json)?\s*(.*?)\s*```", content, flags=re.IGNORECASE | re.DOTALL)
    clean_content = match.group(1) if match else content.strip()

    # 1. Standard full parse
    try:
        parsed = json.loads(clean_content)
        if isinstance(parsed, dict) and isinstance(parsed.get("suggestions"), list):
            return parsed["suggestions"]
    except Exception:
        pass

    # 2. Repair truncated JSON if num_predict capped output mid-stream
    last_brace = clean_content.rfind("}")
    if last_brace != -1:
        truncated_prefix = clean_content[:last_brace + 1]
        for candidate_ending in ("]\n}", "\n  ]\n}", "\n}"):
            try:
                candidate = truncated_prefix + candidate_ending
                parsed = json.loads(candidate)
                if isinstance(parsed, dict) and isinstance(parsed.get("suggestions"), list) and parsed["suggestions"]:
                    return parsed["suggestions"]
            except Exception:
                continue

    # 3. Regex rescue for completed { "section": ... } objects
    rescued = []
    for item_match in re.finditer(r'\{[^{}]*"(?:section|original_content)"[^{}]*\}', clean_content, flags=re.DOTALL):
        try:
            item = json.loads(item_match.group(0))
            if isinstance(item, dict) and "section" in item:
                rescued.append(item)
        except Exception:
            continue
    if rescued:
        return rescued

    raise ValueError("Ollama returned an invalid suggestions response.")


def call_ollama_suggestions(resume_text: str, job_description: str, instructions: str = "") -> tuple[list[dict], list[str]]:
    requirements = extract_keywords(job_description)
    resume_words = extract_keywords(resume_text)
    missing_keywords = sorted(requirements - resume_words)
    source_sections = resume_sections(resume_text.splitlines())
    # Keep the prompt inside the actual local runner context. The scoring and
    # review screens still retain the complete source text; this is only the
    # bounded evidence sent to the local model.
    ranked_sources = sorted(
        (
            (section, item)
            for section, entries in source_sections.items()
            for item in entries
        ),
        key=lambda candidate: (
            len(extract_keywords(candidate[1]["text"]).intersection(requirements)),
            len(candidate[1]["text"]),
        ),
        reverse=True,
    )[:6]
    compact_sections: dict[str, list[dict]] = {section: [] for section in SECTIONS}
    for section, item in ranked_sources:
        compact_sections[section].append({
            "index": item["index"],
            "paragraph_index": item["index"],
            "text": item["text"][:900],
        })
    context = {
        "resume_sections": {
            section: [
                {
                    "paragraph_index": item["paragraph_index"],
                    "text": item["text"],
                    "max_characters": len(item["text"]),
                    "max_words": len(item["text"].split()),
                }
                for item in compact_sections.get(section, [])
            ]
            for section in SECTIONS
        },
        "job_description": job_description,
        "job_description_terms": sorted(requirements)[:60],
        "missing_job_description_keywords": missing_keywords[:40],
        "user_instructions": instructions,
    }
    context_json = json.dumps(context, ensure_ascii=False, separators=(",", ":"))
    prompt = f"""You are an expert resume tailoring assistant and career editor. Analyze the complete JSON context included below in this single request. Do not make separate calls or assume content omitted from that context.

CRITICAL INSTRUCTIONS:
1. Review the bounded, most relevant source paragraphs in resume_sections. Return only the {OLLAMA_MAX_SUGGESTIONS} most relevant suggestions with directly supported evidence across every Experience bullet, Projects entry, and Education entry. Skip generic or short fragments that cannot be concisely improved. Never combine separate sections into one suggestion, and do not invent a section or source paragraph.
2. Every proposal must be meaningfully different from its source paragraph; never return the original unchanged, "NO CHANGE", or a punctuation-only edit. If no truthful improvement is possible, omit it.
3. Treat user_instructions as the requested editing guidance and apply it to every relevant suggestion. Explicitly honor requested emphasis, wording, and focus where compatible with the resume and layout constraints. User instructions never authorize invented facts or unsupported claims. Compare each source paragraph against all missing_job_description_keywords and job_description in the context. Naturally use missing terms only when the resume_sections explicitly support the skill or experience. Pay special attention to supported evidence for terms such as analytical capacity, critical thinking, attention to detail, macros, and mathematics. Do not infer skills or invent qualifications.
4. CRITICAL LAYOUT CONSTRAINT: The final output must fit on exactly ONE page. Keep proposed_content concise; aim for no more than 80% of max_characters; omit that section instead of returning an over-budget rewrite. Do not add words just to repeat job-description keywords. Return no line breaks in replacement text.
5. Do not invent skills, employers, duties, achievements, metrics, degrees, or credentials. Do not keyword-stuff or promise a score. Keep each explanation brief.
6. Keep every explanation and impact estimate brief (under 15 words). Return valid JSON only, matching this exact object shape:
{{
  "suggestions": [
    {{
      "section": "One of: {", ".join(SECTIONS)}",
      "original_content": "Exact source paragraph to replace",
      "proposed_content": "Complete rewritten paragraph, different from original_content",
      "why_this_may_help": "Brief factual explanation",
      "job_requirement": "Relevant supported job-description term(s)",
      "expected_impact": "Brief, non-guaranteed estimate"
    }}
  ]
}}
Use all six keys for each suggestion. proposed_content must be a complete replacement, not an instruction or fragment."""
    unified_prompt = (
        f"{prompt}\n\n"
        "COMPLETE_RESUME_AND_JOB_CONTEXT_JSON:\n"
        f"{context_json}"
    )
    request = Request(
        OLLAMA_API_URL,
        data=json.dumps({
            "model": OLLAMA_MODEL,
            "prompt": unified_prompt,
            "stream": False,
            "format": "json",
            "keep_alive": "30m",
            "options": {
                "temperature": 0.35,
                "num_predict": OLLAMA_MAX_RESPONSE_TOKENS,
                "num_ctx": OLLAMA_NUM_CTX,
            },
        }).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        data = request_ollama_json(request)
        content = data["response"]
        raw_suggestions = _extract_suggestions_from_ollama_response(content)
    except (KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=502, detail="Ollama returned an invalid suggestions response. No resume changes were applied.") from exc
    if not isinstance(raw_suggestions, list):
        raise HTTPException(status_code=502, detail="Ollama returned an invalid suggestions list. No resume changes were applied.")

    paragraphs = resume_text.splitlines()
    sections = resume_sections(paragraphs)
    validated = []
    warnings = []
    seen_sections = set()
    for raw in raw_suggestions[:OLLAMA_MAX_SUGGESTIONS]:
        if not isinstance(raw, dict):
            continue
        section = raw.get("section")
        original = raw.get("original_content", raw.get("original_text"))
        proposed = raw.get("proposed_content", raw.get("proposed_text"))
        if section not in SECTIONS or section in seen_sections or not isinstance(original, str) or not isinstance(proposed, str):
            continue
        original_clean = original.strip()
        proposed_clean = proposed.strip()
        if not original_clean or not proposed_clean or len(proposed_clean) > 10000:
            continue
        if proposed_clean == original_clean:
            warnings.append(
                f"The Ollama suggestion for {section} was unchanged, so it was omitted."
            )
            continue

        candidates = sections.get(section, [])
        matching = None
        for item in candidates:
            if item["text"] == original:
                matching = item
                break
        if matching is None:
            for item in candidates:
                if item["text"].strip() == original_clean:
                    matching = item
                    break
        if matching is None:
            for item in candidates:
                if original_clean in item["text"].strip() or item["text"].strip() in original_clean:
                    matching = item
                    break
        if matching is None and candidates:
            best_match = max(
                candidates,
                key=lambda item: SequenceMatcher(None, item["text"].strip(), original_clean).ratio(),
            )
            if SequenceMatcher(None, best_match["text"].strip(), original_clean).ratio() >= 0.45:
                matching = best_match

        if matching is None:
            continue
        max_budget = len(matching["text"])
        if len(proposed_clean) > max_budget:
            words = proposed_clean.split()
            while words and len(" ".join(words)) > max_budget:
                words.pop()
            trimmed = " ".join(words).rstrip(" ,;.")
            if trimmed and not trimmed.endswith((".", "!", "?")):
                if len(trimmed) + 1 <= max_budget:
                    trimmed += "."
            if len(trimmed) <= max_budget and len(trimmed) > 5:
                proposed_clean = trimmed
            else:
                warnings.append(
                    f"The Ollama suggestion for {section} exceeded the {max_budget}-character budget, so it was omitted."
                )
                continue
        seen_sections.add(section)
        validated.append({
            "id": suggestion_id(section),
            "section": section,
            "paragraph_index": matching["index"],
            "original_text": matching["text"],
            "suggested_text": proposed_clean,
            "proposed_text": proposed_clean,
            "reason": str(raw.get("why_this_may_help", raw.get("reason", "Review this suggestion against your source resume.")))[:1200],
            "requirement": str(raw.get("job_requirement", raw.get("requirement", "")))[:500],
            "expected_impact": str(raw.get("expected_impact", "Potentially improves visibility of verified experience."))[:500],
            "status": "pending",
            "source": "ollama",
        })
    return validated, warnings


ORIGINAL_CALL_OLLAMA_SUGGESTIONS = call_ollama_suggestions



def call_ollama_single_section(
    section: str,
    original_text: str,
    job_description: str,
    custom_instruction: str = "",
) -> tuple[dict | None, list[str]]:
    requirements = extract_keywords(job_description)
    resume_words = extract_keywords(original_text)
    missing = sorted(requirements - resume_words)[:15]
    max_chars = len(original_text)

    prompt = f"""You are an expert resume editor. Tailor this single resume text block for the target job.
SECTION: {section}
ORIGINAL TEXT ({max_chars} characters):
"{original_text}"

TARGET KEYWORDS TO INCORPORATE IF TRUTHFULLY SUPPORTED:
{", ".join(missing[:10])}

CRITICAL CONSTRAINT: Output length MUST NOT EXCEED {max_chars} characters. You must replace, rephrase, or condense words from the original text so the total length is at most {max_chars} characters. Do not simply append.
{f"USER GUIDANCE: {custom_instruction}" if custom_instruction else ""}

Return valid JSON only matching this exact schema:
{{
  "proposed_content": "Rewritten text strictly under {max_chars} characters",
  "why_this_may_help": "Brief one-sentence reason",
  "job_requirement": "Target keyword addressed",
  "expected_impact": "Brief one-sentence impact"
}}"""

    request = Request(
        OLLAMA_API_URL,
        data=json.dumps({
            "model": OLLAMA_MODEL,
            "prompt": prompt,
            "stream": False,
            "format": "json",
            "keep_alive": "30m",
            "options": {
                "temperature": 0.25,
                "num_predict": 150,
                "num_ctx": OLLAMA_NUM_CTX,
            },
        }).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        data = request_ollama_json(request)
        content = data.get("response", "")
        match = re.search(r"```(?:json)?\s*(.*?)\s*```", content, flags=re.IGNORECASE | re.DOTALL)
        clean = match.group(1) if match else content.strip()
        parsed = json.loads(clean)
        if isinstance(parsed, dict) and "suggestions" in parsed and isinstance(parsed["suggestions"], list) and parsed["suggestions"]:
            parsed = parsed["suggestions"][0]
        proposed = str(parsed.get("proposed_content", parsed.get("proposed_text", ""))).strip()
        if not proposed:
            return None, ["Model returned empty proposed text."]
        if len(proposed) > max_chars:
            words = proposed.split()
            while words and len(" ".join(words)) > max_chars:
                words.pop()
            trimmed = " ".join(words).rstrip(" ,;.") + "."
            if len(trimmed) <= max_chars and len(trimmed) > 5:
                proposed = trimmed
            else:
                return None, [f"Model rewrite ({len(proposed)} chars) exceeded the {max_chars}-character budget."]

        result = {
            "id": suggestion_id(section),
            "section": section,
            "original_text": original_text,
            "proposed_text": proposed,
            "suggested_text": proposed,
            "reason": str(parsed.get("why_this_may_help", "Tailored to align with target job requirements."))[:500],
            "requirement": str(parsed.get("job_requirement", ", ".join(missing[:3])))[:200],
            "expected_impact": str(parsed.get("expected_impact", "Enhances alignment with key job requirements."))[:300],
            "status": "pending",
            "source": "ollama",
        }
        return result, []
    except Exception as exc:
        return None, [f"Section tailoring error: {exc}"]


def clean_filename(filename: str) -> str:
    base = Path(filename.replace("\\", "/")).name
    cleaned = re.sub(r"[^A-Za-z0-9._ -]", "", base).strip(" .")
    return cleaned[:120] or "resume"


def invalidate_exports(workflow_id: str, payload: dict) -> None:
    folder = DATA_ROOT / workflow_id
    for key in ("tailored.docx", "tailored.pdf"):
        path = folder / key
        if path.exists():
            path.unlink()
    payload["finalized"] = False


def inspect_pdf_layout(path: Path) -> dict:
    try:
        pdf = fitz.open(path)
    except (RuntimeError, ValueError, fitz.FileDataError) as exc:
        raise HTTPException(status_code=500, detail="The generated PDF could not be checked for layout issues.") from exc
    blank_pages = []
    overflow_blocks = 0
    overlapping_blocks = 0
    for page_number, page in enumerate(pdf, start=1):
        if not page.get_text("text").strip():
            blank_pages.append(page_number)
        blocks = page.get_text("blocks")
        bounds = page.rect
        text_blocks = [
            (block[0], block[1], block[2], block[3])
            for block in blocks
            if len(block) > 4 and isinstance(block[4], str) and block[4].strip()
        ]
        for x0, y0, x1, y1 in text_blocks:
            if x0 < -0.5 or y0 < -0.5 or x1 > bounds.width + 0.5 or y1 > bounds.height + 0.5:
                overflow_blocks += 1
        for index, first in enumerate(text_blocks):
            for second in text_blocks[index + 1:]:
                intersection_width = min(first[2], second[2]) - max(first[0], second[0])
                intersection_height = min(first[3], second[3]) - max(first[1], second[1])
                if intersection_width > 1 and intersection_height > 1:
                    overlapping_blocks += 1
    page_count = len(pdf)
    pdf.close()
    return {
        "page_count": page_count,
        "blank_pages": blank_pages,
        "overflow_blocks": overflow_blocks,
        "overlapping_blocks": overlapping_blocks,
    }


def check_layout(payload: dict, source_path: Path, tailored_pdf: Path) -> list[str]:
    original_pdf = None
    original_page_count = payload["page_count"]
    warnings = []
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if payload["extension"] == ".docx" and soffice:
        try:
            result = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(source_path.parent), str(source_path)],
                capture_output=True,
                text=True,
                timeout=45,
                check=False,
            )
        except subprocess.TimeoutExpired:
            result = None
            warnings.append("Comparing the original DOCX layout timed out; review pagination manually.")
        original_pdf = source_path.parent / "original.pdf"
        if result is not None and result.returncode == 0 and original_pdf.is_file():
            original_page_count = inspect_pdf_layout(original_pdf)["page_count"]
        else:
            original_pdf = None
            warnings.append("The original DOCX page count could not be checked; review pagination manually.")
    tailored = inspect_pdf_layout(tailored_pdf)
    if original_page_count is not None and tailored["page_count"] != original_page_count:
        warnings.append(
            f"Page count changed from {original_page_count} to {tailored['page_count']}; review pagination before sending."
        )
    if tailored["blank_pages"]:
        warnings.append(f"Possible blank page(s): {', '.join(str(page) for page in tailored['blank_pages'])}.")
    if tailored["overflow_blocks"]:
        warnings.append(f"{tailored['overflow_blocks']} text block(s) extend beyond the PDF page bounds.")
    if tailored["overlapping_blocks"]:
        warnings.append(f"{tailored['overlapping_blocks']} text block overlap(s) were detected.")
    if original_page_count is None:
        if not warnings:
            warnings.append("Original and tailored page counts could not be compared automatically; LibreOffice is not installed.")
    if payload["extension"] == ".pdf":
        warnings.append("PDF-source layout cannot be preserved. Check the reflowed export for alignment, fonts, links, and page breaks.")
    else:
        warnings.append("Visual alignment, hyperlinks, and text wrapping still require a manual check in a document viewer.")
    if original_pdf and original_pdf.exists():
        original_pdf.unlink()
    return warnings


def replace_docx_paragraph(paragraph: Paragraph, new_text: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    old_text = "".join(node.text or "" for node in text_nodes)
    if any(child.tag == qn("w:hyperlink") for child in paragraph._p.iterchildren()):
        if new_text != old_text:
            raise HTTPException(
                status_code=422,
                detail="Text in an existing hyperlink cannot be changed without altering its link target.",
            )
        return
    if old_text != paragraph.text:
        raise HTTPException(
            status_code=422,
            detail="This paragraph contains inline layout elements that cannot be safely replaced.",
        )
    if any(
        element.tag in {
            qn("w:tab"),
            qn("w:br"),
            qn("w:drawing"),
            qn("w:object"),
            qn("w:fldChar"),
            qn("w:instrText"),
        }
        for element in paragraph._p.iter()
    ):
        raise HTTPException(
            status_code=422,
            detail="This paragraph contains non-text layout elements that cannot be safely replaced.",
        )

    # Clean additive append (pure addition of matched skills):
    # Preserves 100% of existing runs, fonts, colors, font sizes, and alignments without character splitting!
    old_stripped = old_text.rstrip().rstrip(".")
    new_stripped = new_text.rstrip().rstrip(".")
    if text_nodes and new_stripped.startswith(old_stripped) and len(new_stripped) > len(old_stripped):
        appended_text = new_stripped[len(old_stripped):]
        text_nodes[-1].text = (text_nodes[-1].text or "") + appended_text
        if text_nodes[-1].text[:1].isspace() or text_nodes[-1].text[-1:].isspace():
            text_nodes[-1].set(qn("xml:space"), "preserve")
        return

    node_indexes: list[int] = []
    for index, node in enumerate(text_nodes):
        node_indexes.extend([index] * len(node.text or ""))
    char_node_indexes: list[int | None] = [None] * len(new_text)
    for tag, old_start, old_end, new_start, new_end in SequenceMatcher(
        a=old_text, b=new_text, autojunk=False
    ).get_opcodes():
        old_count = old_end - old_start
        new_count = new_end - new_start
        if tag == "equal":
            for offset in range(new_count):
                char_node_indexes[new_start + offset] = node_indexes[old_start + offset]
        elif tag == "replace" and old_count:
            for offset in range(new_count):
                source_offset = min(old_count - 1, offset * old_count // max(new_count, 1))
                char_node_indexes[new_start + offset] = node_indexes[old_start + source_offset]
        elif tag == "insert" and node_indexes:
            source_index = old_start - 1 if old_start else old_start
            char_node_indexes[new_start:new_end] = [node_indexes[source_index]] * new_count

    if not text_nodes and new_text:
        raise HTTPException(
            status_code=422,
            detail="This empty paragraph has no existing text location available for an in-place replacement.",
        )
    if text_nodes and any(index is None for index in char_node_indexes):
        fallback_index = node_indexes[-1] if node_indexes else 0
        char_node_indexes = [
            fallback_index if index is None else index
            for index in char_node_indexes
        ]

    updated_text = ["" for _ in text_nodes]
    for character, node_index in zip(new_text, char_node_indexes):
        if node_index is not None:
            updated_text[node_index] += character
    for node, text in zip(text_nodes, updated_text):
        node.text = text
        if text[:1].isspace() or text[-1:].isspace():
            node.set(qn("xml:space"), "preserve")
        else:
            node.attrib.pop(qn("xml:space"), None)

    embed_hyperlinks(paragraph)


HYPERLINK_PATTERN = re.compile(
    r"(?i)(?:https?://|www\.)[^\s<>]+|(?:linkedin|github)\.com(?:/[^\s<>]*)?|"
    r"[\w.+-]+@[\w.-]+\.[a-z]{2,}"
)


def normalized_hyperlink_target(text: str) -> str:
    candidate = text.rstrip(".,;:!?)]}")
    if "@" in candidate and not candidate.casefold().startswith(("http://", "https://")):
        return f"mailto:{candidate}"
    if candidate.casefold().startswith("www."):
        return f"https://{candidate}"
    if candidate.casefold().startswith(("http://", "https://")):
        return candidate
    return f"https://{candidate}"


def style_hyperlink_run(run_element) -> None:
    properties = run_element.find(qn("w:rPr"))
    if properties is None:
        properties = OxmlElement("w:rPr")
        run_element.insert(0, properties)
    color = properties.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        properties.append(color)
    color.set(qn("w:val"), "0563C1")
    underline = properties.find(qn("w:u"))
    if underline is None:
        underline = OxmlElement("w:u")
        properties.append(underline)
    underline.set(qn("w:val"), "single")


def embed_hyperlinks(paragraph: Paragraph) -> None:
    for run_element in list(paragraph._p.iterchildren(qn("w:r"))):
        text_nodes = run_element.findall(f".//{qn('w:t')}")
        if len(text_nodes) != 1:
            continue
        text_node = text_nodes[0]
        source_text = text_node.text or ""
        matches = []
        for match in HYPERLINK_PATTERN.finditer(source_text):
            target_text = match.group().rstrip(".,;:!?)]}")
            if target_text:
                matches.append((match.start(), match.start() + len(target_text), target_text))
        if not matches:
            continue
        if any(
            child.tag not in {qn("w:rPr"), qn("w:t")}
            for child in run_element.iterchildren()
        ):
            raise HTTPException(
                status_code=422,
                detail="A generated URL shares a run with non-text content and cannot be linked safely.",
            )

        parent = run_element.getparent()
        insertion_index = parent.index(run_element)
        segments = []
        offset = 0
        for start, end, link_text in matches:
            if start > offset:
                segments.append((False, source_text[offset:start]))
            segments.append((True, link_text))
            offset = end
        if offset < len(source_text):
            segments.append((False, source_text[offset:]))

        parent.remove(run_element)
        for is_link, segment in segments:
            if not segment:
                continue
            if not is_link:
                replacement_run = deepcopy(run_element)
                replacement_run.find(qn("w:t")).text = segment
                parent.insert(insertion_index, replacement_run)
            else:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(
                    qn("r:id"),
                    paragraph.part.relate_to(
                        normalized_hyperlink_target(segment),
                        RT.HYPERLINK,
                        is_external=True,
                    ),
                )
                hyperlink_run = deepcopy(run_element)
                hyperlink_text = hyperlink_run.find(qn("w:t"))
                hyperlink_text.text = segment
                if segment[:1].isspace() or segment[-1:].isspace():
                    hyperlink_text.set(qn("xml:space"), "preserve")
                style_hyperlink_run(hyperlink_run)
                hyperlink.append(hyperlink_run)
                parent.insert(insertion_index, hyperlink)
            insertion_index += 1

    active_text = "".join(
        "".join(node.text or "" for node in hyperlink.xpath(".//w:t"))
        for hyperlink in paragraph._p.xpath(".//w:hyperlink")
    )
    for match in HYPERLINK_PATTERN.finditer(paragraph_xml_text(paragraph)):
        link_text = match.group().rstrip(".,;:!?)]}")
        if link_text and link_text not in active_text:
            raise HTTPException(
                status_code=422,
                detail="A generated contact URL spans multiple formatting runs and could not be linked safely.",
            )


def create_docx_from_pdf_text(texts: list[str], destination: Path) -> None:
    doc = Document()
    doc.sections[0].top_margin = Inches(0.65)
    doc.sections[0].bottom_margin = Inches(0.65)
    for text in texts:
        if text:
            paragraph = doc.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(4)
    try:
        doc.save(destination)
    except OSError as exc:
        raise HTTPException(status_code=500, detail="The DOCX file could not be written. Check available disk space and try again.") from exc


def approved_paragraphs(payload: dict) -> list[str]:
    texts = list(payload["paragraphs"])
    for suggestion in payload.get("suggestions", []):
        if suggestion.get("status") == "approved":
            idx = suggestion.get("paragraph_index")
            original = (suggestion.get("original_text") or "").strip()
            proposed = suggestion.get("proposed_text")
            if proposed is None:
                continue
            replaced = False
            if idx is not None and 0 <= idx < len(texts):
                if not original or texts[idx].strip() == original:
                    texts[idx] = proposed
                    replaced = True
            if not replaced and original:
                for i, text in enumerate(texts):
                    if text.strip() == original:
                        texts[i] = proposed
                        replaced = True
                        break
                if not replaced:
                    for i, text in enumerate(texts):
                        if original in text.strip() or text.strip() in original:
                            texts[i] = proposed
                            replaced = True
                            break
            if not replaced and idx is not None and 0 <= idx < len(texts):
                texts[idx] = proposed
    return texts


def validate_approved_rewrites(payload: dict) -> None:
    for suggestion in payload.get("suggestions", []):
        if suggestion.get("status") != "approved":
            continue
        proposed = normalize_export_text(suggestion.get("proposed_text") or "")
        original = normalize_export_text(suggestion.get("original_text") or "")
        suggested = normalize_export_text(suggestion.get("suggested_text") or "")
        if not proposed or (not suggested and proposed == original):
            raise HTTPException(
                status_code=422,
                detail="The approved suggestion has no rewritten text. Edit it with a truthful change or reject it.",
            )


def update_tailored_state(payload: dict) -> None:
    tailored_paras = approved_paragraphs(payload)
    tailored_text = "\n".join(tailored_paras).strip()
    score_details = score_resume(
        tailored_text,
        payload["job_description"],
        job_title=payload.get("job_title", ""),
        confirmed_tools=payload.get("confirmed_tools", []),
        skipped_tools=payload.get("skipped_tools", []),
    )
    score = score_details["score"]
    payload["tailored_text"] = tailored_text
    payload["tailored_score_details"] = score_details
    payload["tailored_score"] = score
    payload["final_text"] = tailored_text
    payload["final_score_details"] = score_details
    payload["final_score"] = score


def normalize_export_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"^\s{0,3}#{1,6}\s+", "", text)
    text = re.sub(r"[•●▪◦‣⁃]", " - ", text)
    text = re.sub(r"\*\*(.*?)\*\*|__(.*?)__", lambda match: match.group(1) or match.group(2), text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)|(?<!_)_([^_]+)_(?!_)", lambda match: match.group(1) or match.group(2), text)
    text = text.replace("`", "")
    text = text.translate(str.maketrans({
        "–": "-",
        "—": "-",
        "−": "-",
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
    }))
    text = "".join(
        character for character in text
        if unicodedata.category(character) not in {"Cc", "Cf", "Cs"}
    )
    return re.sub(r"\s+", " ", text).strip()


def pdf_compatible_text(text: str) -> str:
    normalized = normalize_export_text(text)
    return normalized.encode("cp1252", errors="replace").decode("cp1252")


def normalize_download_text(text: str) -> list[str]:
    normalized = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    lines = []
    for line in normalized.split("\n"):
        line = re.sub(r"^\s{0,3}#{1,6}\s+", "", line)
        line = re.sub(r"^\s*[-*+•●▪◦‣⁃]\s+", "- ", line)
        line = re.sub(r"\*\*(.*?)\*\*|__(.*?)__", lambda match: match.group(1) or match.group(2), line)
        line = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)|(?<!_)_([^_]+)_(?!_)", lambda match: match.group(1) or match.group(2), line)
        line = line.replace("`", "")
        line = "".join(
            " " if unicodedata.category(character) == "Cc" else character
            for character in line
            if unicodedata.category(character) not in {"Cf", "Cs"}
        )
        lines.append(re.sub(r"[ \t]+", " ", line).strip())
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    return lines


def validate_docx_text_budget(payload: dict) -> None:
    replacements: dict[int, str] = {}
    for suggestion in payload["suggestions"]:
        if suggestion["status"] != "approved":
            continue
        index = suggestion["paragraph_index"]
        if index < 0 or index >= len(payload["paragraphs"]):
            raise HTTPException(status_code=500, detail="A source paragraph could not be located; no export was produced.")
        if index in replacements:
            raise HTTPException(status_code=422, detail="Multiple approved changes target the same paragraph.")
        original = payload["paragraphs"][index]
        proposed = normalize_export_text(suggestion["proposed_text"])
        suggestion["proposed_text"] = proposed
        is_skills_addition = bool(suggestion.get("is_skills_addition"))
        original_base = original.rstrip(" \t\r\n.,;")
        if is_skills_addition and not proposed.startswith(original_base):
            is_skills_addition = False
        if len(proposed) > len(original) and not is_skills_addition:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"The approved replacement for {suggestion['section']} exceeds its "
                    f"{len(original)}-character layout budget. Shorten it before exporting."
                ),
            )
        if any(character in proposed for character in "\r\n\t"):
            raise HTTPException(
                status_code=422,
                detail="Approved DOCX replacements cannot introduce new line or tab structure.",
            )
        replacements[index] = proposed


def generate_docx(payload: dict, source_path: Path, destination: Path) -> None:
    if payload["extension"] != ".docx":
        create_docx_from_pdf_text(approved_paragraphs(payload), destination)
    else:
        validate_docx_text_budget(payload)
        try:
            doc = Document(source_path)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, PackageNotFoundError) as exc:
            raise HTTPException(status_code=500, detail="The original DOCX could not be reopened for export.") from exc
        paragraphs = flatten_docx_paragraphs(doc)
        for suggestion in payload["suggestions"]:
            if suggestion["status"] != "approved":
                continue
            index = suggestion["paragraph_index"]
            target_p = None
            if 0 <= index < len(paragraphs):
                orig = (suggestion.get("original_text") or "").strip()
                if not orig or paragraph_xml_text(paragraphs[index]).strip() == orig:
                    target_p = paragraphs[index]
            if target_p is None and suggestion.get("original_text"):
                orig = suggestion["original_text"].strip()
                for p in paragraphs:
                    if paragraph_xml_text(p).strip() == orig:
                        target_p = p
                        break
            if target_p is None and 0 <= index < len(paragraphs):
                target_p = paragraphs[index]
            if target_p is None:
                raise HTTPException(status_code=500, detail="A source paragraph could not be located; no export was produced.")
            replace_docx_paragraph(target_p, suggestion["proposed_text"])
        try:
            doc.save(destination)
        except OSError as exc:
            raise HTTPException(status_code=500, detail="The DOCX file could not be written. Check available disk space and try again.") from exc


def create_pdf_from_text(texts: list[str], destination: Path) -> None:
    output = build_pdf_bytes(texts)
    try:
        destination.write_bytes(output)
    except OSError as exc:
        raise HTTPException(status_code=500, detail="The PDF could not be written. Your original resume is unchanged.") from exc


def build_pdf_bytes(texts: list[str]) -> bytes:
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "Resume",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        alignment=TA_LEFT,
        spaceAfter=2,
        textColor=colors.HexColor("#1d2b3c"),
    )
    story = []
    for source_text in texts:
        text = pdf_compatible_text(source_text)
        if not text:
            continue
        elif text.strip().casefold() in {
            alias for aliases in SECTION_ALIASES.values() for alias in aliases
        }:
            story.append(Spacer(1, 3))
            story.append(PdfParagraph(f"<b>{html.escape(text)}</b>", normal))
        else:
            story.append(PdfParagraph(html.escape(text).replace("\n", "<br/>"), normal))
    output = io.BytesIO()
    try:
        SimpleDocTemplate(
            output,
            pagesize=letter,
            leftMargin=0.7 * inch,
            rightMargin=0.7 * inch,
            topMargin=0.55 * inch,
            bottomMargin=0.55 * inch,
            title="ResumeFit AI tailored resume",
        ).build(story)
    except (OSError, ValueError, LayoutError) as exc:
        raise HTTPException(status_code=500, detail="The PDF could not be generated. Your original resume is unchanged.") from exc
    return output.getvalue()


def generate_pdf(payload: dict, tailored_docx: Path, destination: Path) -> None:
    if payload["extension"] == ".docx":
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            try:
                result = subprocess.run(
                    [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(destination.parent), str(tailored_docx)],
                    capture_output=True,
                    text=True,
                    timeout=45,
                    check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise HTTPException(status_code=504, detail="DOCX-to-PDF conversion timed out. Retry or export the DOCX instead.") from exc
            generated = destination.parent / f"{tailored_docx.stem}.pdf"
            if result.returncode != 0 or not generated.exists():
                raise HTTPException(status_code=500, detail="LibreOffice could not export this DOCX to PDF.")
            if generated != destination:
                generated.replace(destination)
            return
    create_pdf_from_text(approved_paragraphs(payload), destination)


def response_payload(record: WorkflowRecord, payload: dict) -> dict:
    update_tailored_state(payload)
    ai_provider = next(
        (
            provider
            for provider in ("ollama", "gemini")
            if any(
                suggestion.get("source") == provider
                for suggestion in payload["suggestions"]
            )
        ),
        None,
    )
    return {
        "id": record.id,
        "job_id": record.id,
        "workflow_id": record.id,
        "status": payload.get("status", "completed"),
        "progress": payload.get("progress", 100),
        "filename": record.filename,
        "extension": record.extension,
        "created_at": record.created_at,
        "job_title": payload["job_title"],
        "company_name": payload["company_name"],
        "demo_mode": payload.get("demo_mode", ai_provider is None),
        "ai_provider": ai_provider,
        "limitations": payload["limitations"],
        "analysis": payload["analysis"],
        "suggestions": payload["suggestions"],
        "suggestion_warnings": payload["suggestion_warnings"],
        "instructions": payload["instructions"],
        "finalized": payload["finalized"],
        "final_score": payload.get("final_score"),
        "final_score_details": payload.get("final_score_details"),
        "final_text": payload.get("final_text"),
        "tailored_text": payload.get("tailored_text", payload.get("final_text")),
        "tailored_score": payload.get("tailored_score", payload.get("final_score")),
        "tailored_score_details": payload.get("tailored_score_details", payload.get("final_score_details")),
        "layout_review": payload["layout_review"],
        "original_text": "\n".join(payload["paragraphs"]),
        "download_urls": {
            "original": f"/api/workflows/{record.id}/original",
            "docx": f"/api/workflows/{record.id}/download/docx",
            "pdf": f"/api/workflows/{record.id}/download/pdf",
        } if payload["finalized"] else None,
    }


analysis_jobs_lock = threading.Lock()
analysis_jobs: dict[str, dict] = {}


def get_empty_profile() -> dict:
    return {
        "full_name": "",
        "headline": "",
        "email": "",
        "phone": "",
        "location": "",
        "career_stage": "Fresher / Entry-Level",
        "eligible_job_titles": [],
        "summary": "",
        "skills": [],
        "categorized_skills": {},
        "experience": [],
        "education": [],
        "projects": [],
        "certifications": [],
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def get_default_candidate_profile() -> dict:
    return {
        "full_name": "Jaya Sri Palli",
        "headline": "Full Stack Developer | Python Enthusiast",
        "email": "jayasripalli@example.com",
        "phone": "+91 9876543210",
        "location": "India",
        "career_stage": "Fresher / Entry-Level",
        "eligible_job_titles": [
            "Full Stack Developer",
            "Python Developer",
            "Backend Developer",
            "Junior Machine Learning Engineer",
            "Python Data Analyst",
            "Associate Software Engineer",
            "Graduate Engineer Trainee",
        ],
        "summary": "Motivated Information Technology graduate with hands-on full stack development and machine learning internship experience. Proficient in Python, React.js, Django, Flask, Node.js, and MongoDB.",
        "skills": [
            "Python", "JavaScript", "SQL", "HTML", "CSS", "C",
            "React", "Django", "Flask", "Node.js", "Express.js",
            "MongoDB", "SQLite", "MySQL", "PostgreSQL",
            "Git", "GitHub", "Postman", "VS Code", "Docker", "Render",
            "Scikit-learn", "NLP", "Pandas", "Machine Learning", "Data Analytics",
            "Data Structures", "Algorithms", "REST APIs", "Object-Oriented Programming"
        ],
        "categorized_skills": {
            "Languages": ["Python", "JavaScript", "SQL", "HTML", "CSS", "C"],
            "Frameworks & Libraries": ["React", "Django", "Flask", "Node.js", "Express.js"],
            "Databases": ["MongoDB", "SQLite", "MySQL", "PostgreSQL"],
            "Developer Tools & Cloud": ["Git", "GitHub", "Postman", "VS Code", "Docker", "Render"],
            "AI / ML & Analytics": ["Scikit-learn", "NLP", "Pandas", "Machine Learning", "Data Analytics"],
            "Core CS Concepts": ["Data Structures", "Algorithms", "REST APIs", "Object-Oriented Programming (OOP)"]
        },
        "experience": [
            {
                "title": "Full Stack Developer Intern",
                "company": "Larsen & Toubro (L&T)",
                "location": "Remote",
                "dates": "05/2025 - 07/2025",
                "bullets": [
                    "Engineered responsive full-stack modules using React.js and Django REST Framework.",
                    "Implemented MongoDB document schemas and integrated secure JWT authentication endpoints."
                ]
            },
            {
                "title": "Python Developer Intern",
                "company": "Infosys Springboard",
                "location": "Remote",
                "dates": "01/2025 - 03/2025",
                "bullets": [
                    "Constructed interactive YouTube Analytics Dashboard using Python, Pandas, and data processing pipelines."
                ]
            }
        ],
        "education": [
            {
                "institution": "SVCET",
                "degree": "B.Tech in Information Technology",
                "field_of_study": "Information Technology",
                "graduation_year": "2026 (CGPA: 8.77)"
            }
        ],
        "projects": [
            {
                "name": "Fake News Detection System",
                "description": "Built NLP text classification model using Scikit-learn achieving 90% accuracy; deployed full stack application with Flask API and React frontend.",
                "technologies": ["Python", "NLP", "Scikit-learn", "Flask", "React"]
            },
            {
                "name": "AI Health Chatbot",
                "description": "Engineered responsive health chatbot assistant utilizing Django backend, NLP response matching, and SQLite database storage.",
                "technologies": ["Django", "NLP", "SQLite", "Bootstrap"]
            }
        ],
        "certifications": [
            "Infosys Certified Python Full Stack Developer",
            "L&T EduTech Certified MongoDB Associate"
        ],
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def sanitize_candidate_profile(raw: dict) -> dict:
    empty = get_empty_profile()
    if not isinstance(raw, dict):
        return empty
    return {
        "full_name": str(raw.get("full_name") or "")[:100].strip(),
        "headline": str(raw.get("headline") or "")[:150].strip(),
        "email": str(raw.get("email") or "")[:100].strip(),
        "phone": str(raw.get("phone") or "")[:40].strip(),
        "location": str(raw.get("location") or "")[:100].strip(),
        "career_stage": str(raw.get("career_stage") or "Fresher / Entry-Level")[:60].strip(),
        "eligible_job_titles": [str(t)[:80].strip() for t in raw.get("eligible_job_titles", []) if str(t).strip()][:20],
        "summary": str(raw.get("summary") or "")[:3000].strip(),
        "skills": [str(s)[:60].strip() for s in raw.get("skills", []) if str(s).strip()][:60],
        "categorized_skills": {str(k)[:60]: [str(s)[:60].strip() for s in v if str(s).strip()] for k, v in raw.get("categorized_skills", {}).items() if isinstance(v, list)},
        "experience": [
            {
                "title": str(e.get("title") or "")[:100].strip(),
                "company": str(e.get("company") or "")[:100].strip(),
                "location": str(e.get("location") or "")[:100].strip(),
                "dates": str(e.get("dates") or "")[:60].strip(),
                "bullets": [str(b)[:600].strip() for b in e.get("bullets", []) if str(b).strip()][:15],
            }
            for e in raw.get("experience", [])
            if isinstance(e, dict)
        ][:20],
        "education": [
            {
                "institution": str(ed.get("institution") or "")[:120].strip(),
                "degree": str(ed.get("degree") or "")[:120].strip(),
                "field_of_study": str(ed.get("field_of_study") or "")[:120].strip(),
                "graduation_year": str(ed.get("graduation_year") or "")[:30].strip(),
            }
            for ed in raw.get("education", [])
            if isinstance(ed, dict)
        ][:10],
        "projects": [
            {
                "name": str(pr.get("name") or "")[:100].strip(),
                "description": str(pr.get("description") or "")[:800].strip(),
                "technologies": [str(t)[:40].strip() for t in pr.get("technologies", []) if str(t).strip()][:20],
            }
            for pr in raw.get("projects", [])
            if isinstance(pr, dict)
        ][:15],
        "certifications": [str(c)[:120].strip() for c in raw.get("certifications", []) if str(c).strip()][:20],
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def extract_profile_from_paragraphs(paragraphs: list[str]) -> dict:
    sections = resume_sections(paragraphs)
    all_text = "\n".join(paragraphs)

    email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", all_text)
    email = email_match.group(0) if email_match else ""

    phone_match = re.search(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", all_text)
    phone = phone_match.group(0) if phone_match else ""

    name = ""
    for p in paragraphs[:5]:
        p_clean = p.strip()
        if p_clean and p_clean not in SECTIONS and not re.search(r"@[a-zA-Z0-9-]+\.", p_clean) and len(p_clean) < 60:
            name = p_clean
            break
    if not name and paragraphs:
        name = paragraphs[0].strip()

    summary_paras = [item["text"] for item in sections.get("Professional Summary", [])]
    summary = "\n".join(summary_paras).strip()

    headline = ""
    if len(paragraphs) > 1:
        p2 = paragraphs[1].strip()
        if p2 != name and p2 not in SECTIONS and len(p2) < 80 and not re.search(r"@[a-zA-Z0-9-]+\.", p2):
            headline = p2

    skills = []
    for item in sections.get("Skills", []):
        for chunk in re.split(r"[,•|\n;]+", item["text"]):
            cleaned = chunk.strip(" -•*")
            if cleaned and len(cleaned) < 50 and cleaned not in skills:
                skills.append(cleaned)

    experience = []
    for item in sections.get("Experience", []) + sections.get("Internships", []):
        raw = item["text"].strip()
        if raw:
            experience.append({
                "title": "",
                "company": "",
                "location": "",
                "dates": "",
                "bullets": [raw],
            })

    education = []
    for item in sections.get("Education", []):
        raw = item["text"].strip()
        if raw:
            education.append({
                "institution": "",
                "degree": raw,
                "field_of_study": "",
                "graduation_year": "",
            })

    projects = []
    for item in sections.get("Projects", []):
        raw = item["text"].strip()
        if raw:
            projects.append({
                "name": raw[:60],
                "description": raw,
                "technologies": [],
            })

    certifications = [
        item["text"].strip()
        for item in sections.get("Certifications", [])
        if item["text"].strip()
    ]

    return {
        "full_name": name,
        "headline": headline,
        "email": email,
        "phone": phone,
        "location": "",
        "summary": summary,
        "skills": skills,
        "experience": experience,
        "education": education,
        "projects": projects,
        "certifications": certifications,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def profile_to_paragraphs(profile: dict) -> list[str]:
    paragraphs: list[str] = []
    name = profile.get("full_name", "").strip()
    if name:
        paragraphs.append(name)
    contact_parts = [
        profile.get(k, "").strip() for k in ("email", "phone", "location") if profile.get(k, "").strip()
    ]
    if contact_parts:
        paragraphs.append(" | ".join(contact_parts))
    headline = profile.get("headline", "").strip()
    if headline:
        paragraphs.append(headline)

    summary = profile.get("summary", "").strip()
    if summary:
        paragraphs.append("Professional Summary")
        paragraphs.append(summary)

    skills = profile.get("skills", [])
    if skills:
        paragraphs.append("Skills")
        paragraphs.append(", ".join(skills))

    exp_list = profile.get("experience", [])
    if exp_list:
        paragraphs.append("Experience")
        for entry in exp_list:
            header_parts = [entry.get(k, "").strip() for k in ("title", "company", "dates") if entry.get(k, "").strip()]
            if header_parts:
                paragraphs.append(" - ".join(header_parts))
            for bullet in entry.get("bullets", []):
                if bullet.strip():
                    paragraphs.append(bullet.strip())

    proj_list = profile.get("projects", [])
    if proj_list:
        paragraphs.append("Projects")
        for entry in proj_list:
            p_name = entry.get("name", "").strip()
            p_desc = entry.get("description", "").strip()
            if p_name and p_desc and p_name != p_desc:
                paragraphs.append(f"{p_name}: {p_desc}")
            elif p_desc:
                paragraphs.append(p_desc)
            elif p_name:
                paragraphs.append(p_name)

    edu_list = profile.get("education", [])
    if edu_list:
        paragraphs.append("Education")
        for entry in edu_list:
            parts = [entry.get(k, "").strip() for k in ("degree", "institution", "graduation_year") if entry.get(k, "").strip()]
            if parts:
                paragraphs.append(" - ".join(parts))

    certs = profile.get("certifications", [])
    if certs:
        paragraphs.append("Certifications")
        for cert in certs:
            if cert.strip():
                paragraphs.append(cert.strip())

    return paragraphs


@app.get("/api/profile")
def get_candidate_profile():
    with Session(engine, expire_on_commit=False) as db:
        record = db.get(CandidateProfileRecord, "default")
        if record is not None:
            return json.loads(record.payload)
    master = app.state.master_resume
    if master and master.get("paragraphs"):
        profile = extract_profile_from_paragraphs(master["paragraphs"])
        return profile
    return get_default_candidate_profile()


@app.put("/api/profile")
@app.post("/api/profile")
def save_candidate_profile(body: dict):
    profile = sanitize_candidate_profile(body)
    payload_json = json.dumps(profile, ensure_ascii=False)
    now = datetime.now(timezone.utc).isoformat()
    with Session(engine, expire_on_commit=False) as db:
        record = db.get(CandidateProfileRecord, "default")
        if record is None:
            record = CandidateProfileRecord(id="default", payload=payload_json, updated_at=now)
            db.add(record)
        else:
            record.payload = payload_json
            record.updated_at = now
        db.commit()
    return profile


@app.post("/api/profile/extract")
async def extract_candidate_profile_endpoint(file: UploadFile | None = File(default=None)):
    if file is None:
        master = app.state.master_resume
        if master is None or not master.get("paragraphs"):
            raise HTTPException(status_code=422, detail="No resume was uploaded and no master resume is available.")
        paragraphs = list(master["paragraphs"])
    else:
        filename = clean_filename(file.filename or "")
        extension = Path(filename).suffix.casefold()
        if extension not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=415, detail="Unsupported file type. Upload a PDF or DOCX resume.")
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="Resume uploads are limited to 10 MB.")
        if not content:
            raise HTTPException(status_code=422, detail="The uploaded file is empty.")
        with tempfile.NamedTemporaryFile(suffix=extension, delete=False) as tmp:
            tmp.write(content)
            tmp_path = Path(tmp.name)
        try:
            paragraphs, _ = extract_resume(tmp_path, extension)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()
    profile = extract_profile_from_paragraphs(paragraphs)
    save_candidate_profile(profile)
    return profile


@app.get("/api/jobs/{job_id}")
def get_job_status(job_id: str):
    with analysis_jobs_lock:
        job = analysis_jobs.get(job_id)
        if job:
            job_copy = dict(job)
            if job_copy.get("status") == "completed":
                try:
                    record, payload = get_workflow(job_copy["workflow_id"])
                    job_copy["workflow"] = response_payload(record, payload)
                except HTTPException:
                    pass
            return job_copy

    try:
        record, payload = get_workflow(job_id)
        return {
            "job_id": job_id,
            "workflow_id": job_id,
            "status": payload.get("status", "completed"),
            "progress": payload.get("progress", 100),
            "message": "Analysis complete." if payload.get("status") == "completed" else "Processing...",
            "error": payload.get("job_error"),
            "workflow": response_payload(record, payload),
        }
    except HTTPException:
        raise HTTPException(status_code=404, detail="Analysis job not found.")


@app.get("/api/workflows/{workflow_id}/status")
def get_workflow_status(workflow_id: str):
    return get_job_status(workflow_id)


@app.get("/api/health")
def health():
    master = app.state.master_resume
    return {
        "status": "ok",
        "mode": "ollama" if OLLAMA_API_URL else "demo",
        "ai_configured": bool(OLLAMA_API_URL),
        "ai_model": OLLAMA_MODEL if OLLAMA_API_URL else None,
        "ollama_timeout_seconds": OLLAMA_TIMEOUT_SECONDS,
        "default_resume_available": master is not None,
        "default_resume_filename": master["filename"] if master else None,
    }


@app.get("/api/resume/default")
def download_default_resume():
    master = app.state.master_resume
    if master is None:
        raise HTTPException(status_code=404, detail="No saved default resume is configured.")
    return FileResponse(
        master["path"],
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=master["filename"],
    )


@app.get("/api/workflows")
def list_workflows():
    with Session(engine, expire_on_commit=False) as db:
        records = db.scalars(select(WorkflowRecord).order_by(WorkflowRecord.created_at.desc())).all()
        result = []
        for record in records:
            payload = json.loads(record.payload)
            result.append({
                "id": record.id,
                "filename": record.filename,
                "job_title": payload["job_title"],
                "company_name": payload["company_name"],
                "created_at": record.created_at,
                "score": payload["analysis"]["score"],
                "finalized": payload["finalized"],
            })
    return result


@app.post("/api/analyze")
async def analyze_resume(
    file: UploadFile | None = File(default=None),
    job_description: str = Form(..., min_length=20, max_length=40000),
    job_title: str = Form(default="", max_length=160),
    company_name: str = Form(default="", max_length=160),
    use_profile: bool = Form(default=False),
    background: bool = Form(default=False),
):
    master = app.state.master_resume
    workflow_id = uuid.uuid4().hex
    folder = DATA_ROOT / workflow_id
    folder.mkdir(parents=True, exist_ok=False)

    if file is None:
        if use_profile:
            # Load candidate profile
            profile_data = None
            with Session(engine, expire_on_commit=False) as db:
                p_rec = db.get(CandidateProfileRecord, "default")
                if p_rec:
                    profile_data = json.loads(p_rec.payload)
            if not profile_data and master and master.get("paragraphs"):
                profile_data = extract_profile_from_paragraphs(master["paragraphs"])
            if not profile_data or not profile_data.get("full_name"):
                folder.rmdir()
                raise HTTPException(
                    status_code=422,
                    detail="No candidate profile exists yet. Create or save your profile first before tailoring.",
                )
            paragraphs = profile_to_paragraphs(profile_data)
            clean_name = clean_filename(profile_data.get("full_name") or "candidate")
            filename = f"{clean_name}-profile.docx"
            extension = ".docx"
            source_path = folder / f"original{extension}"
            create_docx_from_pdf_text(paragraphs, source_path)
            page_count = 1
        elif master is not None:
            filename = master["filename"]
            extension = ".docx"
            source_path = folder / f"original{extension}"
            try:
                shutil.copy2(master["path"], source_path)
            except OSError as exc:
                shutil.rmtree(folder)
                raise HTTPException(status_code=500, detail="The default master resume could not be copied for this workflow.") from exc
            paragraphs = list(master["paragraphs"])
            page_count = master["page_count"]
        else:
            # Check if candidate profile is saved in DB
            profile_data = None
            with Session(engine, expire_on_commit=False) as db:
                p_rec = db.get(CandidateProfileRecord, "default")
                if p_rec:
                    profile_data = json.loads(p_rec.payload)
            if profile_data and profile_data.get("full_name"):
                paragraphs = profile_to_paragraphs(profile_data)
                clean_name = clean_filename(profile_data.get("full_name") or "candidate")
                filename = f"{clean_name}-profile.docx"
                extension = ".docx"
                source_path = folder / f"original{extension}"
                create_docx_from_pdf_text(paragraphs, source_path)
                page_count = 1
            else:
                folder.rmdir()
                raise HTTPException(
                    status_code=422,
                    detail=f"No resume was uploaded and the default master resume is missing. Place a DOCX at {MASTER_RESUME_PATH} or create a Candidate Profile.",
                )
    else:
        filename = clean_filename(file.filename or "")
        extension = Path(filename).suffix.casefold()
        if extension not in ALLOWED_EXTENSIONS:
            folder.rmdir()
            raise HTTPException(status_code=415, detail="Unsupported file type. Upload a PDF or DOCX resume.")
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(content) > MAX_UPLOAD_BYTES:
            folder.rmdir()
            raise HTTPException(status_code=413, detail="Resume uploads are limited to 10 MB.")
        if not content:
            folder.rmdir()
            raise HTTPException(status_code=422, detail="The selected resume file is empty.")
        if extension == ".pdf" and not content.startswith(b"%PDF-"):
            folder.rmdir()
            raise HTTPException(status_code=415, detail="The selected file does not appear to be a valid PDF.")
        if extension == ".docx" and not content.startswith(b"PK"):
            folder.rmdir()
            raise HTTPException(status_code=415, detail="The selected file does not appear to be a valid DOCX.")
        source_path = folder / f"original{extension}"
        try:
            source_path.write_bytes(content)
        except OSError as exc:
            shutil.rmtree(folder)
            raise HTTPException(status_code=500, detail="The uploaded resume could not be saved. Check available disk space and try again.") from exc
        try:
            paragraphs, page_count = await run_in_threadpool(extract_resume, source_path, extension)
        except HTTPException:
            shutil.rmtree(folder)
            raise

    if not job_description.strip():
        shutil.rmtree(folder)
        raise HTTPException(status_code=422, detail="Paste the complete job description before analyzing.")

    original_text = "\n".join(paragraphs).strip()
    if len(original_text) < 20:
        shutil.rmtree(folder)
        raise HTTPException(status_code=422, detail="This resume contains too little readable text to analyze.")

    # Protection against duplicate running jobs
    req_hash = hashlib.sha256(
        f"{original_text}:{job_description.strip()}:{job_title.strip()}:{company_name.strip()}".encode("utf-8")
    ).hexdigest()

    with analysis_jobs_lock:
        stale_keys = []
        for jid, existing_job in list(analysis_jobs.items()):
            if existing_job.get("request_hash") == req_hash:
                status = existing_job.get("status")
                w_id = existing_job.get("workflow_id")
                try:
                    rec, pld = get_workflow(w_id)
                    shutil.rmtree(folder)
                    return response_payload(rec, pld)
                except HTTPException:
                    if status in ("pending", "processing"):
                        shutil.rmtree(folder)
                        return {
                            "id": w_id,
                            "job_id": existing_job.get("job_id", w_id),
                            "workflow_id": w_id,
                            "status": status,
                            "progress": existing_job.get("progress", 50),
                            "message": existing_job.get("message", "Analysis is currently running."),
                        }
                    stale_keys.append(jid)
        for jid in stale_keys:
            analysis_jobs.pop(jid, None)

    try:
        analysis = score_resume(original_text, job_description, job_title=job_title.strip())
    except HTTPException:
        shutil.rmtree(folder)
        raise

    limitations = [
        "Match scores are estimated JD vocabulary coverage, not an actual ATS score.",
        "Every suggested change requires your explicit approval. Resume text is never changed automatically.",
    ]
    if extension == ".pdf":
        limitations.append("PDF text and the extracted-content preview are available, but PDF edits cannot preserve original layout. DOCX/PDF exports from a PDF source use a basic text layout.")
    else:
        limitations.append("DOCX exports update existing paragraph text in place, reject replacements longer than the source text budget, and must render to one page. Precise DOCX-to-PDF pagination checks depend on LibreOffice being installed.")

    now_iso = datetime.now(timezone.utc).isoformat()

    if not OLLAMA_API_URL:
        # Synchronous demo mode
        suggestions = make_demo_suggestions(paragraphs, job_description)
        payload = {
            "extension": extension,
            "job_title": job_title.strip(),
            "company_name": company_name.strip(),
            "job_description": job_description.strip(),
            "paragraphs": paragraphs,
            "page_count": page_count,
            "analysis": analysis,
            "confirmed_tools": [],
            "skipped_tools": [],
            "suggestions": suggestions,
            "instructions": "",
            "suggestion_warnings": [],
            "demo_mode": True,
            "limitations": limitations,
            "finalized": False,
            "final_score": None,
            "final_score_details": None,
            "final_text": None,
            "layout_review": [],
            "status": "completed",
            "progress": 100,
        }
        update_tailored_state(payload)
        record = WorkflowRecord(
            id=workflow_id,
            payload=json.dumps(payload, ensure_ascii=False),
            filename=filename,
            extension=extension,
            created_at=now_iso,
        )
        try:
            with Session(engine, expire_on_commit=False) as db:
                db.add(record)
                db.commit()
        except SQLAlchemyError as exc:
            shutil.rmtree(folder)
            raise HTTPException(status_code=500, detail="The tailoring session could not be saved.") from exc

        with analysis_jobs_lock:
            analysis_jobs[workflow_id] = {
                "job_id": workflow_id,
                "workflow_id": workflow_id,
                "status": "completed",
                "progress": 100,
                "message": "Analysis complete.",
                "error": None,
                "created_at": now_iso,
                "updated_at": now_iso,
                "request_hash": req_hash,
            }
        return response_payload(record, payload)

    if background:
        # Local Ollama background mode
        initial_payload = {
            "extension": extension,
            "job_title": job_title.strip(),
            "company_name": company_name.strip(),
            "job_description": job_description.strip(),
            "paragraphs": paragraphs,
            "page_count": page_count,
            "analysis": analysis,
            "confirmed_tools": [],
            "skipped_tools": [],
            "suggestions": [],
            "instructions": "",
            "suggestion_warnings": [],
            "demo_mode": False,
            "limitations": limitations,
            "finalized": False,
            "final_score": None,
            "final_score_details": None,
            "final_text": None,
            "layout_review": [],
            "status": "processing",
            "progress": 25,
        }
        update_tailored_state(initial_payload)
        record = WorkflowRecord(
            id=workflow_id,
            payload=json.dumps(initial_payload, ensure_ascii=False),
            filename=filename,
            extension=extension,
            created_at=now_iso,
        )
        try:
            with Session(engine, expire_on_commit=False) as db:
                db.add(record)
                db.commit()
        except SQLAlchemyError as exc:
            shutil.rmtree(folder)
            raise HTTPException(status_code=500, detail="The tailoring session could not be saved.") from exc

        with analysis_jobs_lock:
            analysis_jobs[workflow_id] = {
                "job_id": workflow_id,
                "workflow_id": workflow_id,
                "status": "processing",
                "progress": 25,
                "message": "Resume analyzed. Connecting to local Ollama (llama3.2)...",
                "error": None,
                "created_at": now_iso,
                "updated_at": now_iso,
                "request_hash": req_hash,
            }

        def _bg_ollama_analysis():
            try:
                suggestions = []
                model_warnings = []

                if call_ollama_suggestions != ORIGINAL_CALL_OLLAMA_SUGGESTIONS:
                    with analysis_jobs_lock:
                        if workflow_id in analysis_jobs:
                            analysis_jobs[workflow_id]["progress"] = 50
                            analysis_jobs[workflow_id]["message"] = "Generating tailored suggestions with local Ollama..."
                            analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()
                    suggestions, model_warnings = call_ollama_suggestions(original_text, job_description)
                else:
                    section_map = resume_sections(paragraphs)

                    # 1. Tailor Skills section - strictly additive, preserving 100% of candidate's skills
                    if "Skills" in section_map and section_map["Skills"]:
                        with analysis_jobs_lock:
                            if workflow_id in analysis_jobs:
                                analysis_jobs[workflow_id]["progress"] = 40
                                analysis_jobs[workflow_id]["message"] = "Tailoring Skills to match job requirements..."
                                analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()
                        skills_items = section_map["Skills"]
                        jd_words = extract_keywords(job_description)
                        all_resume_words = extract_keywords(original_text)
                        missing_from_resume = sorted(jd_words - all_resume_words)
                        res = build_additive_skills_suggestion(skills_items, missing_from_resume)
                        if res:
                            res["source"] = "ollama"
                            suggestions.append(res)

                    # 2. Tailor Experience section
                    if "Experience" in section_map and section_map["Experience"]:
                        with analysis_jobs_lock:
                            if workflow_id in analysis_jobs:
                                analysis_jobs[workflow_id]["progress"] = 70
                                analysis_jobs[workflow_id]["message"] = "Tailoring Experience bullets with quantifiable impact..."
                                analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()
                        exp_item = max(section_map["Experience"], key=lambda it: len(it["text"]))
                        res, w = call_ollama_single_section("Experience", exp_item["text"], job_description)
                        if res:
                            res["paragraph_index"] = exp_item["index"]
                            suggestions.append(res)
                        elif w:
                            model_warnings.extend(w)

                    # 3. Tailor Professional Summary if available and < 3 suggestions
                    if len(suggestions) < 3 and "Professional Summary" in section_map and section_map["Professional Summary"]:
                        with analysis_jobs_lock:
                            if workflow_id in analysis_jobs:
                                analysis_jobs[workflow_id]["progress"] = 85
                                analysis_jobs[workflow_id]["message"] = "Tailoring Professional Summary headline..."
                                analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()
                        summary_item = section_map["Professional Summary"][0]
                        res, w = call_ollama_single_section("Professional Summary", summary_item["text"], job_description)
                        if res:
                            res["paragraph_index"] = summary_item["index"]
                            suggestions.append(res)
                        elif w:
                            model_warnings.extend(w)

                suggestion_warnings = []
                if not suggestions:
                    suggestions, suggestion_warnings = fallback_to_manual_suggestions([], paragraphs, job_description)
                    suggestion_warnings = model_warnings + suggestion_warnings
                else:
                    suggestion_warnings = model_warnings

                with analysis_jobs_lock:
                    if workflow_id in analysis_jobs:
                        analysis_jobs[workflow_id]["progress"] = 95
                        analysis_jobs[workflow_id]["message"] = "Validating suggestions against one-page layout constraints..."
                        analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()

                with Session(engine, expire_on_commit=False) as db:
                    rec = db.get(WorkflowRecord, workflow_id)
                    if rec:
                        pld = json.loads(rec.payload)
                        pld["suggestions"] = suggestions
                        pld["suggestion_warnings"] = suggestion_warnings
                        pld["demo_mode"] = False
                        pld["status"] = "completed"
                        pld["progress"] = 100
                        update_tailored_state(pld)
                        rec.payload = json.dumps(pld, ensure_ascii=False)
                        db.commit()

                with analysis_jobs_lock:
                    if workflow_id in analysis_jobs:
                        analysis_jobs[workflow_id]["status"] = "completed"
                        analysis_jobs[workflow_id]["progress"] = 100
                        analysis_jobs[workflow_id]["message"] = "Analysis complete."
                        analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()

            except Exception as exc:
                err_detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
                is_timeout = (
                    (isinstance(exc, HTTPException) and exc.status_code == 504)
                    or "timeout" in err_detail.lower()
                    or "504" in err_detail
                    or "finish generating" in err_detail.lower()
                )
                fallback_sugs, fallback_warns = fallback_to_manual_suggestions(
                    [], paragraphs, job_description
                )
                if is_timeout:
                    msg = (
                        f"AI analysis timed out while waiting for local response ({err_detail}). "
                        "Fallback suggestions provided for manual tailoring so you can continue editing."
                    )
                else:
                    msg = (
                        f"AI analysis encountered an issue ({err_detail}). "
                        "Fallback suggestions provided for manual tailoring so you can continue editing."
                    )
                suggestion_warnings = [msg] + fallback_warns

                with Session(engine, expire_on_commit=False) as db:
                    rec = db.get(WorkflowRecord, workflow_id)
                    if rec:
                        pld = json.loads(rec.payload)
                        pld["suggestions"] = fallback_sugs
                        pld["suggestion_warnings"] = suggestion_warnings
                        pld["demo_mode"] = False
                        pld["status"] = "completed"
                        pld["progress"] = 100
                        pld["job_error"] = None
                        pld["limitations"].append(
                            "AI generation timed out or was unavailable. Manual review cards are provided."
                        )
                        update_tailored_state(pld)
                        rec.payload = json.dumps(pld, ensure_ascii=False)
                        db.commit()

                with analysis_jobs_lock:
                    if workflow_id in analysis_jobs:
                        analysis_jobs[workflow_id]["status"] = "completed"
                        analysis_jobs[workflow_id]["progress"] = 100
                        analysis_jobs[workflow_id]["message"] = (
                            "Analysis completed (switched to manual suggestions due to slow AI generation)."
                            if is_timeout
                            else "Analysis completed (switched to manual suggestions)."
                        )
                        analysis_jobs[workflow_id]["error"] = None
                        analysis_jobs[workflow_id]["updated_at"] = datetime.now(timezone.utc).isoformat()

        threading.Thread(target=_bg_ollama_analysis, daemon=True).start()
        return response_payload(record, initial_payload)

    # Synchronous Ollama execution
    try:
        suggestions = []
        model_warnings = []
        if call_ollama_suggestions != ORIGINAL_CALL_OLLAMA_SUGGESTIONS:
            suggestions, model_warnings = await run_in_threadpool(
                call_ollama_suggestions, original_text, job_description
            )
        else:
            section_map = resume_sections(paragraphs)
            if "Skills" in section_map and section_map["Skills"]:
                skills_item = section_map["Skills"][0]
                res, w = await run_in_threadpool(call_ollama_single_section, "Skills", skills_item["text"], job_description)
                if res:
                    res["paragraph_index"] = skills_item["index"]
                    suggestions.append(res)
                elif w:
                    model_warnings.extend(w)
            if "Experience" in section_map and section_map["Experience"]:
                exp_item = max(section_map["Experience"], key=lambda it: len(it["text"]))
                res, w = await run_in_threadpool(call_ollama_single_section, "Experience", exp_item["text"], job_description)
                if res:
                    res["paragraph_index"] = exp_item["index"]
                    suggestions.append(res)
                elif w:
                    model_warnings.extend(w)
            if len(suggestions) < 3 and "Professional Summary" in section_map and section_map["Professional Summary"]:
                summary_item = section_map["Professional Summary"][0]
                res, w = await run_in_threadpool(call_ollama_single_section, "Professional Summary", summary_item["text"], job_description)
                if res:
                    res["paragraph_index"] = summary_item["index"]
                    suggestions.append(res)
                elif w:
                    model_warnings.extend(w)

        suggestion_warnings = []
        if not suggestions:
            suggestions, suggestion_warnings = fallback_to_manual_suggestions(
                [], paragraphs, job_description
            )
            suggestion_warnings = model_warnings + suggestion_warnings
        else:
            suggestion_warnings = model_warnings
    except HTTPException:
        shutil.rmtree(folder, ignore_errors=True)
        raise

    payload = {
        "extension": extension,
        "job_title": job_title.strip(),
        "company_name": company_name.strip(),
        "job_description": job_description.strip(),
        "paragraphs": paragraphs,
        "page_count": page_count,
        "analysis": analysis,
        "suggestions": suggestions,
        "instructions": "",
        "suggestion_warnings": suggestion_warnings,
        "demo_mode": False,
        "limitations": limitations,
        "finalized": False,
        "final_score": None,
        "final_score_details": None,
        "final_text": None,
        "layout_review": [],
        "status": "completed",
        "progress": 100,
    }
    update_tailored_state(payload)
    record = WorkflowRecord(
        id=workflow_id,
        payload=json.dumps(payload, ensure_ascii=False),
        filename=filename,
        extension=extension,
        created_at=now_iso,
    )
    try:
        with Session(engine, expire_on_commit=False) as db:
            db.add(record)
            db.commit()
    except SQLAlchemyError as exc:
        shutil.rmtree(folder)
        raise HTTPException(status_code=500, detail="The tailoring session could not be saved.") from exc

    with analysis_jobs_lock:
        analysis_jobs[workflow_id] = {
            "job_id": workflow_id,
            "workflow_id": workflow_id,
            "status": "completed",
            "progress": 100,
            "message": "Analysis complete.",
            "error": None,
            "created_at": now_iso,
            "updated_at": now_iso,
            "request_hash": req_hash,
        }

    return response_payload(record, payload)


@app.get("/api/workflows/{workflow_id}")
def get_workflow_details(workflow_id: str):
    record, payload = get_workflow(workflow_id)
    return response_payload(record, payload)


@app.patch("/api/workflows/{workflow_id}/suggestions/{item_id}")
def update_suggestion(workflow_id: str, item_id: str, body: SuggestionAction):
    record, payload = get_workflow(workflow_id)
    suggestion = next((item for item in payload["suggestions"] if item["id"] == item_id), None)
    if suggestion is None:
        raise HTTPException(status_code=404, detail="This suggestion was not found.")
    if body.action == "edit":
        if body.proposed_text is None or not body.proposed_text.strip():
            raise HTTPException(status_code=422, detail="Enter replacement text before saving an edit.")
        suggestion["proposed_text"] = normalize_export_text(body.proposed_text)
        suggestion["status"] = "pending"
        suggestion["is_skills_addition"] = False
    elif body.action == "undo":
        suggestion["proposed_text"] = suggestion["suggested_text"]
        suggestion["status"] = "pending"
        if suggestion["suggested_text"].startswith(suggestion["original_text"].rstrip(" \t\r\n.,;")):
            suggestion["is_skills_addition"] = True
    else:
        if body.action == "approve" and body.proposed_text and body.proposed_text.strip():
            suggestion["proposed_text"] = normalize_export_text(body.proposed_text)
        if body.action == "approve":
            validate_approved_rewrites({
                **payload,
                "suggestions": [
                    {**item, "status": "approved"}
                    if item["id"] == item_id else item
                    for item in payload["suggestions"]
                ],
            })
        suggestion["status"] = "approved" if body.action == "approve" else "rejected"
    invalidate_exports(workflow_id, payload)
    update_tailored_state(payload)
    save_workflow(record, payload)
    return response_payload(record, payload)


@app.post("/api/workflows/{workflow_id}/suggestions/bulk")
def bulk_suggestions(workflow_id: str, body: BulkSuggestionAction):
    record, payload = get_workflow(workflow_id)
    if body.action == "approve":
        validate_approved_rewrites({
            **payload,
            "suggestions": [
                {**suggestion, "status": "approved"}
                for suggestion in payload["suggestions"]
            ],
        })
    for suggestion in payload["suggestions"]:
        suggestion["status"] = "approved" if body.action == "approve" else "rejected"
    invalidate_exports(workflow_id, payload)
    update_tailored_state(payload)
    save_workflow(record, payload)
    return response_payload(record, payload)


@app.post("/api/workflows/{workflow_id}/instructions")
def apply_instructions(workflow_id: str, body: InstructionRequest):
    record, payload = get_workflow(workflow_id)
    instruction = body.instruction.strip()
    suggestion_warnings = []
    if OLLAMA_API_URL:
        suggestions, model_warnings = call_ollama_suggestions(
            "\n".join(payload["paragraphs"]), payload["job_description"], instruction,
        )
        if not suggestions:
            suggestions, suggestion_warnings = fallback_to_manual_suggestions(
                [], payload["paragraphs"], payload["job_description"],
            )
            suggestion_warnings = model_warnings + suggestion_warnings
            for suggestion in suggestions:
                suggestion["reason"] += (
                    " Your custom instruction was saved; apply it by editing this text "
                    "yourself because no safe AI rewrite was returned."
                )
        mode = False
    else:
        suggestions, suggestion_warnings = fallback_to_manual_suggestions(
            [], payload["paragraphs"], payload["job_description"],
        )
        for suggestion in suggestions:
            suggestion["reason"] += f" Your instruction: “{instruction[:240]}”. Demo mode will not rewrite or invent content; edit this suggestion yourself."
        mode = True
    for suggestion in suggestions:
        suggestion["status"] = "pending"
    payload["suggestions"] = suggestions
    payload["instructions"] = instruction
    payload["demo_mode"] = mode
    payload["suggestion_warnings"] = suggestion_warnings
    invalidate_exports(workflow_id, payload)
    update_tailored_state(payload)
    save_workflow(record, payload)
    return response_payload(record, payload)


@app.post("/api/workflows/{workflow_id}/confirm-tool")
def confirm_workflow_tool(workflow_id: str, request: ToolConfirmationRequest):
    record, payload = get_workflow(workflow_id)
    tool_name = request.tool.strip().casefold()

    confirmed = payload.setdefault("confirmed_tools", [])
    skipped = payload.setdefault("skipped_tools", [])

    if request.status == "approved" or request.status == "related":
        if tool_name not in confirmed:
            confirmed.append(tool_name)
        if tool_name in skipped:
            skipped.remove(tool_name)
    elif request.status == "skipped":
        if tool_name not in skipped:
            skipped.append(tool_name)
        if tool_name in confirmed:
            confirmed.remove(tool_name)

    # Automatically synchronize confirmed tool additions into Skills suggestions additively
    canonical_name = get_canonical_skill_name(tool_name)
    tool_cat = get_skill_category_target(tool_name)
    section_map = resume_sections(payload.get("paragraphs", []))
    skills_items = section_map.get("Skills", [])

    target_item = None
    for item in skills_items:
        if tool_cat in item["text"].casefold():
            target_item = item
            break
    if target_item is None and skills_items:
        target_item = skills_items[-1]

    if target_item is not None:
        existing_sugg = None
        for s in payload.setdefault("suggestions", []):
            if s.get("paragraph_index") == target_item["index"]:
                existing_sugg = s
                break

        orig_text = target_item["text"]
        current_proposed = existing_sugg["proposed_text"] if existing_sugg else orig_text

        if request.status in ("approved", "related"):
            if canonical_name.casefold() not in current_proposed.casefold():
                new_proposed = current_proposed.rstrip().rstrip(",;.") + ", " + canonical_name
                if existing_sugg:
                    existing_sugg["proposed_text"] = new_proposed
                    existing_sugg["suggested_text"] = new_proposed
                    existing_sugg["status"] = "approved"
                    existing_sugg["is_skills_addition"] = True
                else:
                    payload["suggestions"].append({
                        "id": suggestion_id("Skills", target_item["index"]),
                        "section": "Skills",
                        "paragraph_index": target_item["index"],
                        "original_text": orig_text,
                        "proposed_text": new_proposed,
                        "suggested_text": new_proposed,
                        "reason": f"Added confirmed requirement {canonical_name} while keeping all original skills.",
                        "requirement": canonical_name,
                        "expected_impact": "Direct match for confirmed technical requirement.",
                        "status": "approved",
                        "source": "manual",
                        "is_skills_addition": True,
                    })
        elif request.status == "skipped":
            if existing_sugg and canonical_name.casefold() in existing_sugg["proposed_text"].casefold():
                cleaned_text = re.sub(
                    rf"(?:,\s*{re.escape(canonical_name)}|{re.escape(canonical_name)}\s*,\s*)",
                    "",
                    existing_sugg["proposed_text"],
                    flags=re.IGNORECASE,
                )
                if len(cleaned_text.strip()) < len(orig_text.strip()):
                    cleaned_text = orig_text
                existing_sugg["proposed_text"] = cleaned_text
                existing_sugg["suggested_text"] = cleaned_text
                if cleaned_text.strip() == orig_text.strip():
                    existing_sugg["status"] = "pending"

    invalidate_exports(workflow_id, payload)

    original_text = "\n".join(payload.get("paragraphs", [])).strip()
    analysis = score_resume(
        original_text,
        payload["job_description"],
        job_title=payload.get("job_title", ""),
        confirmed_tools=confirmed,
        skipped_tools=skipped,
    )
    payload["analysis"] = analysis
    update_tailored_state(payload)
    save_workflow(record, payload)
    return response_payload(record, payload)


@app.get("/api/workflows/{workflow_id}/discovered-tools")
def get_workflow_discovered_tools(workflow_id: str):
    record, payload = get_workflow(workflow_id)
    return {
        "discovered_tools": payload.get("analysis", {}).get("discovered_tools", []),
        "confirmed_tools": payload.get("confirmed_tools", []),
        "skipped_tools": payload.get("skipped_tools", []),
        "tier1_role_match": payload.get("analysis", {}).get("tier1_role_match", {}),
        "career_stage": payload.get("analysis", {}).get("career_stage", "Fresher / Entry-Level"),
    }


@app.post("/api/workflows/{workflow_id}/tailor-section")
def tailor_specific_section(workflow_id: str, body: TailorSectionRequest):
    record, payload = get_workflow(workflow_id)
    section = body.section.strip()
    if section not in SECTIONS:
        raise HTTPException(status_code=422, detail=f"Invalid section '{section}'.")

    target_idx = body.paragraph_index
    target_text = ""
    if target_idx is not None and 0 <= target_idx < len(payload["paragraphs"]):
        target_text = payload["paragraphs"][target_idx]
    else:
        sections = resume_sections(payload["paragraphs"])
        if section in sections and sections[section]:
            best = sections[section][0]
            target_idx = best["index"]
            target_text = best["text"]

    if not target_text:
        raise HTTPException(status_code=404, detail=f"Could not locate text for section '{section}'.")

    result, warnings = call_ollama_single_section(
        section=section,
        original_text=target_text,
        job_description=payload["job_description"],
        custom_instruction=body.custom_instruction,
    )

    if not result:
        err_msg = warnings[0] if warnings else "Could not generate tailored suggestion for this section."
        raise HTTPException(status_code=502, detail=err_msg)

    result["paragraph_index"] = target_idx

    existing_idx = next(
        (i for i, s in enumerate(payload["suggestions"]) if s["section"] == section and s.get("paragraph_index") == target_idx),
        None,
    )
    if existing_idx is not None:
        payload["suggestions"][existing_idx] = result
    else:
        payload["suggestions"].append(result)

    invalidate_exports(workflow_id, payload)
    update_tailored_state(payload)
    save_workflow(record, payload)
    return response_payload(record, payload)


def build_final_resume(workflow_id: str, record: WorkflowRecord, payload: dict) -> None:
    unresolved = [item["section"] for item in payload["suggestions"] if item["status"] == "pending"]
    if unresolved:
        raise HTTPException(status_code=409, detail="Review every suggestion first: approve or reject each one before generating the final resume.")
    validate_approved_rewrites(payload)
    folder = DATA_ROOT / workflow_id
    source_path = folder / f"original{record.extension}"
    docx_path = folder / "tailored.docx"
    pdf_path = folder / "tailored.pdf"
    generate_docx(payload, source_path, docx_path)
    generate_pdf(payload, docx_path, pdf_path)
    if inspect_pdf_layout(pdf_path)["page_count"] != 1:
        for path in (docx_path, pdf_path):
            if path.exists():
                path.unlink()
        raise HTTPException(
            status_code=422,
            detail="The tailored resume must fit on exactly one page. Shorten approved replacements and try again.",
        )
    update_tailored_state(payload)
    payload["layout_review"] = check_layout(payload, source_path, pdf_path)
    payload["finalized"] = True
    save_workflow(record, payload)


@app.post("/api/workflows/{workflow_id}/finalize")
def finalize_workflow(workflow_id: str):
    record, payload = get_workflow(workflow_id)
    build_final_resume(workflow_id, record, payload)
    return response_payload(record, payload)


@app.get("/api/workflows/{workflow_id}/original")
def download_original(workflow_id: str):
    record, _ = get_workflow(workflow_id)
    source = DATA_ROOT / workflow_id / f"original{record.extension}"
    return FileResponse(source, media_type="application/pdf" if record.extension == ".pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=record.filename)


@app.get("/api/workflows/{workflow_id}/download/{file_format}")
def download_final(workflow_id: str, file_format: Literal["pdf", "docx"]):
    record, payload = get_workflow(workflow_id)
    if not payload["finalized"]:
        raise HTTPException(status_code=409, detail="Approve or reject all suggestions before downloading the tailored resume.")
    filename = f"{Path(record.filename).stem}-resumefit.{file_format}"
    media_type = "application/pdf" if file_format == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return FileResponse(DATA_ROOT / workflow_id / f"tailored.{file_format}", media_type=media_type, filename=filename)


@app.post("/api/download/{file_format}")
def download_resume(file_format: Literal["pdf", "docx"], body: DownloadRequest):
    record, payload = get_workflow(body.workflow_id)
    if any(item["status"] == "pending" for item in payload["suggestions"]):
        raise HTTPException(
            status_code=409,
            detail="Approve or reject every suggestion before downloading the tailored resume.",
        )
    validate_approved_rewrites(payload)

    lines = normalize_download_text(body.text)
    expected_lines = normalize_download_text(
        "\n".join(approved_paragraphs(payload))
    )
    if not lines or lines != expected_lines:
        raise HTTPException(
            status_code=409,
            detail="The preview text is out of date or does not match approved changes. Refresh the workflow and try again.",
        )

    folder = DATA_ROOT / body.workflow_id
    source_path = folder / f"original{record.extension}"
    if not source_path.is_file():
        raise HTTPException(status_code=404, detail="The original resume template is unavailable.")
    media_type = (
        "application/pdf"
        if file_format == "pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    filename = f"{Path(record.filename).stem}-resumefit.{file_format}"
    try:
        with tempfile.TemporaryDirectory(prefix="resumefit-export-", dir=folder) as temporary_folder:
            temporary_path = Path(temporary_folder)
            docx_path = temporary_path / "tailored.docx"
            pdf_path = temporary_path / "tailored.pdf"
            generate_docx(payload, source_path, docx_path)
            generate_pdf(payload, docx_path, pdf_path)
            if inspect_pdf_layout(pdf_path)["page_count"] != 1:
                raise HTTPException(
                    status_code=422,
                    detail="The tailored resume must fit on exactly one page. Shorten the approved text and try again.",
                )
            output_path = pdf_path if file_format == "pdf" else docx_path
            content = output_path.read_bytes()
    except OSError as exc:
        raise HTTPException(status_code=500, detail="The tailored resume could not be compiled.") from exc

    return StreamingResponse(
        iter([content]),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/workflows/{workflow_id}/download/{file_format}")
def export_final(
    workflow_id: str,
    file_format: Literal["pdf", "docx"],
    body: ResumeExportRequest,
):
    record, payload = get_workflow(workflow_id)
    submitted = {item.id: item for item in body.suggestions}
    stored_ids = {item["id"] for item in payload["suggestions"]}
    if len(submitted) != len(body.suggestions) or not set(submitted.keys()).issubset(stored_ids):
        raise HTTPException(
            status_code=422,
            detail="The submitted suggestion state does not match this workflow. Refresh the review and try again.",
        )
    if any(item.status == "pending" for item in body.suggestions):
        raise HTTPException(
            status_code=409,
            detail="Approve or reject every suggestion before downloading the tailored resume.",
        )

    for suggestion in payload["suggestions"]:
        if suggestion["id"] in submitted:
            state = submitted[suggestion["id"]]
            suggestion["status"] = state.status
            suggestion["proposed_text"] = normalize_export_text(state.proposed_text)

    validate_docx_text_budget(payload)
    validate_approved_rewrites(payload)
    invalidate_exports(workflow_id, payload)
    update_tailored_state(payload)
    save_workflow(record, payload)
    build_final_resume(workflow_id, record, payload)

    filename = f"{Path(record.filename).stem}-resumefit.{file_format}"
    media_type = (
        "application/pdf"
        if file_format == "pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    return FileResponse(
        DATA_ROOT / workflow_id / f"tailored.{file_format}",
        media_type=media_type,
        filename=filename,
    )


@app.delete("/api/workflows/{workflow_id}")
def delete_workflow(workflow_id: str):
    with Session(engine, expire_on_commit=False) as db:
        record = db.get(WorkflowRecord, workflow_id)
        if record is None:
            raise HTTPException(status_code=404, detail="This tailoring session was not found.")
        db.delete(record)
        db.commit()
    folder = DATA_ROOT / workflow_id
    if folder.parent.resolve() != DATA_ROOT.resolve():
        raise HTTPException(status_code=500, detail="The stored resume location is invalid; files were not deleted.")
    with analysis_jobs_lock:
        analysis_jobs.pop(workflow_id, None)
    if folder.exists():
        shutil.rmtree(folder)
    return Response(status_code=204)


frontend_build = ROOT / "dist"
if frontend_build.is_dir():
    from fastapi.staticfiles import StaticFiles

    app.mount("/", StaticFiles(directory=frontend_build, html=True), name="frontend")
